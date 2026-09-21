from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.models.planilla import PlanillaOutput
from app.models.practice_planilla import PracticePlanillaOutput
from app.routers.billing_publication import _rows_from_calculation_snapshot
from app.services.monetary_snapshot import build_calculation_snapshot, calculation_snapshot_rows
from app.services.payment_overrides import normalize_payment_overrides
from app.services.planilla_generator import (
    DATA_ROW_START,
    PlanillaGenerator,
    PlanillaRow,
    _build_month_blocks,
)
from app.services.practice_planilla_generator import PracticePlanillaGenerator


def _published_row(activity_kind: str) -> PlanillaRow:
    return PlanillaRow(
        teacher_ci=f"PUBLISHED-{activity_kind.upper()}",
        designation_id=7,
        teacher_name=f"Published {activity_kind.title()} Teacher",
        email=None,
        phone=None,
        subject=f"Published {activity_kind.title()}",
        semester="1",
        group_code="A",
        teacher_type=None,
        gender=None,
        sap_code=None,
        invoice_retention=None,
        account_number=None,
        academic_level=None,
        profession=None,
        specialty=None,
        bank=None,
        base_monthly_hours=10,
        payable_hours=10,
        total_hours=10,
        total_theory_hours=10,
        rate_per_hour=70 if activity_kind == "theory" else 50,
        calculated_payment=700 if activity_kind == "theory" else 500,
        final_payment=700 if activity_kind == "theory" else 500,
        source_kind="published",
        source_id=7,
        source_key="published:7",
        published_schedule_assignment_id=7,
        publication_id=3,
        published_block_id=5,
        activity_kind=activity_kind,
        effective_from=date(2027, 3, 1),
        effective_to=date(2027, 3, 31),
    )


@pytest.mark.parametrize(
    ("generator_class", "output_model", "activity_kind", "sheet_name", "override_amount"),
    [
        (PlanillaGenerator, PlanillaOutput, "theory", "Planilla", Decimal("321.45")),
        (
            PracticePlanillaGenerator,
            PracticePlanillaOutput,
            "practice",
            "Planilla Prácticas",
            Decimal("210.25"),
        ),
    ],
)
def test_published_override_is_identical_in_workbook_snapshot_total_and_billing(
    db_session,
    tmp_path,
    monkeypatch,
    generator_class,
    output_model,
    activity_kind,
    sheet_name,
    override_amount,
):
    row = _published_row(activity_kind)
    generator = generator_class(output_dir=str(tmp_path))
    if activity_kind == "theory":
        monkeypatch.setattr(generator, "_build_planilla_data", lambda *args, **kwargs: ([row], [], []))
    else:
        monkeypatch.setattr(generator, "_build_planilla_data", lambda *args, **kwargs: ([row], []))
    overrides = {f"{row.teacher_ci}:published:7": override_amount}

    result = generator.generate(
        db_session,
        month=3,
        year=2027,
        payment_overrides=overrides,
        discount_mode="full",
    )

    blocks = _build_month_blocks(3, 2027, None, None)
    adjusted_column = generator._get_summary_cols(blocks)["pago_ajustado"]
    worksheet = load_workbook(result.file_path, data_only=True)[sheet_name]
    assert Decimal(str(worksheet.cell(DATA_ROW_START, adjusted_column).value)) == override_amount

    stored = db_session.query(output_model).filter_by(month=3, year=2027).one()
    snapshot = stored.calculation_snapshot
    assert Decimal(snapshot["designations"][0]["amount"]) == override_amount
    assert Decimal(snapshot["total"]) == override_amount
    assert Decimal(str(stored.total_payment)) == override_amount
    assert Decimal(str(result.total_payment)) == override_amount
    assert stored.payment_overrides_json == {
        f"{row.teacher_ci}:published:7": format(override_amount, ".2f")
    }

    billing_rows = _rows_from_calculation_snapshot(snapshot, stored.total_payment)
    assert len(billing_rows) == 1
    assert Decimal(str(billing_rows[0].final_payment)) == override_amount
    assert billing_rows[0].has_admin_override is True


def test_legacy_generator_row_override_key_remains_compatible():
    row = _published_row("theory")
    row.source_kind = "legacy"
    row.source_id = 7
    row.source_key = "legacy:7"
    row.published_schedule_assignment_id = None
    overrides = normalize_payment_overrides({f"{row.teacher_ci}:7": "123.45"})

    assert PlanillaGenerator()._get_row_override(row, overrides, [row]) == Decimal("123.45")
    assert PracticePlanillaGenerator()._get_row_override(row, overrides, [row]) == Decimal("123.45")


def test_legacy_override_does_not_mark_same_id_published_snapshot_row():
    legacy = _published_row("theory")
    legacy.source_kind = "legacy"
    legacy.source_key = "legacy:7"
    legacy.published_schedule_assignment_id = None
    published = _published_row("practice")
    published.teacher_ci = legacy.teacher_ci
    overrides = {f"{legacy.teacher_ci}:7": Decimal("123.45")}
    snapshot = build_calculation_snapshot(
        rows=[legacy, published],
        row_amounts=[Decimal("123.45"), Decimal("500.00")],
        month=3,
        year=2027,
        start_date=None,
        end_date=None,
        discount_mode="full",
        payment_overrides=overrides,
        excluded_days=[],
    )

    immutable_rows = calculation_snapshot_rows(snapshot, snapshot["total"])
    billing_rows = _rows_from_calculation_snapshot(snapshot, snapshot["total"])
    assert [row.has_admin_override for row in immutable_rows] == [True, False]
    assert [row.has_admin_override for row in billing_rows] == [True, False]
