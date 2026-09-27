from __future__ import annotations

import json
import hashlib
from datetime import date, time
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from fastapi.testclient import TestClient

from app.main import app
from app.models.academic_management import (
    AcademicGroup, AcademicProgram, AcademicSubject, Classroom, SubjectOffering,
    TeacherAvailability,
)
from app.models.teacher import Teacher
from app.services.designation_bootstrap_preview import (
    _schedule,
    build_designation_bootstrap_preview,
)
from app.services.designation_bootstrap_resolution import (
    DesignationBootstrapResolutionError,
    RESOLUTION_TOKEN_TTL_SECONDS,
    _mask_ci,
    build_designation_bootstrap_alias_artifact,
    build_designation_bootstrap_resolution_context,
)


def _bytes(workbook: Workbook) -> bytes:
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _official(
    *, theory_schedule="LUNES: 08:00 AM-09:30 AM",
    practice_schedule="MARTES: 10:00 AM-11:30 AM", duplicate=False,
    duplicate_header=False, merged=False, blank_row=False,
    practice_name="Private Practice Name", practice_subject="Subject A PRACTICA",
    include_theory=True, include_practice=True,
):
    workbook = Workbook()
    theory = workbook.active
    theory.title = "DOCENTES TEORICOS"
    theory.append(["N°", None, "MATERIAS", "SEMESTRE", "M-04", "AULA", "HORARIO", "CARGA HORARIA", "MES", "SEMANA", "REAL"])
    if include_theory:
        theory.append([1, "Private Theory Name", "Subject A", "PRIMERO", "M-01", "101", theory_schedule, 60, 12, 3, 57])
    if duplicate and include_theory:
        theory.append([2, "Private Theory Name", "Subject A", "PRIMERO", "M-01", "101", theory_schedule, 60, 12, 3, 57])
    if blank_row:
        theory.append([None] * 11)
        theory.cell(theory.max_row + 1, 11, None)
    if duplicate_header:
        theory.cell(1, 12, "HORARIO")
    if merged:
        theory.merge_cells("B2:C2")
    practice = workbook.create_sheet("DOCENTES ASISTENCIALES")
    practice.append(["N°", "DOCENTE", "MATERIAS", "SEMESTRE", "GRUPO", "HORARIO", "CARGA HORARIA ", "REAL"])
    if include_practice:
        practice.append([1, practice_name, practice_subject, "PRIMERO", "T-01", practice_schedule, 60, 12])
    return _bytes(workbook)


def _salary(
    *, hours=(5, 7), ci="990001", subject="Subject A",
    name="Private Practice Name", duplicate_header=False, merged=False,
    extra_assignment=False, semester="1",
):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SEPTIEMBRE 2026"
    sheet.append(["private title"])
    sheet.append([])
    sheet.append([])
    sheet.append([])
    sheet.append([])
    sheet.append([None, "NOMBRE COMPLETO", "Número de teléfono", "Correo electrónico", "Nº C.I.", "MATERIA", "TIPO DE CONTRATO", "SEMESTRE", "TOTAL HORAS", "MONTO TOTAL"])
    if duplicate_header:
        sheet.cell(6, 11, "TOTAL HORAS")
    for index, value in enumerate(hours):
        row_ci = ci[index] if isinstance(ci, tuple) else ci
        row_semester = semester[index] if isinstance(semester, tuple) else semester
        sheet.append([None, name, "70000000", "private@example.test", row_ci, subject, "SERVICE", row_semester, value, value * 50])
    if extra_assignment:
        sheet.append([None, "Orphan Private Name", None, None, "990009", "Subject B", "SERVICE", "2", 4, 200])
    if merged:
        sheet.merge_cells("B7:C7")
    return _bytes(workbook)


def _seed(db_session, *, availability=True, subject=True, classroom=True, practice_teacher=True, groups=True):
    program = AcademicProgram(code="MED", name="Medicine", active=True)
    db_session.add(program)
    db_session.flush()
    academic_subject = None
    if subject:
        academic_subject = AcademicSubject(code="SUB-A", name="Subject A", active=True)
        db_session.add(academic_subject)
        db_session.flush()
        db_session.add(SubjectOffering(
            subject_id=academic_subject.id, program_id=program.id, academic_period="II/2026",
            semester=1, theory_hours=60, practice_hours=60, active=True,
        ))
    db_session.add(Teacher(ci="880001", full_name="Private Theory Name"))
    if groups:
        db_session.add_all([
            AcademicGroup(program_id=program.id, academic_period="II/2026", semester=1, shift="morning", code="M-1", active=True),
            AcademicGroup(program_id=program.id, academic_period="II/2026", semester=1, shift="afternoon", code="T-1", active=True),
        ])
    if classroom:
        db_session.add(Classroom(code="101", name="Room 101", campus="Test", capacity=20, classroom_type="classroom", resources=[], active=True))
    if practice_teacher:
        db_session.add(Teacher(ci="990001", full_name="Private Practice Name"))
    db_session.flush()
    if availability:
        db_session.add_all([
            TeacherAvailability(teacher_ci="880001", academic_period="II/2026", weekday="monday", start_time=time(8), end_time=time(10), active=True),
            TeacherAvailability(teacher_ci="990001", academic_period="II/2026", weekday="tuesday", start_time=time(9), end_time=time(12), active=True),
        ])
    db_session.flush()
    return program


def _preview(db_session, official=None, salary=None, alias=None):
    return build_designation_bootstrap_preview(
        db_session, official_content=official or _official(), salary_content=salary or _salary(),
        academic_period="II/2026", effective_date=date(2026, 8, 21), program_identity="MED",
        alias_content=alias,
    )


def _alias(official: bytes, salary: bytes, *, teachers=None, subjects=None, **overrides):
    payload = {
        "schema_version": "designation-alias-v1",
        "official_sha256": hashlib.sha256(official).hexdigest(),
        "salary_sha256": hashlib.sha256(salary).hexdigest(),
        "academic_period": "II/2026",
        "effective_date": "2026-08-21",
        "teacher_aliases": teachers or {},
        "subject_aliases": subjects or {},
    }
    payload.update(overrides)
    return json.dumps(payload, sort_keys=True).encode()


def _codes(result):
    return {item["code"]: item["count"] for item in result["blockers"]}


def _resolution_sources():
    return (
        _official(
            practice_name="Official Practice Alias",
            practice_subject="Official Clinical Subject PRACTICA",
        ),
        _salary(
            name="Salary Practice Alias",
            subject="Salary Clinical Subject",
            ci="12345678",
        ),
    )


def _resolution_context(official: bytes, salary: bytes, *, now=1_000):
    return build_designation_bootstrap_resolution_context(
        official_content=official,
        salary_content=salary,
        academic_period="II/2026",
        effective_date=date(2026, 8, 21),
        signing_key="test-resolution-signing-key",
        now=now,
    )


def _resolution_selections(context):
    return {
        "teacher_selections": [
            {
                "official_teacher_key": item["official_teacher_key"],
                "candidate_token": item["candidates"][0]["candidate_token"],
            }
            for item in context["teacher_resolutions"]
        ],
        "subject_selections": [
            {
                "salary_subject_key": item["salary_subject_key"],
                "semester": item["semester"],
                "official_subject_key": item["candidates"][0]["official_subject_key"],
            }
            for item in context["subject_resolutions"]
        ],
    }


def test_resolution_context_is_deterministic_minimized_and_masks_ci():
    official, salary = _resolution_sources()
    first = _resolution_context(official, salary)
    repeated = _resolution_context(official, salary)

    assert repeated == first
    assert len(first["teacher_resolutions"]) == 1
    assert len(first["subject_resolutions"]) == 1
    teacher = first["teacher_resolutions"][0]
    subject = first["subject_resolutions"][0]
    assert teacher["official_teacher_display"] == "Official Practice Alias"
    assert teacher["candidates"][0]["salary_teacher_display"] == "Salary Practice Alias"
    assert teacher["candidates"][0]["masked_ci"] == "••••5678"
    assert subject["semester"] == 1
    assert subject["candidates"] == [{
        "official_subject_display": "Official Clinical Subject PRACTICA",
        "official_subject_key": "OFFICIAL CLINICAL SUBJECT",
    }]
    serialized = json.dumps(first, ensure_ascii=False)
    for forbidden in ("12345678", "70000000", "private@example.test", "MONTO TOTAL"):
        assert forbidden not in serialized
    assert set(first) == {
        "resolution_token", "expires_in_seconds", "teacher_resolutions", "subject_resolutions",
    }


@pytest.mark.parametrize("ci", ["Q", "QZ", "QZX", "QZXV"])
def test_short_ci_is_never_disclosed_by_serialized_mask(ci):
    serialized = json.dumps({"masked_ci": _mask_ci(ci)}, ensure_ascii=False)

    assert ci not in serialized


def test_resolution_token_binds_sources_expires_and_rejects_invalid_selection():
    official, salary = _resolution_sources()
    context = _resolution_context(official, salary)
    selections = _resolution_selections(context)

    with pytest.raises(DesignationBootstrapResolutionError, match="stale_resolution_sources"):
        build_designation_bootstrap_alias_artifact(
            official_content=official,
            salary_content=_salary(name="Different Salary Source"),
            academic_period="II/2026",
            effective_date=date(2026, 8, 21),
            resolution_token=context["resolution_token"],
            signing_key="test-resolution-signing-key",
            now=1_001,
            **selections,
        )

    invalid = json.loads(json.dumps(selections))
    invalid["teacher_selections"][0]["candidate_token"] = "0" * 32
    with pytest.raises(DesignationBootstrapResolutionError, match="invalid_alias_selection"):
        build_designation_bootstrap_alias_artifact(
            official_content=official,
            salary_content=salary,
            academic_period="II/2026",
            effective_date=date(2026, 8, 21),
            resolution_token=context["resolution_token"],
            signing_key="test-resolution-signing-key",
            now=1_001,
            **invalid,
        )

    with pytest.raises(DesignationBootstrapResolutionError, match="expired_resolution_token"):
        build_designation_bootstrap_alias_artifact(
            official_content=official,
            salary_content=salary,
            academic_period="II/2026",
            effective_date=date(2026, 8, 21),
            resolution_token=context["resolution_token"],
            signing_key="test-resolution-signing-key",
            now=1_000 + RESOLUTION_TOKEN_TTL_SECONDS + 1,
            **selections,
        )


def test_generated_alias_artifact_is_accepted_by_authoritative_preview(db_session):
    official, salary = _resolution_sources()
    context = _resolution_context(official, salary)
    artifact = build_designation_bootstrap_alias_artifact(
        official_content=official,
        salary_content=salary,
        academic_period="II/2026",
        effective_date=date(2026, 8, 21),
        resolution_token=context["resolution_token"],
        signing_key="test-resolution-signing-key",
        now=1_001,
        **_resolution_selections(context),
    )

    result = build_designation_bootstrap_preview(
        db_session,
        official_content=official,
        salary_content=salary,
        academic_period="II/2026",
        effective_date=date(2026, 8, 21),
        program_identity="MED",
        alias_content=artifact,
    )
    codes = _codes(result)
    assert result["sources"]["alias_present"] is True
    assert "unresolved_teacher_alias" not in codes
    assert "unresolved_subject_alias" not in codes
    assert "teacher_alias_schema_invalid" not in codes
    assert "subject_alias_schema_invalid" not in codes


def test_v2_subject_aliases_preserve_same_source_key_across_semesters(db_session):
    official = _official(
        practice_name="Shared Teacher",
        practice_subject="Official First PRACTICA",
    )
    official_book = load_workbook(BytesIO(official))
    official_book["DOCENTES ASISTENCIALES"].append([
        2, "Shared Teacher", "Official Second PRACTICA", "SEGUNDO", "T-02",
        "MIERCOLES: 10:00 AM-11:30 AM", 60, 12,
    ])
    official = _bytes(official_book)
    salary = _salary(
        name="Shared Teacher",
        subject="Shared Salary Subject",
        ci="990001",
        semester=("1", "2"),
    )
    context = _resolution_context(official, salary)

    artifact = build_designation_bootstrap_alias_artifact(
        official_content=official,
        salary_content=salary,
        academic_period="II/2026",
        effective_date=date(2026, 8, 21),
        resolution_token=context["resolution_token"],
        signing_key="test-resolution-signing-key",
        now=1_001,
        **_resolution_selections(context),
    )
    payload = json.loads(artifact)
    result = _preview(db_session, official=official, salary=salary, alias=artifact)

    assert payload["schema_version"] == "designation-alias-v2"
    assert [(item["salary_subject_key"], item["semester"]) for item in payload["subject_aliases"]] == [
        ("SHARED SALARY SUBJECT", 1),
        ("SHARED SALARY SUBJECT", 2),
    ]
    assert "subject_alias_duplicate_source" not in _codes(result)
    assert result["practice"]["unresolved_subject_alias_count"] == 0


def test_resolution_endpoints_require_admin_and_never_persist(client, db_session):
    official, salary = _resolution_sources()
    files = {
        "official_workbook": ("official.xlsx", official, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        "salary_workbook": ("salary.xlsx", salary, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    }
    data = {"academic_period": "II/2026", "effective_date": "2026-08-21"}

    unauthenticated = TestClient(app)
    response = unauthenticated.post(
        "/api/admin/academic-management/designation-bootstrap/resolution-context",
        files=files,
        data=data,
    )
    unauthenticated.close()
    assert response.status_code == 401

    before = {
        "teachers": db_session.query(Teacher).count(),
        "programs": db_session.query(AcademicProgram).count(),
    }
    response = client.post(
        "/api/admin/academic-management/designation-bootstrap/resolution-context",
        files=files,
        data=data,
    )
    assert response.status_code == 200
    payload = response.json()
    assert "12345678" not in response.text
    assert payload["teacher_resolutions"][0]["candidates"][0]["masked_ci"] == "••••5678"
    assert db_session.query(Teacher).count() == before["teachers"]
    assert db_session.query(AcademicProgram).count() == before["programs"]


def test_alias_artifact_endpoint_rejects_stale_sources_and_preview_accepts_body(client):
    official, salary = _resolution_sources()
    files = {
        "official_workbook": ("official.xlsx", official, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        "salary_workbook": ("salary.xlsx", salary, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    }
    scope = {"academic_period": "II/2026", "effective_date": "2026-08-21"}
    context = client.post(
        "/api/admin/academic-management/designation-bootstrap/resolution-context",
        files=files,
        data=scope,
    ).json()
    selections = _resolution_selections(context)

    artifact_response = client.post(
        "/api/admin/academic-management/designation-bootstrap/alias-artifact",
        files=files,
        data={
            **scope,
            "resolution_token": context["resolution_token"],
            "selections": json.dumps(selections),
        },
    )
    assert artifact_response.status_code == 200
    assert artifact_response.headers["cache-control"] == "no-store"
    assert artifact_response.headers["content-type"].startswith("application/json")

    preview_response = client.post(
        "/api/admin/academic-management/designation-bootstrap/preview",
        files={
            **files,
            "alias_resolution": ("designation-alias-v2.json", artifact_response.content, "application/json"),
        },
        data={**scope, "program_identity": "MED"},
    )
    assert preview_response.status_code == 200
    preview_codes = {item["code"] for item in preview_response.json()["blockers"]}
    assert "unresolved_teacher_alias" not in preview_codes
    assert "unresolved_subject_alias" not in preview_codes

    stale_response = client.post(
        "/api/admin/academic-management/designation-bootstrap/alias-artifact",
        files={
            "official_workbook": files["official_workbook"],
            "salary_workbook": ("salary.xlsx", _salary(name="Changed Source"), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        },
        data={
            **scope,
            "resolution_token": context["resolution_token"],
            "selections": json.dumps(selections),
        },
    )
    assert stale_response.status_code == 409
    assert stale_response.json()["detail"] == {"code": "stale_resolution_sources"}


def test_exact_join_aggregates_salary_rows_and_response_is_pii_free(db_session):
    _seed(db_session)
    result = _preview(db_session)

    assert result["practice"]["official_row_count"] == 1
    assert result["practice"]["salary_row_count"] == 2
    assert result["practice"]["salary_parsed_row_count"] == 2
    assert result["practice"]["join_covered_count"] == 1
    assert result["practice"]["join_missing_count"] == 0
    assert result["practice"]["salary_hours_noncomparable_count"] == 1
    assert result["planned"] == {
        "theory_block_count": 1, "practice_block_count": 1,
        "theory_assignment_count": 1, "practice_assignment_count": 1,
    }
    assert _codes(result)["missing_classroom"] == 1
    serialized = json.dumps(result, default=str)
    for private_value in ("Private Theory Name", "Private Practice Name", "990001", "70000000", "private@example.test"):
        assert private_value not in serialized


def test_empty_official_import_has_explicit_plan_blockers(db_session):
    _seed(db_session)

    result = _preview(
        db_session,
        official=_official(include_theory=False, include_practice=False),
        salary=_salary(hours=()),
    )

    codes = _codes(result)
    assert codes["no_parsed_official_rows"] == 1
    assert codes["no_planned_blocks"] == 1
    assert codes["no_planned_assignments"] == 1
    assert result["can_apply"] is False


def test_theory_only_import_does_not_require_salary_corroboration(db_session):
    _seed(db_session)

    result = _preview(
        db_session,
        official=_official(include_practice=False),
        salary=_salary(),
    )

    codes = _codes(result)
    assert "orphan_salary_assignment" not in codes
    assert "missing_salary_subject" not in codes
    assert result["practice"]["join_missing_count"] == 0
    assert result["can_apply"] is True


def test_digest_is_deterministic_and_binds_source_and_database_state(db_session):
    _seed(db_session)
    first = _preview(db_session)
    repeated = _preview(db_session)
    assert repeated["preview_digest"] == first["preview_digest"]
    assert repeated["database_state_fingerprint"] == first["database_state_fingerprint"]

    changed_source = _preview(db_session, salary=_salary(hours=(5, 8)))
    assert changed_source["preview_digest"] != first["preview_digest"]
    teacher = db_session.get(Teacher, "880001")
    teacher.full_name = "Changed Private Name"
    db_session.flush()
    changed_state = _preview(db_session)
    assert changed_state["database_state_fingerprint"] != first["database_state_fingerprint"]
    assert changed_state["preview_digest"] != first["preview_digest"]


def test_ambiguity_hours_and_invalid_schedules_are_reason_coded(db_session):
    _seed(db_session)
    db_session.add(Teacher(ci="880002", full_name="Private Theory Name"))
    db_session.flush()
    result = _preview(
        db_session,
        official=_official(theory_schedule="", practice_schedule="MARTES: 12:00 PM-10:00 AM"),
        salary=_salary(hours=(5, 6), ci="990001"),
    )
    codes = _codes(result)
    assert codes["blank_schedule"] == 1
    assert codes["impossible_interval"] == 1
    assert result["can_apply"] is False


def test_salary_payment_hours_are_informational_not_comparable(db_session):
    result = _preview(db_session, salary=_salary(hours=(5, 6)))
    assert result["practice"]["salary_payment_occurrence_count"] == 2
    assert result["practice"]["salary_payment_hours_total"] == 11
    assert result["practice"]["salary_hours_noncomparable_count"] == 1
    assert "practice_hours_mismatch" not in _codes(result)


def test_conflicting_salary_identities_make_practice_join_ambiguous(db_session):
    _seed(db_session)
    result = _preview(db_session, salary=_salary(ci=("990001", "990002")))
    assert result["practice"]["join_ambiguous_count"] == 1
    assert _codes(result)["practice_join_ambiguous"] == 1


def test_duplicate_and_missing_resolution_and_availability_blockers(db_session):
    _seed(db_session, availability=False, classroom=False, practice_teacher=False, groups=False)
    result = _preview(db_session, official=_official(duplicate=True))
    codes = _codes(result)
    assert codes["duplicate_block"] >= 1
    assert codes["schedule_overlap"] >= 1
    assert codes["missing_classroom"] == 3
    assert codes["missing_group"] == 3
    assert codes["missing_teacher"] >= 1
    assert codes["missing_availability"] >= 1


def test_missing_subject_also_blocks_offering(db_session):
    _seed(db_session, subject=False)
    result = _preview(db_session)
    codes = _codes(result)
    assert codes["missing_subject"] == 2
    assert codes["missing_offering"] == 2
    assert codes["missing_salary_subject"] == 2


def test_source_join_is_database_independent_and_requires_nonblank_ci(db_session):
    result = _preview(db_session)
    assert result["practice"]["join_covered_count"] == 1
    assert result["practice"]["join_missing_count"] == 0
    assert _codes(result)["missing_program"] == 1

    blank_ci = _preview(db_session, salary=_salary(ci=""))
    assert blank_ci["practice"]["join_covered_count"] == 0
    assert blank_ci["practice"]["join_missing_ci_count"] == 1
    assert _codes(blank_ci)["practice_join_missing_ci"] == 1


def test_orphan_salary_assignment_and_duplicate_headers_block(db_session):
    orphan = _preview(db_session, salary=_salary(extra_assignment=True))
    assert orphan["practice"]["orphan_salary_count"] == 1
    assert _codes(orphan)["orphan_salary_assignment"] == 1

    duplicate_salary = _preview(db_session, salary=_salary(duplicate_header=True))
    assert _codes(duplicate_salary)["duplicate_salary_header"] == 1
    duplicate_official = _preview(db_session, official=_official(duplicate_header=True))
    assert _codes(duplicate_official)["theory_header_structure"] == 1


@pytest.mark.parametrize("raw", [
    "junk LUNES: 08:00 AM-09:30 AM",
    "LUNES: 08:00 AM-09:30 AM junk",
])
def test_schedule_rejects_junk_and_invalid_meridiem(raw):
    slots, errors, _warnings = _schedule(raw)
    assert not slots
    assert errors


def test_schedule_applies_only_forensic_normalizations_with_lineage():
    slots, errors, warnings = _schedule("JUVES   15: 10 PM-- 17: 25 PM")
    assert errors == []
    assert len(slots) == 1
    assert slots[0].weekday == "thursday"
    assert slots[0].start == time(15, 10)
    assert slots[0].end == time(17, 25)
    assert slots[0].raw == "JUVES   15: 10 PM-- 17: 25 PM"
    assert slots[0].normalized == "JUEVES: 15:10-17:25"
    assert set(warnings) == {
        "day_colon_inserted", "known_day_typo_corrected",
        "double_hyphen_normalized", "repeated_whitespace_normalized",
        "redundant_meridiem_on_24h_range",
    }

    twelve_hour, errors, warnings = _schedule("LUNES: 01:00 PM-02:30 PM")
    assert errors == [] and warnings == []
    assert (twelve_hour[0].start, twelve_hour[0].end) == (time(13), time(14, 30))


@pytest.mark.parametrize(("raw", "expected"), [
    ("LUNES: 12:00 AM-01:00 AM", (time(0), time(1))),
    ("LUNES: 12:00 PM-01:00 PM", (time(12), time(13))),
])
def test_schedule_preserves_true_twelve_hour_boundaries(raw, expected):
    slots, errors, warnings = _schedule(raw)
    assert errors == []
    assert warnings == []
    assert (slots[0].start, slots[0].end) == expected


@pytest.mark.parametrize(("raw", "expected"), [
    ("LUNES: 12:00 PM-13:30 PM", (time(12), time(13, 30))),
    ("LUNES: 11:00 AM-13:00 PM", (time(11), time(13))),
    ("LUNES: 13:00 AM-14:00 PM", (time(13), time(14))),
    ("LUNES: 13:00 PM-14:00 PM", (time(13), time(14))),
])
def test_schedule_strips_only_numerically_redundant_meridiem(raw, expected):
    slots, errors, warnings = _schedule(raw)
    assert errors == []
    assert warnings == ["redundant_meridiem_on_24h_range"]
    assert (slots[0].start, slots[0].end) == expected


@pytest.mark.parametrize("raw", [
    "LUNES: 12:00 AM-13:30 AM",
    "LUNES: 12:00 AM-13:30 PM",
    "LUNES: 11:00 PM-13:00 AM",
    "LUNES: 11:00 PM-13:00 PM",
])
def test_schedule_blocks_mixed_ranges_that_could_reinterpret_meridiem(raw):
    slots, errors, warnings = _schedule(raw)
    assert slots == ()
    assert errors == ["ambiguous_meridiem"]
    assert warnings == []


@pytest.mark.parametrize("raw, code", [
    ("LUNES: 10:00 AM-09:00 AM", "impossible_interval"),
    ("LUNES: 10:00 AM-10:00 AM", "impossible_interval"),
    ("LUNES: 01:00 AM-10:00 AM", "impossible_interval"),
    ("LUNES: XX:00 AM-10:00 AM", "invalid_schedule"),
    ("LUNES: 08:00 AM-10 AM", "invalid_schedule"),
    ("LUNES: 08:00 AM-10:00", "ambiguous_meridiem"),
    ("LUNES: 11:00 PM-01:00 AM", "impossible_interval"),
])
def test_schedule_keeps_unsafe_intervals_blocked(raw, code):
    slots, errors, _warnings = _schedule(raw)
    assert not slots
    assert code in errors


def test_strict_aliases_resolve_source_join_and_bind_digest(db_session):
    official = _official(
        practice_name="Official Practice Name",
        practice_subject="Official Subject PRACTICA",
    )
    salary = _salary(name="Salary Practice Name", subject="Salary Subject", ci="990001")
    without_alias = _preview(db_session, official=official, salary=salary)
    assert without_alias["practice"]["unresolved_teacher_alias_count"] == 1
    assert without_alias["practice"]["unresolved_subject_alias_count"] == 1

    alias = _alias(
        official, salary,
        teachers={
            "OFFICIAL PRACTICE NAME": {
                "salary_teacher_key": "SALARY PRACTICE NAME", "salary_ci": "990001",
            },
        },
        subjects={
            "SALARY SUBJECT": {
                "official_subject_key": "OFFICIAL SUBJECT",
                "semester": 1, "activity_type": "practice",
            },
        },
    )
    resolved = _preview(db_session, official=official, salary=salary, alias=alias)
    assert resolved["practice"]["join_covered_count"] == 1
    assert resolved["practice"]["unresolved_teacher_alias_count"] == 0
    assert resolved["practice"]["unresolved_subject_alias_count"] == 0
    assert resolved["sources"]["alias_sha256"] == hashlib.sha256(alias).hexdigest()
    assert resolved["preview_digest"] != without_alias["preview_digest"]
    serialized = json.dumps(resolved, default=str)
    assert "990001" not in serialized
    assert "OFFICIAL PRACTICE NAME" not in serialized


def test_subject_alias_drives_downstream_salary_catalog_resolution(db_session):
    official = _official(practice_subject="Official Subject PRACTICA")
    salary = _salary(subject="Salary Subject")
    program = _seed(db_session, subject=False)
    subject = AcademicSubject(code="OFF-SUB", name="Official Subject", active=True)
    db_session.add(subject)
    db_session.flush()
    db_session.add(SubjectOffering(
        subject_id=subject.id, program_id=program.id, academic_period="II/2026",
        semester=1, theory_hours=0, practice_hours=60, active=True,
    ))
    db_session.flush()
    alias = _alias(official, salary, subjects={
        "SALARY SUBJECT": {
            "official_subject_key": "OFFICIAL SUBJECT",
            "semester": 1,
            "activity_type": "practice",
        },
    })

    result = _preview(db_session, official=official, salary=salary, alias=alias)

    assert "missing_salary_subject" not in _codes(result)
    assert "ambiguous_salary_subject" not in _codes(result)


@pytest.mark.parametrize("override, code", [
    ({"official_sha256": "0" * 64}, "alias_official_hash_mismatch"),
    ({"salary_sha256": "0" * 64}, "alias_salary_hash_mismatch"),
    ({"academic_period": "I/2026"}, "alias_period_mismatch"),
    ({"effective_date": "2026-08-20"}, "alias_effective_date_mismatch"),
])
def test_alias_binding_drift_blocks(db_session, override, code):
    official, salary = _official(), _salary()
    result = _preview(db_session, official=official, salary=salary, alias=_alias(official, salary, **override))
    assert _codes(result)[code] == 1


def test_alias_orphan_cross_semester_and_many_to_one_block(db_session):
    official = _official(practice_name="Official One", practice_subject="Official Subject PRACTICA")
    workbook = load_workbook(BytesIO(official))
    workbook["DOCENTES ASISTENCIALES"].append([
        2, "Official Two", "Official Subject PRACTICA", "PRIMERO", "T-02",
        "MARTES: 12:00 PM-01:30 PM", 60, 12,
    ])
    official = _bytes(workbook)
    salary = _salary(name="Salary One", subject="Salary Subject", ci="990001")

    many_teacher = _alias(official, salary, teachers={
        "OFFICIAL ONE": {"salary_teacher_key": "SALARY ONE", "salary_ci": "990001"},
        "OFFICIAL TWO": {"salary_teacher_key": "SALARY ONE", "salary_ci": "990001"},
    })
    assert _codes(_preview(db_session, official=official, salary=salary, alias=many_teacher))["teacher_alias_many_to_one"] == 1

    orphan = _alias(official, salary, teachers={
        "MISSING OFFICIAL": {"salary_teacher_key": "SALARY ONE", "salary_ci": "990001"},
    })
    assert _codes(_preview(db_session, official=official, salary=salary, alias=orphan))["teacher_alias_orphan"] == 1

    cross = _alias(official, salary, subjects={
        "SALARY SUBJECT": {
            "official_subject_key": "OFFICIAL SUBJECT", "semester": 2,
            "activity_type": "practice",
        },
    })
    assert _codes(_preview(db_session, official=official, salary=salary, alias=cross))["subject_alias_cross_semester"] == 1

    salary_book = load_workbook(BytesIO(salary))
    salary_book.active.append([None, "Salary One", None, None, "990001", "Second Salary Subject", "SERVICE", "1", 1, 50])
    salary_two_subjects = _bytes(salary_book)
    many_subject = _alias(official, salary_two_subjects, subjects={
        "SALARY SUBJECT": {"official_subject_key": "OFFICIAL SUBJECT", "semester": 1, "activity_type": "practice"},
        "SECOND SALARY SUBJECT": {"official_subject_key": "OFFICIAL SUBJECT", "semester": 1, "activity_type": "practice"},
    })
    result = _preview(db_session, official=official, salary=salary_two_subjects, alias=many_subject)
    assert _codes(result)["subject_alias_many_to_one"] == 1


def test_alias_targets_cannot_collide_with_direct_source_mappings(db_session):
    official = _official(
        practice_name="Official Alias",
        practice_subject="Official Alias Subject PRACTICA",
    )
    official_book = load_workbook(BytesIO(official))
    official_book["DOCENTES ASISTENCIALES"].append([
        2, "Direct Teacher", "Direct Subject PRACTICA", "PRIMERO", "T-02",
        "MARTES: 12:00 PM-01:30 PM", 60, 12,
    ])
    official = _bytes(official_book)
    salary = _salary(name="Direct Teacher", subject="Direct Subject", ci="990001")
    salary_book = load_workbook(BytesIO(salary))
    salary_book.active.append([
        None, "Alias Salary Teacher", None, None, "990002", "Alias Salary Subject",
        "SERVICE", "1", 1, 50,
    ])
    salary = _bytes(salary_book)
    alias = _alias(
        official,
        salary,
        teachers={
            "OFFICIAL ALIAS": {
                "salary_teacher_key": "DIRECT TEACHER",
                "salary_ci": "990001",
            },
        },
        subjects={
            "ALIAS SALARY SUBJECT": {
                "official_subject_key": "DIRECT SUBJECT",
                "semester": 1,
                "activity_type": "practice",
            },
        },
    )

    codes = _codes(_preview(db_session, official=official, salary=salary, alias=alias))

    assert codes["teacher_alias_many_to_one"] == 1
    assert codes["subject_alias_many_to_one"] == 1


def test_alias_json_is_closed_and_versioned(db_session):
    official, salary = _official(), _salary()
    malformed = b"{not-json"
    assert _codes(_preview(db_session, official=official, salary=salary, alias=malformed))["alias_invalid_json"] == 1

    payload = json.loads(_alias(official, salary))
    payload["unexpected"] = True
    closed = json.dumps(payload).encode()
    assert _codes(_preview(db_session, official=official, salary=salary, alias=closed))["alias_schema_invalid"] == 1

    wrong_version = _alias(official, salary, schema_version="other")
    assert _codes(_preview(db_session, official=official, salary=salary, alias=wrong_version))["alias_schema_version"] == 1

    duplicate_key = _alias(official, salary).replace(
        b'"teacher_aliases": {}', b'"teacher_aliases": {}, "teacher_aliases": {}',
    )
    assert _codes(_preview(db_session, official=official, salary=salary, alias=duplicate_key))["alias_duplicate_key"] == 1


def test_real_file_schedule_recovery_aggregate(db_session):
    source_root = Path(__file__).resolve().parents[4]
    official_path = source_root / "DESIGNACION OFICIAL MATERIAS Y DOCENTES II-2026.xlsx"
    salary_path = source_root / "Planilla_Salario_Practicas_Septiembre_2026.xlsx"
    if not official_path.exists() or not salary_path.exists():
        pytest.skip("authoritative audit workbooks are not available")
    result = build_designation_bootstrap_preview(
        db_session,
        official_content=official_path.read_bytes(), salary_content=salary_path.read_bytes(),
        academic_period="II/2026", effective_date=date(2026, 8, 21),
        program_identity="MED 510-01",
    )
    assert result["theory"]["parsed_row_count"] == 417
    assert result["theory"]["error_row_count"] == 28
    assert result["practice"]["official_parsed_row_count"] == 36
    assert result["practice"]["official_error_row_count"] == 1
    assert result["practice"]["unresolved_teacher_alias_count"] == 28
    assert result["practice"]["unresolved_subject_alias_count"] == 37
    assert {item["code"]: item["count"] for item in result["warnings"]}[
        "redundant_meridiem_on_24h_range"
    ] == 454
    assert _codes(result)["ambiguous_meridiem"] == 1


def test_merged_and_skipped_rows_are_accounted_for(db_session):
    merged = _preview(db_session, official=_official(merged=True), salary=_salary(merged=True))
    codes = _codes(merged)
    assert codes["merged_required_official_cell"] == 1
    assert codes["merged_required_salary_cell"] == 1
    assert merged["theory"]["row_count"] == (
        merged["theory"]["parsed_row_count"] + merged["theory"]["skipped_row_count"]
        + merged["theory"]["error_row_count"]
    )

    workbook = load_workbook(BytesIO(_salary()))
    workbook.active.append([None, "TOTAL", None, None, None, None, None, None, None, 600])
    skipped = _preview(db_session, salary=_bytes(workbook))
    assert skipped["practice"]["salary_skipped_row_count"] >= 1


def test_structural_failures_and_normalized_collisions_block(db_session):
    malformed = Workbook()
    malformed.active.title = "WRONG"
    structural = _preview(db_session, official=_bytes(malformed))
    assert _codes(structural)["official_sheet_structure"] == 1

    program = _seed(db_session)
    collision_subject = AcademicSubject(code="SUB-B", name="subject a", active=True)
    db_session.add(collision_subject)
    db_session.flush()
    db_session.add(SubjectOffering(
        subject_id=collision_subject.id, program_id=program.id, academic_period="II/2026",
        semester=1, theory_hours=1, practice_hours=0, active=True,
    ))
    db_session.add_all([
        Teacher(ci="88 0001", full_name="PRIVATE THEORY NAME"),
        AcademicGroup(program_id=program.id, academic_period="II/2026", semester=1, shift="morning", code="M-01", active=True),
        Classroom(code="OTHER-ROOM", name="101", campus="Test", capacity=20, classroom_type="classroom", resources=[], active=True),
    ])
    db_session.flush()
    collisions = _preview(db_session)
    codes = _codes(collisions)
    assert codes["ambiguous_subject"] >= 1
    assert codes["ambiguous_offering"] >= 1
    assert codes["ambiguous_group"] >= 1
    assert codes["ambiguous_classroom"] >= 1
    assert codes["ambiguous_teacher"] >= 1


def test_program_collision_stops_cross_program_resolution(db_session):
    _seed(db_session)
    db_session.add(AcademicProgram(code="OTHER", name="MED", active=True))
    db_session.flush()
    result = _preview(db_session)
    assert _codes(result)["ambiguous_program"] == 1
    assert result["resolution"]["subject_resolved_count"] == 0


def test_service_does_not_autoflush_pending_state(db_session):
    pending = Teacher(ci="PENDING", full_name="Pending Private Teacher")
    db_session.add(pending)
    result = _preview(db_session)
    assert pending in db_session.new
    assert _codes(result)["missing_program"] == 1


def test_admin_endpoint_is_read_only_and_aggregate_only(client, db_session):
    _seed(db_session)
    before = db_session.new.copy(), db_session.dirty.copy(), db_session.deleted.copy()
    response = client.post(
        "/api/admin/academic-management/designation-bootstrap/preview",
        data={"academic_period": "II/2026", "effective_date": "2026-08-21", "program_identity": "MED"},
        files={
            "official_workbook": ("official.xlsx", _official(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
            "salary_workbook": ("salary.xlsx", _salary(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        },
    )
    assert response.status_code == 200
    assert response.json()["practice"]["salary_row_count"] == 2
    assert before == (db_session.new.copy(), db_session.dirty.copy(), db_session.deleted.copy())


def _endpoint(client, *, official=None, salary=None, alias=None, official_name="official.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", data=None):
    files = {
        "official_workbook": (official_name, _official() if official is None else official, mime),
        "salary_workbook": ("salary.xlsx", _salary() if salary is None else salary, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    }
    if alias is not None:
        files["alias_resolution"] = ("aliases.json", alias, "application/json")
    return client.post(
        "/api/admin/academic-management/designation-bootstrap/preview",
        data=data or {"academic_period": "II/2026", "effective_date": "2026-08-21", "program_identity": "MED"},
        files=files,
    )


def test_endpoint_rejects_unauthorized_invalid_and_unsafe_requests(client):
    authorization = client.headers.pop("Authorization")
    try:
        assert _endpoint(client).status_code in {401, 403}
    finally:
        client.headers["Authorization"] = authorization

    assert _endpoint(client, official=b"").status_code == 400
    assert _endpoint(client, official=b"not-a-workbook").status_code == 400
    assert _endpoint(client, official_name="official.bin").status_code == 415
    assert _endpoint(client, mime="application/octet-stream").status_code == 415
    assert _endpoint(client, official=b"x" * (20 * 1024 * 1024 + 1)).status_code == 413


@pytest.mark.parametrize("data", [
    {"academic_period": "", "effective_date": "2026-08-21", "program_identity": "MED"},
    {"academic_period": "invalid", "effective_date": "2026-08-21", "program_identity": "MED"},
    {"academic_period": "II/2026", "effective_date": "not-a-date", "program_identity": "MED"},
    {"academic_period": "II/2026", "effective_date": "2026-08-21", "program_identity": "   "},
])
def test_endpoint_rejects_malformed_form_fields(client, data):
    assert _endpoint(client, data=data).status_code == 422


def test_wrong_effective_date_is_a_stable_preview_blocker(client):
    response = _endpoint(client, data={
        "academic_period": "II/2026", "effective_date": "2026-08-20", "program_identity": "MED",
    })
    assert response.status_code == 200
    assert response.json()["can_apply"] is False
    assert any(item["code"] == "invalid_effective_date" for item in response.json()["blockers"])


def test_endpoint_accepts_bound_alias_json_part(client):
    official = _official(practice_name="Official Practice", practice_subject="Official Subject PRACTICA")
    salary = _salary(name="Salary Practice", subject="Salary Subject")
    alias = _alias(
        official, salary,
        teachers={"OFFICIAL PRACTICE": {"salary_teacher_key": "SALARY PRACTICE", "salary_ci": "990001"}},
        subjects={"SALARY SUBJECT": {"official_subject_key": "OFFICIAL SUBJECT", "semester": 1, "activity_type": "practice"}},
    )
    response = _endpoint(client, official=official, salary=salary, alias=alias)
    assert response.status_code == 200
    assert response.json()["practice"]["join_covered_count"] == 1


def test_workbook_dimension_limit_fails_closed(db_session):
    workbook = load_workbook(BytesIO(_official()))
    workbook["DOCENTES TEORICOS"].cell(5001, 1, "bounded")
    result = _preview(db_session, official=_bytes(workbook))
    assert _codes(result)["workbook_dimensions_exceeded"] == 1
