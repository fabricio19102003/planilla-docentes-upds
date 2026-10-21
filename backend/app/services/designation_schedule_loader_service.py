"""Load the current theory designations into a new schedule draft (designation-module-v2, T7b).

Designations do not store the classroom, so the room comes from the AULA column of the
designation workbook (sheet DOCENTES TEORICOS), resolved by number against the official
catalog. Missing subjects, offerings, groups, catalog rooms and teacher availability
(06:30–23:05 on each class day) are created in apply (L15). The draft is never published:
clashes are reported (T8a) for the user to fix in the planner. A row that cannot become a
block is a blocker and nothing is written.
"""
from __future__ import annotations

import hashlib
import io
import json
from dataclasses import dataclass, field
from datetime import date, time

import openpyxl
from sqlalchemy.orm import Session

from app.models.academic_management import (
    AcademicGroup,
    AcademicProgram,
    AcademicScheduleAssignment,
    AcademicScheduleBlock,
    AcademicScheduleDraft,
    AcademicSubject,
    Classroom,
    SubjectOffering,
    TeacherAvailability,
)
from app.models.designation import Designation
from app.models.teacher import Teacher
from app.services import academic_management_service as academic
from app.services.academic_schedule_conflicts_service import analyze
from app.services.academic_schedule_view_service import SubjectResolver, ViewBlock, _fold, normalized_group
from app.services.app_settings_service import get_academic_hour_minutes
from app.services.classroom_catalog_service import CAMPUS, CatalogRoom, catalog_room
from app.services.effective_schedule_service import canonical_weekday
from app.utils.academic_semester import semester_number

SHEET = "DOCENTES TEORICOS"
AVAILABILITY_START, AVAILABILITY_END = time(6, 30), time(23, 5)
ENGLISH_DAY = {
    "lunes": "monday", "martes": "tuesday", "miercoles": "wednesday",
    "jueves": "thursday", "viernes": "friday", "sabado": "saturday",
}
DAY_ORDER = {day: index for index, day in enumerate(ENGLISH_DAY)}
BLOCKER_MESSAGES = {
    "missing_classroom": "La fila de la designación no tiene un número de aula del catálogo (columna AULA).",
    "invalid_semester": "El semestre de la designación no es reconocible.",
    "invalid_hours": "Las horas académicas de cada clase deben ser un número entero positivo.",
    "invalid_time": "El horario debe estar entre 06:00 y 23:30, en pasos de 5 minutos y durar de 15 minutos a 6 horas.",
    "invalid_weekday": "El día del horario no es de lunes a sábado.",
    "duplicate_assignment": "Otro docente ya dicta este mismo bloque en las mismas fechas.",
    "unknown_teacher": "El docente de la designación no existe.",
    "draft_exists": "Ya hay un borrador activo con este nombre: archívelo o elija otra fecha.",
}


class LoaderInputError(ValueError):
    """The workbook or request cannot be read; nothing was planned."""


class LoaderBlocked(Exception):
    def __init__(self, code: str, preview: dict) -> None:
        super().__init__(code)
        self.code = code
        self.preview = preview


@dataclass
class PlannedBlock:
    designation: Designation
    semester: int
    group_code: str
    subject_name: str
    subject: AcademicSubject | None
    room: CatalogRoom
    weekday: str
    start: time
    end: time
    academic_hours: int | None
    effective_from: date
    effective_to: date | None
    assignments: list[tuple[str, date, date | None, int]] = field(default_factory=list)



def _slot_hours(raw: object) -> tuple[int | None, bool]:
    """Whole positive academic hours; empty or 0 means "not set" (legacy rule applies)."""
    if raw is None or raw == "" or raw == 0:
        return None, True
    try:
        value = float(str(raw).strip())
    except ValueError:
        return None, False
    if value <= 0 or not value.is_integer():
        return None, False
    return int(value), True

def _cell_text(value: object) -> str:
    return " ".join(str(value or "").split())


def read_rooms(content: bytes) -> tuple[dict[tuple, str], dict[tuple, set[str]]]:
    """AULA by (CI, subject, semester, group) and the rooms seen for (subject, semester, group)."""
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many types for corrupt files
        raise LoaderInputError("El archivo no es un Excel válido.") from exc
    if SHEET not in workbook.sheetnames:
        raise LoaderInputError(f"El Excel debe tener la hoja {SHEET}.")
    rows = workbook[SHEET].iter_rows(values_only=True)
    header = [_fold(cell) for cell in next(rows, ())]
    try:
        subject_col, semester_col = header.index("materias"), header.index("semestre")
        room_col, ci_col = header.index("aula"), header.index("ci")
    except ValueError as exc:
        raise LoaderInputError("La hoja debe tener las columnas MATERIAS, SEMESTRE, AULA y CI.") from exc
    group_col = semester_col + 1
    exact: dict[tuple, str] = {}
    by_group: dict[tuple, set[str]] = {}
    for row in rows:
        if not row or len(row) <= max(subject_col, room_col, ci_col, group_col) or not row[subject_col]:
            continue
        room = catalog_room(row[room_col])
        key = (_fold(row[subject_col]), semester_number(_cell_text(row[semester_col])),
               normalized_group(row[group_col]))
        if room is None:
            continue
        exact[(_cell_text(row[ci_col]), *key)] = room.room
        by_group.setdefault(key, set()).add(room.room)
    return exact, by_group


def _program(db: Session, program_id: int | None) -> AcademicProgram:
    query = db.query(AcademicProgram).filter(AcademicProgram.active.is_(True))
    if program_id is not None:
        program = query.filter(AcademicProgram.id == program_id).first()
        if program is None:
            raise LoaderInputError("El programa académico no existe o está inactivo.")
        return program
    programs = query.order_by(AcademicProgram.id).all()
    if len(programs) != 1:
        raise LoaderInputError("Seleccione el programa académico.")
    return programs[0]


def _valid_time(start: time, end: time) -> bool:
    start_minutes, end_minutes = start.hour * 60 + start.minute, end.hour * 60 + end.minute
    return (academic.BLOCK_DAY_START <= start_minutes and end_minutes <= academic.BLOCK_DAY_END
            and not start_minutes % academic.BLOCK_STEP_MINUTES and not end_minutes % academic.BLOCK_STEP_MINUTES
            and academic.BLOCK_MIN_MINUTES <= end_minutes - start_minutes <= academic.BLOCK_MAX_MINUTES)


def _overlaps(left_from: date, left_to: date | None, right_from: date, right_to: date | None) -> bool:
    return (left_to is None or right_from <= left_to) and (right_to is None or left_from <= right_to)


def _plan(db: Session, *, academic_period: str, effective_from: date, workbook: bytes,
          program_id: int | None) -> dict:
    program = _program(db, program_id)
    exact_rooms, group_rooms = read_rooms(workbook)
    resolver = SubjectResolver(db, academic_period)
    teachers = {ci for (ci,) in db.query(Teacher.ci)}
    designations = db.query(Designation).filter(
        Designation.academic_period == academic_period,
        Designation.designation_type != "practice",
    ).order_by(Designation.id).all()
    designations = [item for item in designations
                    if item.contract_end_date is None or item.contract_end_date >= effective_from]

    blockers: list[dict] = []

    def block(code: str, item: Designation) -> None:
        blockers.append({
            "code": code, "designation_id": item.id, "teacher_ci": item.teacher_ci, "subject": item.subject,
            "semester": item.semester, "group_code": item.group_code, "message": BLOCKER_MESSAGES[code],
        })

    planned: dict[tuple, PlannedBlock] = {}
    for item in designations:
        semester = semester_number(item.semester)
        group_code = normalized_group(item.group_code)
        if semester is None:
            block("invalid_semester", item)
            continue
        if item.teacher_ci not in teachers:
            block("unknown_teacher", item)
            continue
        key = (_fold(item.subject), semester, group_code)
        room_number = exact_rooms.get((item.teacher_ci, *key))
        if room_number is None and len(group_rooms.get(key, ())) == 1:
            room_number = next(iter(group_rooms[key]))
        room = catalog_room(room_number)
        if room is None:
            block("missing_classroom", item)
            continue
        subject = resolver.resolve(item.subject, semester)
        start_from = max(effective_from, item.contract_start_date or effective_from)
        rows_blocked = False
        candidates = []
        for slot in item.schedule_json or []:
            weekday = canonical_weekday(slot.get("dia", ""))
            try:
                start, end = time.fromisoformat(str(slot["hora_inicio"])), time.fromisoformat(str(slot["hora_fin"]))
            except (KeyError, TypeError, ValueError):
                start = end = None
            if weekday not in ENGLISH_DAY:
                block("invalid_weekday", item)
                rows_blocked = True
                break
            if start is None or not _valid_time(start, end):
                block("invalid_time", item)
                rows_blocked = True
                break
            hours, hours_ok = _slot_hours(slot.get("horas_academicas"))
            if not hours_ok:
                block("invalid_hours", item)
                rows_blocked = True
                break
            candidates.append((weekday, start, end, hours))
        if rows_blocked:
            continue
        for weekday, start, end, hours in candidates:
            block_key = (subject.id if subject else key[0], semester, group_code, room.code, weekday, start, end)
            target = planned.get(block_key)
            if target is None:
                target = planned[block_key] = PlannedBlock(
                    designation=item, semester=semester, group_code=group_code, subject_name=item.subject,
                    subject=subject, room=room, weekday=weekday, start=start, end=end, academic_hours=hours,
                    effective_from=start_from, effective_to=item.contract_end_date,
                )
            elif any(_overlaps(a_from, a_to, start_from, item.contract_end_date)
                     for _ci, a_from, a_to, _id in target.assignments):
                block("duplicate_assignment", item)
                break
            target.assignments.append((item.teacher_ci, start_from, item.contract_end_date, item.id))

    draft_name = f"Designación vigente {academic_period} desde {effective_from.isoformat()}"
    if db.query(AcademicScheduleDraft.id).filter(
        AcademicScheduleDraft.program_id == program.id,
        AcademicScheduleDraft.academic_period == academic_period,
        AcademicScheduleDraft.normalized_name == academic._normalized_name(draft_name),
        AcademicScheduleDraft.status == "draft",
    ).first():
        blockers.insert(0, {"code": "draft_exists", "designation_id": None, "teacher_ci": None, "subject": None,
                            "semester": None, "group_code": None, "message": BLOCKER_MESSAGES["draft_exists"]})

    blocks = sorted(
        planned.values(),
        key=lambda item: (DAY_ORDER[item.weekday], item.start, item.semester, item.group_code, item.subject_name),
    )

    offerings = {(offering.subject_id, offering.semester) for offering in db.query(SubjectOffering).filter(
        SubjectOffering.program_id == program.id, SubjectOffering.academic_period == academic_period)}
    curriculum: dict[int, set[int]] = {}
    for subject_id, semester in offerings:
        curriculum.setdefault(subject_id, set()).add(semester)
    warnings: list[dict] = []
    for item in planned.values():
        semesters = curriculum.get(item.subject.id, set()) if item.subject else set()
        if not semesters or item.semester in semesters:
            continue
        for *_rest, designation_id in item.assignments:
            if any(warning["designation_id"] == designation_id for warning in warnings):
                continue
            source = next(row for row in designations if row.id == designation_id)
            listed = ", ".join(str(value) for value in sorted(semesters))
            warnings.append({
                "code": "subject_other_semester", "designation_id": designation_id, "teacher_ci": source.teacher_ci,
                "subject": source.subject, "semester": source.semester, "group_code": source.group_code,
                "message": (f"La materia {item.subject.code} está en la malla del semestre {listed}, pero la designación "
                            f"dice semestre {item.semester}: se creará una oferta en el semestre {item.semester}."),
            })
    warnings.sort(key=lambda warning: warning["designation_id"])
    groups = {(group.semester, normalized_group(group.code)) for group in db.query(AcademicGroup).filter(
        AcademicGroup.program_id == program.id, AcademicGroup.academic_period == academic_period)}
    rooms = {code for (code,) in db.query(Classroom.code)}
    new_subjects = sorted({item.subject_name for item in blocks if item.subject is None})
    new_offerings = {(item.subject.id if item.subject else _fold(item.subject_name), item.semester)
                     for item in blocks if item.subject is None or (item.subject.id, item.semester) not in offerings}
    new_groups = sorted({(item.semester, item.group_code) for item in blocks} - groups)
    new_rooms = sorted({item.room.code for item in blocks} - rooms)
    windows = {(ci, item.weekday) for item in blocks for ci, *_rest in item.assignments}
    new_windows = sorted(window for window in windows if not _covered(db, academic_period, window, blocks))

    view = [ViewBlock(
        key=f"load:{index}", semester=item.semester, group_code=item.group_code,
        subject_code=item.subject.code if item.subject else None, subject_name=item.subject_name,
        activity_type="theory", weekday=item.weekday, start_time=item.start.strftime("%H:%M"),
        end_time=item.end.strftime("%H:%M"),
        teacher_ci=item.assignments[0][0],
        teacher_name=None, classroom_code=item.room.code, classroom_name=item.room.name, source_type="draft",
    ) for index, item in enumerate(blocks)]
    conflicts = analyze(view, rooms=[], load_threshold=40, academic_hour_minutes=get_academic_hour_minutes(db))
    conflicts = {"summary": conflicts["summary"], "clashes": conflicts["clashes"]}

    block_rows = [{
        "designation_id": item.designation.id, "semester": item.semester, "group_code": item.group_code,
        "subject": item.subject_name, "subject_code": item.subject.code if item.subject else None,
        "weekday": item.weekday, "start_time": item.start.strftime("%H:%M"), "end_time": item.end.strftime("%H:%M"),
        "academic_hours": item.academic_hours, "classroom_code": item.room.code,
        "teacher_ci": item.assignments[0][0], "effective_from": item.assignments[0][1].isoformat(),
        "assignments": [{"teacher_ci": ci, "effective_from": a_from.isoformat(),
                         "effective_to": a_to.isoformat() if a_to else None, "designation_id": did}
                        for ci, a_from, a_to, did in item.assignments],
    } for item in blocks]
    digest_source = {
        "period": academic_period, "effective_from": effective_from.isoformat(), "program": program.id,
        "workbook": hashlib.sha256(workbook).hexdigest(), "blocks": block_rows, "blockers": blockers,
        "create": [new_subjects, sorted(map(str, new_offerings)), new_groups, new_rooms, new_windows],
    }
    preview = {
        "academic_period": academic_period, "effective_from": effective_from, "program_id": program.id,
        "draft_name": draft_name,
        "digest": hashlib.sha256(json.dumps(digest_source, sort_keys=True, default=str).encode()).hexdigest(),
        "can_apply": not blockers and bool(blocks),
        "counts": {"designations": len({did for item in blocks for *_r, did in item.assignments}),
                   "blocks": len(blocks), "assignments": sum(len(item.assignments) for item in blocks)},
        "create": {"subjects": new_subjects, "offerings": len(new_offerings),
                   "groups": [f"{semester} {code}" for semester, code in new_groups],
                   "classrooms": new_rooms, "availability": len(new_windows)},
        "blockers": blockers, "warnings": warnings, "blocks": block_rows, "conflicts": conflicts,
    }
    return {"preview": preview, "program": program, "blocks": blocks, "new_windows": new_windows}


def _covered(db: Session, academic_period: str, window: tuple[str, str], blocks: list[PlannedBlock]) -> bool:
    ci, weekday = window
    spans = [(item.start, item.end) for item in blocks
             if item.weekday == weekday and any(a_ci == ci for a_ci, *_r in item.assignments)]
    rows = db.query(TeacherAvailability).filter(
        TeacherAvailability.teacher_ci == ci, TeacherAvailability.academic_period == academic_period,
        TeacherAvailability.weekday == ENGLISH_DAY[weekday], TeacherAvailability.active.is_(True),
    ).all()
    return all(any(row.start_time <= start and end <= row.end_time for row in rows) for start, end in spans)


def preview_load(db: Session, *, academic_period: str, effective_from: date, workbook: bytes,
                 program_id: int | None = None) -> dict:
    return _plan(db, academic_period=academic_period, effective_from=effective_from,
                 workbook=workbook, program_id=program_id)["preview"]


def apply_load(db: Session, *, academic_period: str, effective_from: date, workbook: bytes, digest: str,
               program_id: int | None = None) -> dict:
    plan = _plan(db, academic_period=academic_period, effective_from=effective_from,
                 workbook=workbook, program_id=program_id)
    preview, program = plan["preview"], plan["program"]
    if preview["digest"] != digest:
        raise LoaderBlocked("stale_preview", preview)
    if not preview["can_apply"]:
        raise LoaderBlocked("blocked", preview)

    academic._lock_row(db, AcademicProgram, program.id)
    draft = AcademicScheduleDraft(
        program_id=program.id, academic_period=academic_period, name=preview["draft_name"],
        normalized_name=academic._normalized_name(preview["draft_name"]), status="draft",
    )
    db.add(draft)
    rooms = {room.code: room for room in db.query(Classroom)}
    groups = {(group.semester, normalized_group(group.code)): group for group in db.query(AcademicGroup).filter(
        AcademicGroup.program_id == program.id, AcademicGroup.academic_period == academic_period)}
    offerings = {(offering.subject_id, offering.semester): offering for offering in db.query(SubjectOffering).filter(
        SubjectOffering.program_id == program.id, SubjectOffering.academic_period == academic_period)}
    created_subjects: dict[str, AcademicSubject] = {}
    assignments = 0
    for item in plan["blocks"]:
        subject = item.subject or created_subjects.get(_fold(item.subject_name))
        if subject is None:
            subject = AcademicSubject(code=_subject_code(db, item.semester), name=item.subject_name,
                                      description="Creada al cargar la designación", active=True)
            db.add(subject)
            db.flush()
            created_subjects[_fold(item.subject_name)] = subject
        offering = offerings.get((subject.id, item.semester))
        if offering is None:
            designation = item.designation
            hours = designation.semester_hours or max(1, (designation.weekly_hours or 1) * 20)
            offering = offerings[(subject.id, item.semester)] = SubjectOffering(
                subject_id=subject.id, program_id=program.id, academic_period=academic_period,
                semester=item.semester, theory_hours=hours, practice_hours=0, active=True,
            )
            db.add(offering)
        group = groups.get((item.semester, item.group_code))
        if group is None:
            group = groups[(item.semester, item.group_code)] = AcademicGroup(
                program_id=program.id, academic_period=academic_period, semester=item.semester,
                shift=item.group_code.split("-")[0], code=item.group_code, expected_size=None, active=True,
            )
            db.add(group)
        room = rooms.get(item.room.code)
        if room is None:
            room = rooms[item.room.code] = Classroom(
                code=item.room.code, name=item.room.name, campus=CAMPUS, capacity=item.room.capacity,
                classroom_type="classroom", resources=[], active=True,
            )
            db.add(room)
        db.flush()
        block = AcademicScheduleBlock(
            draft=draft, offering_id=offering.id, group_id=group.id, classroom_id=room.id,
            activity_type="theory", weekday=ENGLISH_DAY[item.weekday], start_time=item.start,
            end_time=item.end, academic_hours=item.academic_hours,
        )
        db.add(block)
        for ci, a_from, a_to, _designation_id in item.assignments:
            db.add(AcademicScheduleAssignment(block=block, teacher_ci=ci, effective_from=a_from, effective_to=a_to))
            assignments += 1
    for ci, weekday in plan["new_windows"]:
        db.add(TeacherAvailability(teacher_ci=ci, academic_period=academic_period, weekday=ENGLISH_DAY[weekday],
                                   start_time=AVAILABILITY_START, end_time=AVAILABILITY_END, active=True))
    db.flush()
    return {
        "draft_id": draft.id, "draft_name": draft.name, "blocks": len(plan["blocks"]), "assignments": assignments,
        "created": preview["create"],
        "designation_ids": sorted({did for item in plan["blocks"] for *_r, did in item.assignments}),
        "conflicts": preview["conflicts"]["summary"],
    }


def _subject_code(db: Session, semester: int) -> str:
    prefix = f"DSG-{semester}-"
    used = {code for (code,) in db.query(AcademicSubject.code).filter(AcademicSubject.code.like(f"{prefix}%"))}
    number = 1
    while f"{prefix}{number:03d}" in used:
        number += 1
    return f"{prefix}{number:03d}"

