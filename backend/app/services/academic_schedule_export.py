"""Excel export of the schedule view: one weekly grid per (semester, group) plus a list."""
from __future__ import annotations

from datetime import date
from io import BytesIO
import hashlib
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

GRID_START = 6 * 60 + 30
GRID_END = 23 * 60 + 30
BAND = 30
FIRST_BAND_ROW = 4
DAYS = [("lunes", "Lunes"), ("martes", "Martes"), ("miercoles", "Miércoles"),
        ("jueves", "Jueves"), ("viernes", "Viernes"), ("sabado", "Sábado")]
ACTIVITY = {"theory": "Teoría", "practice": "Práctica"}
SOURCE = {"current": "Horario vigente", "draft": "Borrador"}
PALETTE = ["DBEAFE", "DCFCE7", "FEF3C7", "FCE7F3", "E0E7FF", "CCFBF1", "FFEDD5", "EDE9FE",
           "FEE2E2", "ECFCCB", "CFFAFE", "F5F5F4"]
THIN = Side(style="thin", color="CBD5E1")
DASHED = Side(style="dashed", color="64748B")


def _minutes(value: str) -> int:
    hour, minute = value.split(":")
    return int(hour) * 60 + int(minute)


def _clock(minutes: int) -> str:
    return f"{minutes // 60:02d}:{minutes % 60:02d}"


def _color(subject: str) -> str:
    return PALETTE[int(hashlib.sha256(subject.encode()).hexdigest(), 16) % len(PALETTE)]


def _rows(block: dict) -> tuple[int, int]:
    start = max(GRID_START, _minutes(block["start_time"]))
    end = min(GRID_END, _minutes(block["end_time"]))
    first = (start - GRID_START) // BAND
    last = max(first, (end - 1 - GRID_START) // BAND)
    return FIRST_BAND_ROW + first, FIRST_BAND_ROW + last


def _text(block: dict) -> str:
    detail = " · ".join(item for item in (ACTIVITY.get(block["activity_type"], block["activity_type"]),
                                          block["teacher_name"] or block["teacher_ci"]) if item)
    lines = [block["subject_name"], detail, f"{block['start_time']}–{block['end_time']}"]
    if block["classroom_name"] or block["classroom_code"]:
        lines.append(block["classroom_name"] or block["classroom_code"])
    return "\n".join(lines)


_INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


def _sheet_title(semester: int | None, group: str, used: set[str]) -> str:
    """Excel-safe, unique sheet title (no []:*?/\\, at most 31 characters)."""
    prefix = f"S{semester}" if semester is not None else "Sin semestre"
    base = _INVALID_SHEET_CHARS.sub("-", f"{prefix} {group}")[:31]
    title, counter = base, 2
    while title in used:
        suffix = f" ({counter})"
        title, counter = base[:31 - len(suffix)] + suffix, counter + 1
    used.add(title)
    return title


def _safe_segment(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9.-]+", "_", value.replace("/", "-")).strip("_")


def export_filename(academic_period: str, semester: int | None, group: str | None) -> str:
    parts = [f"horarios_{_safe_segment(academic_period)}"]
    if semester is not None:
        parts.append(f"sem{semester}")
    if group:
        parts.append(_safe_segment(group))
    return "_".join(part for part in parts if part) + ".xlsx"


def _grid_sheet(workbook: Workbook, view: dict, semester: int | None, group: str, blocks: list[dict],
                used: set[str]) -> None:
    sheet = workbook.create_sheet(_sheet_title(semester, group, used))
    reference = view["reference_date"]
    reference = reference.isoformat() if isinstance(reference, date) else str(reference)
    sheet["A1"] = (f"Horarios {view['academic_period']} — Semestre {semester if semester is not None else 'sin definir'}"
                   f" — Grupo {group} — {SOURCE.get(view['source'], view['source'])} al {reference}")
    sheet["A1"].font = Font(bold=True, size=13)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=7)
    header = ["Hora", *[label for _, label in DAYS]]
    for column, value in enumerate(header, start=1):
        cell = sheet.cell(row=3, column=column, value=value)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E293B")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    bands = (GRID_END - GRID_START) // BAND
    for index in range(bands):
        row = FIRST_BAND_ROW + index
        start = GRID_START + index * BAND
        label = sheet.cell(row=row, column=1, value=f"{_clock(start)}–{_clock(start + BAND)}")
        label.font = Font(size=9, color="475569")
        label.alignment = Alignment(horizontal="center", vertical="center")
        for column in range(1, 8):
            sheet.cell(row=row, column=column).border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
        sheet.row_dimensions[row].height = 22
    sheet.column_dimensions["A"].width = 13
    for column in range(2, 8):
        sheet.column_dimensions[get_column_letter(column)].width = 28

    for column, (weekday, _label) in enumerate(DAYS, start=2):
        day_blocks = sorted((b for b in blocks if b["weekday"] == weekday), key=lambda b: b["start_time"])
        clusters: list[list] = []
        for block in day_blocks:
            first, last = _rows(block)
            if clusters and first <= clusters[-1][1]:
                clusters[-1][1] = max(clusters[-1][1], last)
                clusters[-1][2].append(block)
            else:
                clusters.append([first, last, [block]])
        for first, last, members in clusters:
            cell = sheet.cell(row=first, column=column, value="\n\n".join(_text(b) for b in members))
            practice = all(b["activity_type"] == "practice" for b in members)
            cell.fill = PatternFill("solid", fgColor=_color(members[0]["subject_name"]))
            cell.font = Font(size=9, italic=practice, bold=not practice)
            cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
            side = DASHED if practice else THIN
            if last > first:
                sheet.merge_cells(start_row=first, start_column=column, end_row=last, end_column=column)
            for row in range(first, last + 1):
                sheet.cell(row=row, column=column).border = Border(left=side, right=side, top=side, bottom=side)
    sheet.freeze_panes = "B4"


def schedule_workbook(view: dict) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    groups: dict[tuple, list[dict]] = {}
    for block in view["blocks"]:
        groups.setdefault((block["semester"], block["group_code"]), []).append(block)
    used: set[str] = {"Lista"}
    for (semester, group), blocks in groups.items():
        _grid_sheet(workbook, view, semester, group, blocks, used)
    listing = workbook.create_sheet("Lista")
    listing.append(["Semestre", "Grupo", "Día", "Inicio", "Fin", "Materia", "Actividad", "Docente", "Aula"])
    for cell in listing[1]:
        cell.font = Font(bold=True)
    day_label = dict(DAYS) | {"domingo": "Domingo"}
    for block in view["blocks"]:
        listing.append([
            block["semester"], block["group_code"], day_label.get(block["weekday"], block["weekday"]),
            block["start_time"], block["end_time"], block["subject_name"],
            ACTIVITY.get(block["activity_type"], block["activity_type"]),
            block["teacher_name"] or block["teacher_ci"],
            block["classroom_name"] or block["classroom_code"] or "",
        ])
    for column, width in zip("ABCDEFGHI", (10, 9, 11, 8, 8, 42, 10, 36, 28)):
        listing.column_dimensions[column].width = width
    listing.freeze_panes = "A2"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
