"""Schedule view by semester and group, for the current (paid) schedule and for drafts."""
from __future__ import annotations

from datetime import date
from io import BytesIO

import openpyxl

from app.models.designation import Designation
from app.models.teacher import Teacher

BASE = "/api/admin/academic-management"
VIEW = f"{BASE}/schedule-view"


def _post(client, path, payload):
    response = client.post(f"{BASE}{path}", json=payload)
    assert response.status_code == 201, response.text
    return response.json()


def _catalog(client, db_session):
    program = _post(client, "/programs", {"code": "VIEW", "name": "View Program"})
    anatomy = _post(client, "/subjects", {"code": "MRF 0100", "name": "Anatomía Humana I"})
    offering = _post(client, "/offerings", {
        "subject_id": anatomy["id"], "program_id": program["id"], "academic_period": "I/2027",
        "semester": 1, "theory_hours": 51, "practice_hours": 51,
    })
    group = _post(client, "/groups", {
        "program_id": program["id"], "academic_period": "I/2027", "semester": 1,
        "shift": "M", "code": "M-01", "expected_size": 30,
    })
    room = _post(client, "/classrooms", {
        "code": "LAB-1", "name": "Laboratorio 1", "campus": "Cobija", "capacity": 30,
        "classroom_type": "laboratory", "resources": [],
    })
    db_session.add_all([
        Teacher(ci="V-THEORY", full_name="DOCENTE TEORIA"),
        Teacher(ci="V-PRACTICE", full_name="Docente Practica"),
    ])
    db_session.flush()
    for ci in ("V-THEORY", "V-PRACTICE"):
        availability = client.post(f"{BASE}/availability", json={
            "teacher_ci": ci, "academic_period": "I/2027", "weekday": "tuesday",
            "start_time": "06:30", "end_time": "23:00",
        })
        assert availability.status_code == 201, availability.text
    return program, anatomy, offering, group, room


def _designation(db_session, **overrides):
    values = dict(
        teacher_ci="V-THEORY", subject="ANATOMÍA HUMANA I", semester="PRIMERO", group_code="M-1",
        academic_period="I/2027", designation_type="regular", weekly_hours=3,
        schedule_json=[{"dia": "lunes", "hora_inicio": "06:30", "hora_fin": "08:55", "horas_academicas": 3}],
        contract_start_date=date(2027, 1, 1), contract_end_date=date(2027, 6, 30),
    )
    values.update(overrides)
    designation = Designation(**values)
    db_session.add(designation)
    db_session.commit()
    return designation


def _practice_draft(client, program, offering, group, room, *, name="Prácticas"):
    draft = _post(client, "/schedule-drafts", {
        "program_id": program["id"], "academic_period": "I/2027", "name": name,
    })
    block = _post(client, f"/schedule-drafts/{draft['id']}/blocks", {
        "offering_id": offering["id"], "group_id": group["id"], "classroom_id": room["id"],
        "activity_type": "practice", "weekday": "tuesday", "start_time": "10:00", "end_time": "12:00",
    })
    assigned = client.post(
        f"{BASE}/schedule-drafts/{draft['id']}/blocks/{block['id']}/assignments",
        json={"teacher_ci": "V-PRACTICE", "effective_from": "2027-01-01", "effective_to": None},
    )
    assert assigned.status_code == 201, assigned.text
    return draft, block


def _publish(client, draft_id):
    preview = client.post(f"{BASE}/schedule-drafts/{draft_id}/publication-preview",
                          json={"effective_from": "2027-02-01"})
    assert preview.status_code == 200, preview.text
    published = client.post(f"{BASE}/schedule-drafts/{draft_id}/publish", json={
        "effective_from": "2027-02-01", "preview_digest": preview.json()["digest"],
    })
    assert published.status_code == 201, published.text


def test_current_schedule_merges_paid_theory_and_published_practice(client, db_session):
    program, _subject, offering, group, room = _catalog(client, db_session)
    _designation(db_session)
    draft, _block = _practice_draft(client, program, offering, group, room)
    _publish(client, draft["id"])

    response = client.get(VIEW, params={"academic_period": "I/2027", "date": "2027-03-01"})
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["academic_period"] == "I/2027"
    assert body["source"] == "current"
    assert body["reference_date"] == "2027-03-01"
    assert body["draft_id"] is None
    assert body["semesters"] == [{"semester": 1, "groups": ["M-1"]}]
    assert [
        (b["semester"], b["group_code"], b["subject_code"], b["subject_name"], b["activity_type"],
         b["weekday"], b["start_time"], b["end_time"], b["teacher_ci"], b["teacher_name"],
         b["classroom_code"], b["classroom_name"], b["source_type"])
        for b in body["blocks"]
    ] == [
        (1, "M-1", "MRF 0100", "Anatomía Humana I", "theory", "lunes", "06:30", "08:55",
         "V-THEORY", "DOCENTE TEORIA", None, None, "legacy"),
        (1, "M-1", "MRF 0100", "Anatomía Humana I", "practice", "martes", "10:00", "12:00",
         "V-PRACTICE", "Docente Practica", "LAB-1", "Laboratorio 1", "published"),
    ]
    assert len({b["key"] for b in body["blocks"]}) == 2


def test_current_schedule_respects_contract_dates(client, db_session):
    _catalog(client, db_session)
    _designation(db_session, contract_end_date=date(2027, 2, 15))
    body = client.get(VIEW, params={"academic_period": "I/2027", "date": "2027-03-01"}).json()
    assert body["blocks"] == []
    assert body["semesters"] == []


def test_filters_normalize_group_and_semesters_ignore_filters(client, db_session):
    _catalog(client, db_session)
    _designation(db_session)
    _designation(db_session, semester="Segundo", group_code="T-02", subject="FISIOLOGÍA I",
                 schedule_json=[{"dia": "miercoles", "hora_inicio": "13:30", "hora_fin": "15:00",
                                 "horas_academicas": 2}])
    params = {"academic_period": "I/2027", "date": "2027-03-01", "semester": 1, "group": "M-01"}
    body = client.get(VIEW, params=params).json()
    assert body["semesters"] == [{"semester": 1, "groups": ["M-1"]}, {"semester": 2, "groups": ["T-2"]}]
    assert [(b["semester"], b["group_code"]) for b in body["blocks"]] == [(1, "M-1")]
    other = client.get(VIEW, params={**params, "semester": 2, "group": "T-2"}).json()
    assert [(b["subject_name"], b["subject_code"]) for b in other["blocks"]] == [("FISIOLOGÍA I", None)]


def test_unknown_semester_text_is_null(client, db_session):
    _catalog(client, db_session)
    _designation(db_session, semester="ELECTIVO")
    body = client.get(VIEW, params={"academic_period": "I/2027", "date": "2027-03-01"}).json()
    assert body["semesters"] == [{"semester": None, "groups": ["M-1"]}]
    assert body["blocks"][0]["semester"] is None


def test_ambiguous_subject_name_has_no_code(client, db_session):
    program, _subject, _offering, _group, _room = _catalog(client, db_session)
    twin = _post(client, "/subjects", {"code": "MRF 9999", "name": "Anatomia humana i"})
    _post(client, "/offerings", {
        "subject_id": twin["id"], "program_id": program["id"], "academic_period": "I/2027",
        "semester": 1, "theory_hours": 10, "practice_hours": 0,
    })
    _designation(db_session)
    body = client.get(VIEW, params={"academic_period": "I/2027", "date": "2027-03-01"}).json()
    assert body["blocks"][0]["subject_code"] is None


def test_draft_source_shows_draft_blocks_with_current_teacher(client, db_session):
    program, _subject, offering, group, room = _catalog(client, db_session)
    _designation(db_session)
    draft, _block = _practice_draft(client, program, offering, group, room)
    body = client.get(VIEW, params={
        "academic_period": "I/2027", "source": "draft", "draft_id": draft["id"], "date": "2027-03-01",
    }).json()
    assert body["source"] == "draft"
    assert body["draft_id"] == draft["id"]
    assert [(b["source_type"], b["activity_type"], b["teacher_name"], b["classroom_code"], b["group_code"])
            for b in body["blocks"]] == [("draft", "practice", "Docente Practica", "LAB-1", "M-1")]


def test_draft_source_requires_an_existing_draft(client):
    missing = client.get(VIEW, params={"academic_period": "I/2027", "source": "draft"})
    assert missing.status_code == 422
    unknown = client.get(VIEW, params={"academic_period": "I/2027", "source": "draft", "draft_id": 999})
    assert unknown.status_code == 404


def test_export_has_one_grid_sheet_per_group_and_a_list(client, db_session):
    program, _subject, offering, group, room = _catalog(client, db_session)
    _designation(db_session)
    _designation(db_session, semester="SEGUNDO", group_code="T-2", subject="FISIOLOGÍA I",
                 schedule_json=[{"dia": "miercoles", "hora_inicio": "13:30", "hora_fin": "15:00",
                                 "horas_academicas": 2}])
    response = client.get(f"{VIEW}/export", params={"academic_period": "I/2027", "date": "2027-03-01"})
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert 'filename="horarios_I-2027.xlsx"' in response.headers["content-disposition"]
    workbook = openpyxl.load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == ["S1 M-1", "S2 T-2", "Lista"]
    grid = workbook["S1 M-1"]
    assert "I/2027" in str(grid["A1"].value) and "Semestre 1" in str(grid["A1"].value)
    header = [grid.cell(row=3, column=column).value for column in range(1, 8)]
    assert header == ["Hora", "Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado"]
    assert grid.cell(row=4, column=1).value == "06:30–07:00"
    assert "Anatomía Humana I" in grid.cell(row=4, column=2).value
    assert "DOCENTE TEORIA" in grid.cell(row=4, column=2).value
    # 06:30-08:55 covers five half-hour bands (06:30 .. 08:30).
    assert "B4:B8" in {str(item) for item in grid.merged_cells.ranges}
    listing = workbook["Lista"]
    assert [cell.value for cell in listing[1]] == [
        "Semestre", "Grupo", "Día", "Inicio", "Fin", "Materia", "Actividad", "Docente", "Aula",
    ]
    assert listing.max_row == 3


def test_export_filename_includes_filters(client, db_session):
    _catalog(client, db_session)
    _designation(db_session)
    response = client.get(f"{VIEW}/export", params={
        "academic_period": "I/2027", "date": "2027-03-01", "semester": 1, "group": "M-01",
    })
    assert response.status_code == 200, response.text
    assert 'filename="horarios_I-2027_sem1_M-1.xlsx"' in response.headers["content-disposition"]
    assert openpyxl.load_workbook(BytesIO(response.content)).sheetnames == ["S1 M-1", "Lista"]


def test_designation_names_resolve_to_the_curriculum_subject_of_their_semester(client, db_session):
    program, _anatomy, _offering, _group, _room = _catalog(client, db_session)
    for code, name in (("SCP 0300", "Promoción de la Salud en el Ciclo de Vida I"),
                       ("SCP 0400", "Promoción de la Salud en el Ciclo de Vida II"),
                       ("SCP 0700", "Antropología Médica e Interculturalidad")):
        subject = _post(client, "/subjects", {"code": code, "name": name})
        _post(client, "/offerings", {
            "subject_id": subject["id"], "program_id": program["id"], "academic_period": "I/2027",
            "semester": int(code[-3]), "theory_hours": 17, "practice_hours": 51,
        })
    _designation(db_session, subject="PROMOCIÓN DE LA SALUD I", semester="TERCERO", group_code="M-1")
    _designation(db_session, subject="PROMOCIÓN DE LA SALUD II", semester="CUARTO", group_code="M-2")
    _designation(db_session, subject="ANTROPOLOGÍA MÉDICA E INTERCULTURAL", semester="SEPTIMO", group_code="N-1")
    # Same words in a semester where the subject is not offered: stays unresolved.
    _designation(db_session, subject="PROMOCIÓN DE LA SALUD I", semester="QUINTO", group_code="M-3")
    body = client.get(VIEW, params={"academic_period": "I/2027", "date": "2027-03-01"}).json()
    resolved = {(b["semester"], b["subject_code"], b["subject_name"]) for b in body["blocks"]}
    assert (3, "SCP 0300", "Promoción de la Salud en el Ciclo de Vida I") in resolved
    assert (4, "SCP 0400", "Promoción de la Salud en el Ciclo de Vida II") in resolved
    assert (7, "SCP 0700", "Antropología Médica e Interculturalidad") in resolved
    assert (5, None, "PROMOCIÓN DE LA SALUD I") in resolved


def test_published_blocks_show_the_current_catalog_name(client, db_session):
    program, anatomy, offering, group, room = _catalog(client, db_session)
    draft, _block = _practice_draft(client, program, offering, group, room)
    _publish(client, draft["id"])
    renamed = client.put(f"{BASE}/subjects/{anatomy['id']}", json={"code": "MRF 0100", "name": "Anatomía Humana I (malla)"})
    assert renamed.status_code == 200, renamed.text
    body = client.get(VIEW, params={"academic_period": "I/2027", "date": "2027-03-01"}).json()
    assert [(b["subject_code"], b["subject_name"]) for b in body["blocks"]] == [("MRF 0100", "Anatomía Humana I (malla)")]
