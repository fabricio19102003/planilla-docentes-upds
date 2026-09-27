"""Read-only, digest-bound preview for the II/2026 designation bootstrap."""
from __future__ import annotations

import hashlib
import json
import re
import stat
import unicodedata
import zipfile
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, time
from io import BytesIO
from pathlib import PurePosixPath
from typing import Any, Literal

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from sqlalchemy.orm import Session

from app.models.academic_management import (
    AcademicGroup, AcademicProgram, AcademicSubject, Classroom, SubjectOffering,
    TeacherAvailability, AcademicScheduleAssignment, AcademicScheduleBlock,
    AcademicScheduleDraft, AcademicSchedulePublication,
    AcademicSchedulePublishedAssignment, AcademicSchedulePublishedBlock,
)
from app.models.teacher import Teacher
from app.services.designation_loader import normalize_name
from app.utils.helpers import normalize_group_code

POLICY_VERSION = "designation-bootstrap-preview-v4"
ALIAS_SCHEMA_VERSION = "designation-alias-v2"
LEGACY_ALIAS_SCHEMA_VERSION = "designation-alias-v1"
OFFICIAL_SHEETS = ("DOCENTES TEORICOS", "DOCENTES ASISTENCIALES")
SALARY_NAME = "NOMBRE COMPLETO"
SALARY_CI = "NO C.I."
SALARY_SUBJECT = "MATERIA"
SALARY_SEMESTER = "SEMESTRE"
SALARY_HOURS = "TOTAL HORAS"
SALARY_HEADERS = {SALARY_NAME, SALARY_CI, SALARY_SUBJECT, SALARY_SEMESTER, SALARY_HOURS}
SEMESTERS = {
    "PRIMERO": 1, "SEGUNDO": 2, "TERCERO": 3, "CUARTO": 4,
    "QUINTO": 5, "SEXTO": 6, "SEPTIMO": 7, "OCTAVO": 8,
}
DAYS = {
    "LUNES": "monday", "MARTES": "tuesday", "MIERCOLES": "wednesday",
    "JUEVES": "thursday", "VIERNES": "friday", "SABADO": "saturday",
}
DAY_RE = re.compile(r"(?i)\b(LUNES|MARTES|MI[EÉ]RCOLES|JUEVES|JUVES|VIERNES|S[AÁ]BADO)\b\s*(:?)")
TIME_RE = re.compile(r"(?i)(\d{1,2})\s*:\s*(\d{2})\s*(AM|PM)?\s*-\s*(\d{1,2})\s*:\s*(\d{2})\s*(AM|PM)?")
MAX_ZIP_MEMBERS = 128
MAX_ZIP_MEMBER_BYTES = 20 * 1024 * 1024
MAX_ZIP_TOTAL_BYTES = 64 * 1024 * 1024
MAX_ZIP_RATIO = 200
MAX_WORKSHEETS = 8
MAX_WORKSHEET_ROWS = 5000
MAX_WORKSHEET_COLUMNS = 64
MAX_WORKBOOK_CELLS = 100_000
MAX_CELL_STRING = 10_000
REQUIRED_XLSX_MEMBERS = {"[Content_Types].xml", "_rels/.rels", "xl/workbook.xml"}


class WorkbookValidationError(ValueError):
    def __init__(self, code: str):
        super().__init__(code)
        self.code = code


def _fold(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or "").strip().upper())
    return " ".join("".join(c for c in text if not unicodedata.combining(c)).split()).upper()


def _sha(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def _digest(value: object) -> str:
    return _sha(json.dumps(value, sort_keys=True, separators=(",", ":"), default=str).encode())


def _semester(value: object) -> int | None:
    folded = _fold(value)
    if folded in SEMESTERS:
        return SEMESTERS[folded]
    return int(folded) if folded.isdigit() and 0 < int(folded) <= 20 else None


def _subject_key(value: object) -> str:
    return re.sub(r"\s+PRACTICA$", "", _fold(value))


def _ci_key(value: object) -> str:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", "", str(value or "").strip())


def _clock(hour: str, minute: str, marker: str | None) -> time | None:
    h, m = int(hour), int(minute)
    if m > 59 or h > 23 or (marker and not 1 <= h <= 12):
        return None
    if marker and h <= 12:
        h = h % 12 + (12 if marker.upper() == "PM" else 0)
    return time(h, m)


def preflight_xlsx(content: bytes) -> None:
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            members = archive.infolist()
            if not members or len(members) > MAX_ZIP_MEMBERS:
                raise WorkbookValidationError("unsafe_workbook_archive")
            names = {item.filename for item in members}
            if len(names) != len(members) or not REQUIRED_XLSX_MEMBERS <= names:
                raise WorkbookValidationError("unsupported_workbook")
            total = 0
            for item in members:
                path = PurePosixPath(item.filename)
                mode = item.external_attr >> 16
                if (path.is_absolute() or ".." in path.parts or "\\" in item.filename
                        or item.flag_bits & 0x1 or stat.S_ISLNK(mode)):
                    raise WorkbookValidationError("unsafe_workbook_archive")
                if item.file_size > MAX_ZIP_MEMBER_BYTES:
                    raise WorkbookValidationError("workbook_member_too_large")
                total += item.file_size
                if total > MAX_ZIP_TOTAL_BYTES:
                    raise WorkbookValidationError("workbook_expanded_too_large")
                if item.file_size and item.file_size / max(item.compress_size, 1) > MAX_ZIP_RATIO:
                    raise WorkbookValidationError("workbook_compression_ratio")
    except WorkbookValidationError:
        raise
    except (zipfile.BadZipFile, OSError, ValueError) as exc:
        raise WorkbookValidationError("unsupported_workbook") from exc


def _validate_workbook_bounds(workbook: Any) -> None:
    if len(workbook.worksheets) > MAX_WORKSHEETS:
        raise WorkbookValidationError("workbook_dimensions_exceeded")
    cells = 0
    for sheet in workbook.worksheets:
        if sheet.max_row > MAX_WORKSHEET_ROWS or sheet.max_column > MAX_WORKSHEET_COLUMNS:
            raise WorkbookValidationError("workbook_dimensions_exceeded")
        cells += sheet.max_row * sheet.max_column
        if cells > MAX_WORKBOOK_CELLS:
            raise WorkbookValidationError("workbook_dimensions_exceeded")
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and len(cell.value) > MAX_CELL_STRING:
                    raise WorkbookValidationError("workbook_cell_too_large")


@dataclass(frozen=True)
class Slot:
    weekday: str
    start: time
    end: time
    raw: str
    normalized: str
    normalizations: tuple[str, ...]


@dataclass(frozen=True)
class SourceRow:
    workbook_sha256: str
    sheet: str
    row: int
    raw: tuple[object, ...]


@dataclass(frozen=True)
class OfficialRow:
    activity: Literal["theory", "practice"]
    teacher_raw: str
    teacher_key: str
    subject_raw: str
    subject_key: str
    semester: int | None
    group_raw: str
    group_key: str
    classroom_raw: str | None
    hours: int | None
    slots: tuple[Slot, ...]
    source: SourceRow


@dataclass(frozen=True)
class SalaryRow:
    teacher_raw: str
    teacher_key: str
    ci_raw: str
    ci_key: str
    subject_raw: str
    subject_key: str
    semester: int | None
    hours: int | None
    source: SourceRow


@dataclass(frozen=True)
class PlannedBlock:
    activity: Literal["theory", "practice"]
    teacher_key: str
    teacher_ci: str | None
    subject_key: str
    subject_id: int | None
    offering_id: int | None
    semester: int | None
    group_key: str
    group_id: int | None
    classroom_key: str | None
    classroom_id: int | None
    slot: Slot
    source: SourceRow


@dataclass(frozen=True)
class ParseStats:
    physical: int = 0
    parsed: int = 0
    skipped: int = 0
    errors: int = 0


@dataclass(frozen=True)
class AliasResolution:
    sha256: str | None
    teacher: dict[str, tuple[str, str]]
    subject: dict[tuple[str, int], str]
    errors: Counter[str]


def _closed_json_object(pairs: list[tuple[str, Any]]) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for key, value in pairs:
        if key in result:
            raise ValueError("duplicate JSON key")
        result[key] = value
    return result


def _parse_alias_resolution(
    content: bytes | None, *, official_sha256: str, salary_sha256: str,
    academic_period: str, effective_date: date,
    practice: list[OfficialRow], salary: list[SalaryRow],
) -> AliasResolution:
    if content is None:
        return AliasResolution(None, {}, {}, Counter())
    alias_hash = _sha(content)
    errors: Counter[str] = Counter()
    try:
        payload = json.loads(content.decode("utf-8"), object_pairs_hook=_closed_json_object)
    except ValueError as exc:
        if str(exc) == "duplicate JSON key":
            return AliasResolution(alias_hash, {}, {}, Counter({"alias_duplicate_key": 1}))
        return AliasResolution(alias_hash, {}, {}, Counter({"alias_invalid_json": 1}))
    except UnicodeDecodeError:
        return AliasResolution(alias_hash, {}, {}, Counter({"alias_invalid_json": 1}))
    required = {
        "schema_version", "official_sha256", "salary_sha256", "academic_period",
        "effective_date", "teacher_aliases", "subject_aliases",
    }
    if not isinstance(payload, dict) or set(payload) != required:
        return AliasResolution(alias_hash, {}, {}, Counter({"alias_schema_invalid": 1}))
    schema_version = payload.get("schema_version")
    if schema_version not in {LEGACY_ALIAS_SCHEMA_VERSION, ALIAS_SCHEMA_VERSION}:
        errors["alias_schema_version"] += 1
    if payload.get("official_sha256") != official_sha256:
        errors["alias_official_hash_mismatch"] += 1
    if payload.get("salary_sha256") != salary_sha256:
        errors["alias_salary_hash_mismatch"] += 1
    if payload.get("academic_period") != academic_period:
        errors["alias_period_mismatch"] += 1
    if payload.get("effective_date") != effective_date.isoformat():
        errors["alias_effective_date_mismatch"] += 1
    raw_teachers = payload.get("teacher_aliases")
    raw_subjects = payload.get("subject_aliases")
    if (not isinstance(raw_teachers, dict)
            or (schema_version == LEGACY_ALIAS_SCHEMA_VERSION and not isinstance(raw_subjects, dict))
            or (schema_version == ALIAS_SCHEMA_VERSION and not isinstance(raw_subjects, list))):
        errors["alias_schema_invalid"] += 1
        return AliasResolution(alias_hash, {}, {}, errors)

    official_teachers = {row.teacher_key for row in practice}
    salary_identities = {(row.teacher_key, row.ci_key) for row in salary if row.ci_key}
    teacher_map: dict[str, tuple[str, str]] = {}
    teacher_target_owners = {
        identity: official_key
        for official_key in official_teachers
        for identity in salary_identities
        if identity[0] == official_key
    }
    for official_key, target in raw_teachers.items():
        if (not isinstance(official_key, str) or official_key != normalize_name(official_key)
                or not isinstance(target, dict)
                or set(target) != {"salary_teacher_key", "salary_ci"}):
            errors["teacher_alias_schema_invalid"] += 1
            continue
        salary_teacher = target.get("salary_teacher_key")
        salary_ci = target.get("salary_ci")
        if (not isinstance(salary_teacher, str) or salary_teacher != normalize_name(salary_teacher)
                or not isinstance(salary_ci, str) or not salary_ci
                or salary_ci != _ci_key(salary_ci)):
            errors["teacher_alias_schema_invalid"] += 1
            continue
        identity = (salary_teacher, salary_ci)
        if official_key not in official_teachers or identity not in salary_identities:
            errors["teacher_alias_orphan"] += 1
            continue
        target_owner = teacher_target_owners.get(identity)
        if target_owner is not None and target_owner != official_key:
            errors["teacher_alias_many_to_one"] += 1
            continue
        teacher_target_owners[identity] = official_key
        teacher_map[official_key] = identity

    official_subjects = {(row.subject_key, row.semester) for row in practice if row.semester}
    salary_subjects = {(row.subject_key, row.semester) for row in salary if row.semester}
    subject_map: dict[tuple[str, int], str] = {}
    subject_target_owners = {
        source: source
        for source in salary_subjects & official_subjects
    }
    if schema_version == LEGACY_ALIAS_SCHEMA_VERSION:
        subject_entries = []
        for salary_key, target in raw_subjects.items():
            if (isinstance(target, dict)
                    and set(target) == {"official_subject_key", "semester", "activity_type"}):
                subject_entries.append({"salary_subject_key": salary_key, **target})
            else:
                subject_entries.append(target)
    else:
        subject_entries = raw_subjects
    seen_subject_sources: set[tuple[str, int]] = set()
    for entry in subject_entries:
        if not isinstance(entry, dict) or set(entry) != {
            "salary_subject_key", "official_subject_key", "semester", "activity_type",
        }:
            errors["subject_alias_schema_invalid"] += 1
            continue
        salary_key = entry.get("salary_subject_key")
        target = entry
        if (not isinstance(salary_key, str) or salary_key != _subject_key(salary_key)
                or not isinstance(target, dict)):
            errors["subject_alias_schema_invalid"] += 1
            continue
        official_key = target.get("official_subject_key")
        semester = target.get("semester")
        activity = target.get("activity_type")
        if (not isinstance(official_key, str) or official_key != _subject_key(official_key)
                or isinstance(semester, bool) or not isinstance(semester, int)
                or activity != "practice"):
            errors["subject_alias_schema_invalid"] += 1
            continue
        source = (salary_key, semester)
        destination = (official_key, semester)
        if source in seen_subject_sources:
            errors["subject_alias_duplicate_source"] += 1
            continue
        seen_subject_sources.add(source)
        if source not in salary_subjects or destination not in official_subjects:
            salary_other = any(key == salary_key for key, _semester_value in salary_subjects)
            official_other = any(key == official_key for key, _semester_value in official_subjects)
            errors["subject_alias_cross_semester" if salary_other and official_other else "subject_alias_orphan"] += 1
            continue
        target_owner = subject_target_owners.get(destination)
        if target_owner is not None and target_owner != source:
            errors["subject_alias_many_to_one"] += 1
            continue
        subject_target_owners[destination] = source
        subject_map[source] = official_key
    if errors:
        return AliasResolution(alias_hash, {}, {}, errors)
    return AliasResolution(alias_hash, teacher_map, subject_map, errors)


def _schedule(value: object) -> tuple[tuple[Slot, ...], list[str], list[str]]:
    raw = str(value or "").strip()
    if not raw:
        return (), ["blank_schedule"], []
    matches = list(DAY_RE.finditer(raw))
    if not matches:
        return (), ["invalid_schedule"], []
    if raw[:matches[0].start()].strip(" \n\t;,"):
        return (), ["invalid_schedule"], []
    slots: list[Slot] = []
    errors: list[str] = []
    warnings: list[str] = []
    for index, day_match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(raw)
        raw_segment = raw[day_match.start():end].strip()
        segment = raw[day_match.end():end].strip()
        normalizations: list[str] = []
        if not day_match.group(2):
            normalizations.append("day_colon_inserted")
        day_key = _fold(day_match.group(1))
        if day_key == "JUVES":
            day_key = "JUEVES"
            normalizations.append("known_day_typo_corrected")
        if re.search(r"[ \t]{2,}", raw_segment):
            normalizations.append("repeated_whitespace_normalized")
        if "--" in segment:
            if segment.count("--") != 1:
                errors.append("invalid_schedule")
                continue
            segment = segment.replace("--", "-")
            normalizations.append("double_hyphen_normalized")
        collapsed = re.sub(r"[ \t]{2,}", " ", segment)
        if collapsed != segment:
            segment = collapsed
            if "repeated_whitespace_normalized" not in normalizations:
                normalizations.append("repeated_whitespace_normalized")
        time_matches = list(TIME_RE.finditer(segment))
        if len(time_matches) != 1 or segment[:time_matches[0].start()].strip(" \n\t;,"):
            errors.append("invalid_schedule")
            continue
        match = time_matches[0]
        if segment[match.end():].strip(" \n\t;,"):
            errors.append("invalid_schedule")
            continue
        start_hour, end_hour = int(match.group(1)), int(match.group(4))
        start_marker, end_marker = match.group(3), match.group(6)
        if (start_marker is None) != (end_marker is None):
            errors.append("ambiguous_meridiem")
            continue
        if (start_hour > 12 or end_hour > 12) and start_marker and end_marker:
            marked_endpoints = ((start_hour, start_marker), (end_hour, end_marker))
            markers_are_redundant = all(
                hour > 12
                or (hour < 12 and marker.upper() == "AM")
                or (hour == 12 and marker.upper() == "PM")
                for hour, marker in marked_endpoints
            )
            if not markers_are_redundant:
                errors.append("ambiguous_meridiem")
                continue
            start_marker = end_marker = None
            normalizations.append("redundant_meridiem_on_24h_range")
        start = _clock(match.group(1), match.group(2), start_marker)
        finish = _clock(match.group(4), match.group(5), end_marker)
        if start is None or finish is None or finish <= start:
            errors.append("impossible_interval")
            continue
        duration = finish.hour * 60 + finish.minute - start.hour * 60 - start.minute
        if duration > 360:
            errors.append("impossible_interval")
            continue
        warnings.extend(normalizations)
        normalized = f"{day_key}: {start.strftime('%H:%M')}-{finish.strftime('%H:%M')}"
        slots.append(Slot(DAYS[day_key], start, finish, raw_segment, normalized, tuple(normalizations)))
    return tuple(slots), errors, warnings


def _integer(value: object) -> int | None:
    if isinstance(value, bool) or value in (None, ""):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return int(number) if number.is_integer() and number >= 0 else None


def _merged_required(sheet: Any, row: int, columns: tuple[int, ...]) -> bool:
    for column in columns:
        cell = sheet.cell(row, column)
        if isinstance(cell, MergedCell) or any(cell.coordinate in area for area in sheet.merged_cells.ranges):
            return True
    return False


def _parse_official(
    content: bytes,
) -> tuple[list[OfficialRow], Counter[str], Counter[str], Counter[str], dict[str, ParseStats], int]:
    preflight_xlsx(content)
    workbook = load_workbook(BytesIO(content), data_only=False, keep_links=False)
    _validate_workbook_bounds(workbook)
    errors: Counter[str] = Counter()
    warnings: Counter[str] = Counter()
    activity_error_counts: Counter[str] = Counter()
    stats: dict[str, ParseStats] = {}
    if tuple(workbook.sheetnames) != OFFICIAL_SHEETS:
        errors["official_sheet_structure"] += 1
    rows: list[OfficialRow] = []
    source_hash = _sha(content)
    definitions = {
        "DOCENTES TEORICOS": ("theory", ["N°", "", "MATERIAS", "SEMESTRE", "M-04", "AULA", "HORARIO"]),
        "DOCENTES ASISTENCIALES": ("practice", ["N°", "DOCENTE", "MATERIAS", "SEMESTRE", "GRUPO", "HORARIO"]),
    }
    for sheet_name, (activity, expected) in definitions.items():
        if sheet_name not in workbook.sheetnames:
            continue
        sheet = workbook[sheet_name]
        actual = [_fold(sheet.cell(1, column).value) for column in range(1, len(expected) + 1)]
        required_headers = [_fold(item) for item in expected if item]
        header_values = [_fold(sheet.cell(1, column).value) for column in range(1, sheet.max_column + 1)]
        duplicate_headers = any(header_values.count(header) != 1 for header in required_headers)
        if actual != [_fold(item) for item in expected] or duplicate_headers:
            errors[f"{activity}_header_structure"] += 1
            activity_error_counts[activity] += 1
            stats[activity] = ParseStats(physical=max(sheet.max_row - 1, 0), errors=max(sheet.max_row - 1, 0))
            continue
        physical = parsed = skipped = error_rows = 0
        required_columns = (2, 3, 4, 5, 7) if activity == "theory" else (2, 3, 4, 5, 6)
        for number in range(2, sheet.max_row + 1):
            physical += 1
            values = tuple(sheet.cell(number, column).value for column in range(1, sheet.max_column + 1))
            if not any(item not in (None, "") for item in values):
                skipped += 1
                continue
            if _merged_required(sheet, number, required_columns):
                errors["merged_required_official_cell"] += 1
                activity_error_counts[activity] += 1
                error_rows += 1
                continue
            teacher, subject, semester, group = values[1:5]
            schedule_value = values[6] if activity == "theory" else values[5]
            slots, schedule_errors, schedule_warnings = _schedule(schedule_value)
            errors.update(schedule_errors)
            warnings.update(schedule_warnings)
            activity_error_counts[activity] += len(schedule_errors)
            row_invalid = bool(schedule_errors)
            core_invalid = False
            if schedule_errors:
                slots = ()
            semester_value = _semester(semester)
            if semester_value is None:
                errors["invalid_semester"] += 1
                activity_error_counts[activity] += 1
                row_invalid = True
                core_invalid = True
            if not all(str(item or "").strip() for item in (teacher, subject, group)):
                errors["missing_required_official_value"] += 1
                activity_error_counts[activity] += 1
                row_invalid = True
                core_invalid = True
            if row_invalid:
                error_rows += 1
            else:
                parsed += 1
            if core_invalid:
                continue
            rows.append(OfficialRow(
                activity=activity, teacher_raw=str(teacher or "").strip(),
                teacher_key=normalize_name(str(teacher or "")),
                subject_raw=str(subject or "").strip(), subject_key=_subject_key(subject),
                semester=semester_value, group_raw=str(group or "").strip(),
                group_key=_fold(normalize_group_code(str(group or ""))),
                classroom_raw=str(values[5]).strip() if activity == "theory" and values[5] not in (None, "") else None,
                hours=_integer(values[10] if activity == "theory" else values[7]), slots=slots,
                source=SourceRow(source_hash, sheet_name, number, values),
            ))
        stats[activity] = ParseStats(physical, parsed, skipped, error_rows)
        if physical != parsed + skipped + error_rows:
            raise AssertionError("official row accounting invariant failed")
    workbook.close()
    return rows, errors, warnings, activity_error_counts, stats, len(workbook.sheetnames)


def _parse_salary(content: bytes) -> tuple[list[SalaryRow], Counter[str], ParseStats, int]:
    preflight_xlsx(content)
    workbook = load_workbook(BytesIO(content), data_only=False, keep_links=False)
    _validate_workbook_bounds(workbook)
    errors: Counter[str] = Counter()
    if len(workbook.sheetnames) != 1:
        errors["salary_sheet_structure"] += 1
    sheet = workbook[workbook.sheetnames[0]]
    header_rows = []
    for row in range(1, sheet.max_row + 1):
        values = {_fold(sheet.cell(row, column).value) for column in range(1, sheet.max_column + 1)}
        if SALARY_HEADERS <= values:
            header_rows.append(row)
    if len(header_rows) != 1:
        errors["salary_header_structure"] += 1
        workbook.close()
        return [], errors, ParseStats(errors=max(sheet.max_row - 1, 0)), len(workbook.sheetnames)
    header_row = header_rows[0]
    header_values = [_fold(sheet.cell(header_row, column).value) for column in range(1, sheet.max_column + 1)]
    if any(header_values.count(header) != 1 for header in SALARY_HEADERS):
        errors["duplicate_salary_header"] += 1
        workbook.close()
        return [], errors, ParseStats(physical=max(sheet.max_row - header_row, 0), errors=max(sheet.max_row - header_row, 0)), len(workbook.sheetnames)
    columns = {value: index + 1 for index, value in enumerate(header_values) if value}
    result: list[SalaryRow] = []
    source_hash = _sha(content)
    physical = parsed = skipped = error_rows = 0
    required_columns = tuple(columns[name] for name in SALARY_HEADERS)
    for number in range(header_row + 1, sheet.max_row + 1):
        physical += 1
        values = tuple(sheet.cell(number, column).value for column in range(1, sheet.max_column + 1))
        name = sheet.cell(number, columns[SALARY_NAME]).value
        ci = sheet.cell(number, columns[SALARY_CI]).value
        subject = sheet.cell(number, columns[SALARY_SUBJECT]).value
        semester_raw = sheet.cell(number, columns[SALARY_SEMESTER]).value
        hours_raw = sheet.cell(number, columns[SALARY_HOURS]).value
        if not any(item not in (None, "") for item in (ci, subject, semester_raw, hours_raw)):
            skipped += 1
            continue
        if _merged_required(sheet, number, required_columns):
            errors["merged_required_salary_cell"] += 1
            error_rows += 1
            continue
        semester = _semester(semester_raw)
        hours = _integer(hours_raw)
        row_invalid = False
        if not all(str(item or "").strip() for item in (name, subject)):
            errors["missing_required_salary_value"] += 1
            row_invalid = True
        if semester is None:
            errors["invalid_salary_semester"] += 1
            row_invalid = True
        if hours is None:
            errors["invalid_salary_hours"] += 1
            row_invalid = True
        if row_invalid:
            error_rows += 1
            continue
        parsed += 1
        result.append(SalaryRow(
            teacher_raw=str(name or "").strip(), teacher_key=normalize_name(str(name or "")),
            ci_raw=str(ci or "").strip(), ci_key=_ci_key(ci), subject_raw=str(subject or "").strip(),
            subject_key=_subject_key(subject), semester=semester, hours=hours,
            source=SourceRow(source_hash, sheet.title, number, values),
        ))
    if physical != parsed + skipped + error_rows:
        raise AssertionError("salary row accounting invariant failed")
    workbook.close()
    return result, errors, ParseStats(physical, parsed, skipped, error_rows), len(workbook.sheetnames)


def _unique_map(items: list[Any], keys: Any) -> tuple[dict[Any, Any], set[Any]]:
    grouped: dict[Any, dict[Any, Any]] = defaultdict(dict)
    for item in items:
        for key in keys(item):
            grouped[key][tuple(getattr(item, column.name) for column in item.__table__.primary_key.columns)] = item
    ambiguous = {key for key, values in grouped.items() if len(values) > 1}
    return {key: next(iter(values.values())) for key, values in grouped.items() if len(values) == 1}, ambiguous


def _rows_state(rows: list[Any]) -> list[dict[str, Any]]:
    return [
        {column.name: getattr(row, column.name) for column in row.__table__.columns}
        for row in rows
    ]


def _state_fingerprint(
    db: Session, period: str, program_id: int | None,
    official: list[OfficialRow], salary: list[SalaryRow],
    excluded_draft_id: int | None = None,
) -> str:
    programs = db.query(AcademicProgram).order_by(AcademicProgram.id).all()
    offerings = []
    groups = []
    subjects = []
    if program_id is not None:
        offerings = db.query(SubjectOffering).filter(
            SubjectOffering.program_id == program_id,
            SubjectOffering.academic_period == period,
        ).order_by(SubjectOffering.id).all()
        subject_ids = sorted({item.subject_id for item in offerings})
        subjects = db.query(AcademicSubject).filter(
            AcademicSubject.id.in_(subject_ids)
        ).order_by(AcademicSubject.id).all() if subject_ids else []
        groups = db.query(AcademicGroup).filter(
            AcademicGroup.program_id == program_id,
            AcademicGroup.academic_period == period,
        ).order_by(AcademicGroup.id).all()
    classroom_keys = {_fold(row.classroom_raw) for row in official if row.classroom_raw}
    classrooms = [item for item in db.query(Classroom).order_by(Classroom.id).all()
                  if _fold(item.code) in classroom_keys or _fold(item.name) in classroom_keys]
    teacher_names = {row.teacher_key for row in official}
    teacher_cis = {row.ci_key for row in salary if row.ci_key}
    teachers = [item for item in db.query(Teacher).order_by(Teacher.ci).all()
                if normalize_name(item.full_name) in teacher_names or _ci_key(item.ci) in teacher_cis]
    cis = [item.ci for item in teachers]
    availability = db.query(TeacherAvailability).filter(
        TeacherAvailability.academic_period == period,
        TeacherAvailability.teacher_ci.in_(cis),
    ).order_by(TeacherAvailability.id).all() if cis else []
    drafts = []
    draft_blocks = []
    draft_assignments = []
    publications = []
    published_blocks = []
    published_assignments = []
    if program_id is not None:
        draft_query = db.query(AcademicScheduleDraft).filter(
            AcademicScheduleDraft.program_id == program_id,
            AcademicScheduleDraft.academic_period == period,
        )
        if excluded_draft_id is not None:
            draft_query = draft_query.filter(AcademicScheduleDraft.id != excluded_draft_id)
        drafts = draft_query.order_by(AcademicScheduleDraft.id).all()
        draft_ids = [item.id for item in drafts]
        draft_blocks = db.query(AcademicScheduleBlock).filter(
            AcademicScheduleBlock.draft_id.in_(draft_ids)
        ).order_by(AcademicScheduleBlock.id).all() if draft_ids else []
        block_ids = [item.id for item in draft_blocks]
        draft_assignments = db.query(AcademicScheduleAssignment).filter(
            AcademicScheduleAssignment.block_id.in_(block_ids)
        ).order_by(AcademicScheduleAssignment.id).all() if block_ids else []
        publication_query = db.query(AcademicSchedulePublication).filter(
            AcademicSchedulePublication.program_id == program_id,
            AcademicSchedulePublication.academic_period == period,
        )
        if excluded_draft_id is not None:
            publication_query = publication_query.filter(
                AcademicSchedulePublication.source_draft_id != excluded_draft_id
            )
        publications = publication_query.order_by(AcademicSchedulePublication.id).all()
        publication_ids = [item.id for item in publications]
        published_blocks = db.query(AcademicSchedulePublishedBlock).filter(
            AcademicSchedulePublishedBlock.publication_id.in_(publication_ids)
        ).order_by(AcademicSchedulePublishedBlock.id).all() if publication_ids else []
        published_block_ids = [item.id for item in published_blocks]
        published_assignments = db.query(AcademicSchedulePublishedAssignment).filter(
            AcademicSchedulePublishedAssignment.publication_block_id.in_(published_block_ids)
        ).order_by(AcademicSchedulePublishedAssignment.id).all() if published_block_ids else []
    return _digest({
        "programs": _rows_state(programs), "subjects": _rows_state(subjects),
        "offerings": _rows_state(offerings), "groups": _rows_state(groups),
        "classrooms": _rows_state(classrooms), "teachers": _rows_state(teachers),
        "availability": _rows_state(availability),
        "drafts": _rows_state(drafts), "draft_blocks": _rows_state(draft_blocks),
        "draft_assignments": _rows_state(draft_assignments),
        "publications": _rows_state(publications),
        "published_blocks": _rows_state(published_blocks),
        "published_assignments": _rows_state(published_assignments),
    })


def _build_designation_bootstrap_preview(
    db: Session, *, official_content: bytes, salary_content: bytes,
    academic_period: str, effective_date: date, program_identity: str,
    alias_content: bytes | None, include_private: bool = False,
    excluded_draft_id: int | None = None,
) -> dict[str, Any]:
    blockers: Counter[str] = Counter()
    warnings: Counter[str] = Counter()
    source_hashes = {"official": _sha(official_content), "salary": _sha(salary_content)}
    academic_period = " ".join(str(academic_period or "").split()).upper()
    program_identity = " ".join(str(program_identity or "").split())
    if not academic_period or len(academic_period) > 30:
        blockers["invalid_academic_period"] += 1
    if not program_identity or len(program_identity) > 200:
        blockers["invalid_program_identity"] += 1
    if effective_date != date(2026, 8, 21):
        blockers["invalid_effective_date"] += 1
    try:
        official, official_errors, schedule_warnings, activity_error_counts, official_stats, official_sheet_count = _parse_official(official_content)
    except WorkbookValidationError as exc:
        official, official_errors, schedule_warnings, activity_error_counts, official_stats, official_sheet_count = (
            [], Counter({exc.code: 1}), Counter(), Counter({"theory": 1}), {}, 0,
        )
    except Exception:
        official, official_errors, schedule_warnings, activity_error_counts, official_stats, official_sheet_count = (
            [], Counter({"unsupported_official_workbook": 1}), Counter(), Counter({"theory": 1}), {}, 0,
        )
    try:
        salary, salary_errors, salary_stats, salary_sheet_count = _parse_salary(salary_content)
    except WorkbookValidationError as exc:
        salary, salary_errors, salary_stats, salary_sheet_count = [], Counter({exc.code: 1}), ParseStats(), 0
    except Exception:
        salary, salary_errors, salary_stats, salary_sheet_count = [], Counter({"unsupported_salary_workbook": 1}), ParseStats(), 0
    blockers.update(official_errors)
    blockers.update(salary_errors)
    warnings.update(schedule_warnings)
    if not any(stats.parsed for stats in official_stats.values()):
        blockers["no_parsed_official_rows"] += 1

    programs = db.query(AcademicProgram).filter(
        AcademicProgram.active.is_(True)
    ).order_by(AcademicProgram.id).all()
    matching_programs = [item for item in programs if (
        _fold(item.code) == _fold(program_identity) or _fold(item.name) == _fold(program_identity)
    )]
    program = matching_programs[0] if len(matching_programs) == 1 else None
    blockers["missing_program" if not matching_programs else "ambiguous_program"] += int(len(matching_programs) != 1)
    program_id = program.id if program else None

    practice = [row for row in official if row.activity == "practice"]
    relevant_salary = salary if practice else []
    aliases = _parse_alias_resolution(
        alias_content,
        official_sha256=source_hashes["official"], salary_sha256=source_hashes["salary"],
        academic_period=academic_period, effective_date=effective_date,
        practice=practice, salary=relevant_salary,
    )
    blockers.update(aliases.errors)

    # Source corroboration is deliberately completed before any catalog lookup.
    salary_keys: dict[tuple[str, str, int], list[SalaryRow]] = defaultdict(list)
    for row in relevant_salary:
        if row.teacher_key and row.subject_key and row.semester is not None:
            canonical_subject = aliases.subject.get(
                (row.subject_key, row.semester), row.subject_key,
            )
            salary_keys[(row.teacher_key, canonical_subject, row.semester)].append(row)

    practice_keys: dict[tuple[str, str, int], list[OfficialRow]] = defaultdict(list)
    salary_teacher_keys = {row.teacher_key for row in relevant_salary}
    salary_subject_keys = {(row.subject_key, row.semester) for row in relevant_salary if row.semester}
    unresolved_teacher_alias = unresolved_subject_alias = 0
    for row in practice:
        if row.teacher_key and row.subject_key and row.semester is not None:
            teacher_alias = aliases.teacher.get(row.teacher_key)
            canonical_teacher = teacher_alias[0] if teacher_alias else row.teacher_key
            practice_keys[(canonical_teacher, row.subject_key, row.semester)].append(row)
            if teacher_alias is None and row.teacher_key not in salary_teacher_keys:
                unresolved_teacher_alias += 1
            direct_subject = (row.subject_key, row.semester) in salary_subject_keys
            aliased_subject = any(
                official_key == row.subject_key and semester == row.semester
                for (salary_key, semester), official_key in aliases.subject.items()
            )
            if not direct_subject and not aliased_subject:
                unresolved_subject_alias += 1
    blockers["unresolved_teacher_alias"] += unresolved_teacher_alias
    blockers["unresolved_subject_alias"] += unresolved_subject_alias
    join_covered = join_missing = join_ambiguous = missing_ci = conflicting_ci = 0
    valid_join_ci: dict[tuple[str, str, int], str] = {}
    for row in practice:
        teacher_alias = aliases.teacher.get(row.teacher_key)
        canonical_teacher = teacher_alias[0] if teacher_alias else row.teacher_key
        key = (canonical_teacher, row.subject_key, row.semester or 0)
        matches = salary_keys.get(key, [])
        if not matches:
            join_missing += 1
            continue
        cis = {item.ci_key for item in matches if item.ci_key}
        if any(not item.ci_key for item in matches) or not cis:
            missing_ci += 1
        elif len(cis) > 1:
            conflicting_ci += 1
            join_ambiguous += 1
        else:
            canonical_ci = next(iter(cis))
            if teacher_alias and canonical_ci != teacher_alias[1]:
                conflicting_ci += 1
                join_ambiguous += 1
            else:
                join_covered += 1
                valid_join_ci[key] = canonical_ci
    blockers["practice_join_missing"] += join_missing
    blockers["practice_join_ambiguous"] += join_ambiguous
    blockers["practice_join_missing_ci"] += missing_ci
    blockers["practice_join_conflicting_ci"] += conflicting_ci
    orphan_salary = len(set(salary_keys) - set(practice_keys))
    blockers["orphan_salary_assignment"] += orphan_salary
    comparable_source_keys = set(practice_keys) & set(salary_keys)
    salary_hours_noncomparable = len(comparable_source_keys)
    warnings["salary_payment_hours_noncomparable"] += salary_hours_noncomparable

    subject_map: dict[str, AcademicSubject] = {}
    offering_map: dict[tuple[str, int], SubjectOffering] = {}
    group_map: dict[tuple[str, int], AcademicGroup] = {}
    classroom_map: dict[str, Classroom] = {}
    ambiguous_subjects: set[str] = set()
    ambiguous_offerings: set[tuple[str, int]] = set()
    ambiguous_groups: set[tuple[str, int]] = set()
    ambiguous_classrooms: set[str] = set()
    if program_id is not None:
        offerings = db.query(SubjectOffering).join(AcademicSubject).filter(
            SubjectOffering.program_id == program_id,
            SubjectOffering.academic_period == academic_period,
            SubjectOffering.active.is_(True),
            AcademicSubject.active.is_(True),
        ).order_by(SubjectOffering.id).all()
        subject_ids = sorted({item.subject_id for item in offerings})
        subjects = db.query(AcademicSubject).filter(
            AcademicSubject.id.in_(subject_ids), AcademicSubject.active.is_(True),
        ).order_by(AcademicSubject.id).all() if subject_ids else []
        subject_map, ambiguous_subjects = _unique_map(
            subjects, lambda item: {_subject_key(item.code), _subject_key(item.name)},
        )
        subjects_by_id = {item.id: item for item in subjects}
        offering_map, ambiguous_offerings = _unique_map(
            offerings, lambda item: {
                (key, item.semester)
                for key in (
                    _subject_key(subjects_by_id[item.subject_id].code),
                    _subject_key(subjects_by_id[item.subject_id].name),
                )
            },
        )
        groups = db.query(AcademicGroup).filter(
            AcademicGroup.program_id == program_id,
            AcademicGroup.academic_period == academic_period,
            AcademicGroup.active.is_(True),
        ).order_by(AcademicGroup.id).all()
        group_map, ambiguous_groups = _unique_map(
            groups, lambda item: {(_fold(normalize_group_code(item.code)), item.semester)},
        )
        classrooms = db.query(Classroom).filter(
            Classroom.active.is_(True)
        ).order_by(Classroom.id).all()
        classroom_map, ambiguous_classrooms = _unique_map(
            classrooms, lambda item: {_fold(item.code), _fold(item.name)},
        )
    blockers["ambiguous_subject"] += len(ambiguous_subjects)
    blockers["ambiguous_offering"] += len(ambiguous_offerings)
    blockers["ambiguous_group"] += len(ambiguous_groups)
    blockers["ambiguous_classroom"] += len(ambiguous_classrooms)

    teachers = db.query(Teacher).order_by(Teacher.ci).all()
    ci_map, ambiguous_teacher_cis = _unique_map(teachers, lambda item: {_ci_key(item.ci)})
    name_map, ambiguous_teacher_names = _unique_map(
        teachers, lambda item: {normalize_name(item.full_name)},
    )
    blockers["ambiguous_teacher_ci"] += len(ambiguous_teacher_cis)
    for row in relevant_salary:
        canonical_subject = aliases.subject.get(
            (row.subject_key, row.semester), row.subject_key,
        )
        if canonical_subject in ambiguous_subjects:
            blockers["ambiguous_salary_subject"] += 1
        elif canonical_subject not in subject_map:
            blockers["missing_salary_subject"] += 1

    plan: list[PlannedBlock] = []
    teacher_resolved = teacher_missing = teacher_ambiguous = 0
    subject_resolved = offering_resolved = group_resolved = classroom_resolved = 0
    availability_missing = 0
    for row in official:
        if row.subject_key in ambiguous_subjects:
            blockers["ambiguous_subject_reference"] += 1
            subject = None
        else:
            subject = subject_map.get(row.subject_key)
        if subject is None and row.subject_key not in ambiguous_subjects:
            blockers["missing_subject"] += 1
        else:
            subject_resolved += int(subject is not None)
        offering_key = (row.subject_key, row.semester) if subject and row.semester else None
        offering = offering_map.get(offering_key) if offering_key not in ambiguous_offerings else None
        if offering_key in ambiguous_offerings:
            blockers["ambiguous_offering_reference"] += 1
        elif offering is None:
            blockers["missing_offering"] += 1
        else:
            offering_resolved += 1
        group_key = (row.group_key, row.semester) if row.semester else None
        group = group_map.get(group_key) if group_key not in ambiguous_groups else None
        if group_key in ambiguous_groups:
            blockers["ambiguous_group_reference"] += 1
        elif group is None:
            blockers["missing_group"] += 1
        else:
            group_resolved += 1
        classroom_key = _fold(row.classroom_raw) if row.classroom_raw else None
        classroom = classroom_map.get(classroom_key) if classroom_key not in ambiguous_classrooms else None
        if classroom_key in ambiguous_classrooms:
            blockers["ambiguous_classroom_reference"] += 1
        elif classroom is None:
            blockers["missing_classroom"] += 1
        else:
            classroom_resolved += 1

        teacher = None
        if row.activity == "practice" and row.semester:
            teacher_alias = aliases.teacher.get(row.teacher_key)
            canonical_teacher = teacher_alias[0] if teacher_alias else row.teacher_key
            ci = valid_join_ci.get((canonical_teacher, row.subject_key, row.semester))
            if ci in ambiguous_teacher_cis:
                teacher_ambiguous += 1
                blockers["ambiguous_teacher"] += 1
            elif ci:
                teacher = ci_map.get(ci)
        else:
            if row.teacher_key in ambiguous_teacher_names:
                teacher_ambiguous += 1
                blockers["ambiguous_teacher"] += 1
            else:
                teacher = name_map.get(row.teacher_key)
        if teacher is None:
            teacher_missing += 1
            blockers["missing_teacher"] += 1
        else:
            teacher_resolved += 1
        for slot in row.slots:
            if teacher is not None and not db.query(TeacherAvailability.id).filter(
                TeacherAvailability.teacher_ci == teacher.ci,
                TeacherAvailability.academic_period == academic_period,
                TeacherAvailability.weekday == slot.weekday,
                TeacherAvailability.start_time <= slot.start,
                TeacherAvailability.end_time >= slot.end,
                TeacherAvailability.active.is_(True),
            ).first():
                availability_missing += 1
            plan.append(PlannedBlock(
                activity=row.activity, teacher_key=row.teacher_key,
                teacher_ci=teacher.ci if teacher else None, subject_key=row.subject_key,
                subject_id=subject.id if subject else None, offering_id=offering.id if offering else None,
                semester=row.semester, group_key=row.group_key, group_id=group.id if group else None,
                classroom_key=classroom_key,
                classroom_id=classroom.id if classroom else None, slot=slot, source=row.source,
            ))
    blockers["missing_availability"] += availability_missing

    if not plan:
        blockers["no_planned_blocks"] += 1
    if not any(item.teacher_ci is not None for item in plan):
        blockers["no_planned_assignments"] += 1

    seen: Counter[tuple[Any, ...]] = Counter()
    for item in plan:
        seen[(item.activity, item.offering_id or (item.subject_key, item.semester),
              item.group_id or item.group_key, item.classroom_id or item.classroom_key,
              item.slot.weekday, item.slot.start, item.slot.end)] += 1
    duplicate_blocks = sum(count - 1 for count in seen.values() if count > 1)
    theory_duplicates = sum(
        count - 1 for key, count in seen.items() if key[0] == "theory" and count > 1
    )
    blockers["duplicate_block"] += duplicate_blocks
    overlaps = 0
    for index, left in enumerate(plan):
        for right in plan[index + 1:]:
            if left.slot.weekday != right.slot.weekday or left.slot.start >= right.slot.end or right.slot.start >= left.slot.end:
                continue
            if ((left.teacher_ci and left.teacher_ci == right.teacher_ci)
                    or (left.group_id and left.group_id == right.group_id)
                    or (left.classroom_id and left.classroom_id == right.classroom_id)):
                overlaps += 1
    blockers["schedule_overlap"] += overlaps
    blockers = Counter({key: value for key, value in blockers.items() if value})
    immutable_plan = tuple(plan)

    state_fingerprint = _state_fingerprint(
        db, academic_period, program_id, official, salary, excluded_draft_id,
    )
    plan_digest_payload = [{
        "activity": item.activity, "teacher": item.teacher_ci, "subject": item.subject_id,
        "offering": item.offering_id, "semester": item.semester, "group": item.group_id,
        "classroom": item.classroom_id, "weekday": item.slot.weekday,
        "start": item.slot.start.isoformat(), "end": item.slot.end.isoformat(),
        "source": [item.source.workbook_sha256, item.source.sheet, item.source.row],
    } for item in sorted(immutable_plan, key=lambda item: (
        item.activity, item.subject_key, item.group_key, item.slot.weekday,
        item.slot.start, item.teacher_key, item.source.row,
    ))]
    preview_digest = _digest({
        "policy": POLICY_VERSION, "sources": source_hashes, "period": academic_period,
        "program": program_id,
        "alias_sha256": aliases.sha256,
        "effective_date": effective_date.isoformat(), "plan": plan_digest_payload,
        "database_state": state_fingerprint,
    })
    theory_rows = [item for item in official if item.activity == "theory"]
    theory_stats = official_stats.get("theory", ParseStats())
    practice_stats = official_stats.get("practice", ParseStats())
    result = {
        "policy_version": POLICY_VERSION,
        "academic_period": academic_period,
        "effective_date": effective_date,
        "sources": {
            "official_sha256": source_hashes["official"], "official_sheet_count": official_sheet_count,
            "salary_sha256": source_hashes["salary"], "salary_sheet_count": salary_sheet_count,
            "alias_sha256": aliases.sha256, "alias_present": alias_content is not None,
        },
        "theory": {
            "row_count": theory_stats.physical, "parsed_row_count": theory_stats.parsed,
            "skipped_row_count": theory_stats.skipped,
            "error_row_count": theory_stats.errors,
            "block_count": sum(len(item.slots) for item in theory_rows),
            "assignment_count": sum(len(item.slots) for item in theory_rows),
            "duplicate_count": theory_duplicates,
            "error_count": activity_error_counts["theory"],
        },
        "practice": {
            "official_row_count": practice_stats.physical,
            "official_parsed_row_count": practice_stats.parsed,
            "official_skipped_row_count": practice_stats.skipped,
            "official_error_row_count": practice_stats.errors,
            "salary_row_count": salary_stats.physical,
            "salary_parsed_row_count": salary_stats.parsed,
            "salary_skipped_row_count": salary_stats.skipped,
            "salary_error_row_count": salary_stats.errors,
            "join_covered_count": join_covered, "join_missing_count": join_missing,
            "join_ambiguous_count": join_ambiguous,
            "join_missing_ci_count": missing_ci,
            "join_conflicting_ci_count": conflicting_ci,
            "orphan_salary_count": orphan_salary,
            "unresolved_teacher_alias_count": unresolved_teacher_alias,
            "unresolved_subject_alias_count": unresolved_subject_alias,
            "salary_payment_occurrence_count": salary_stats.parsed,
            "salary_payment_hours_total": sum(row.hours or 0 for row in salary),
            "salary_hours_noncomparable_count": salary_hours_noncomparable,
        },
        "resolution": {
            "teacher_resolved_count": teacher_resolved, "teacher_missing_count": teacher_missing,
            "teacher_ambiguous_count": teacher_ambiguous, "subject_resolved_count": subject_resolved,
            "offering_resolved_count": offering_resolved, "group_resolved_count": group_resolved,
            "classroom_resolved_count": classroom_resolved,
        },
        "planned": {
            "theory_block_count": sum(item.activity == "theory" for item in immutable_plan),
            "practice_block_count": sum(item.activity == "practice" for item in immutable_plan),
            "theory_assignment_count": sum(item.activity == "theory" and item.teacher_ci is not None for item in immutable_plan),
            "practice_assignment_count": sum(item.activity == "practice" and item.teacher_ci is not None for item in immutable_plan),
        },
        "warnings": [{"code": code, "count": count} for code, count in sorted(warnings.items()) if count],
        "blockers": [{"code": code, "count": count} for code, count in sorted(blockers.items())],
        "can_apply": not blockers,
        "database_state_fingerprint": state_fingerprint,
        "preview_digest": preview_digest,
    }
    if include_private:
        result["_private_plan"] = immutable_plan
        result["_program_id"] = program_id
    return result


def build_designation_bootstrap_preview(
    db: Session, *, official_content: bytes, salary_content: bytes,
    academic_period: str, effective_date: date, program_identity: str,
    alias_content: bytes | None = None,
) -> dict[str, Any]:
    """Build a preview without flushing or mutating caller-owned session state."""
    with db.no_autoflush:
        preview = _build_designation_bootstrap_preview(
            db, official_content=official_content, salary_content=salary_content,
            academic_period=academic_period, effective_date=effective_date,
            program_identity=program_identity, alias_content=alias_content,
            include_private=True,
        )
        # A post-apply preview of the exact same operation must reproduce the
        # original confirmation digest rather than treating its own draft as
        # external state drift. Other target drafts/publications remain bound.
        from app.models.designation_bootstrap import DesignationBootstrapReceipt

        receipt = db.query(DesignationBootstrapReceipt).filter_by(
            official_sha256=preview["sources"]["official_sha256"],
            salary_sha256=preview["sources"]["salary_sha256"],
            alias_sha256=preview["sources"]["alias_sha256"],
            program_id=preview["_program_id"],
            academic_period=academic_period,
            effective_date=effective_date,
        ).one_or_none()
        if receipt is None:
            return {
                key: value for key, value in preview.items()
                if key not in {"_private_plan", "_program_id"}
            }
        replay_preview = _build_designation_bootstrap_preview(
            db, official_content=official_content, salary_content=salary_content,
            academic_period=academic_period, effective_date=effective_date,
            program_identity=program_identity, alias_content=alias_content,
            excluded_draft_id=receipt.draft_id,
        )
        if replay_preview["preview_digest"] == receipt.preview_digest:
            return replay_preview
        return {
            key: value for key, value in preview.items()
            if key not in {"_private_plan", "_program_id"}
        }
