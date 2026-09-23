from __future__ import annotations

import hashlib
import json
from datetime import date, datetime
from pathlib import Path

import pytest
from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Font

from app.models.academic_management import (
    AcademicScheduleAssignment,
    AcademicScheduleBlock,
    AcademicSchedulePublication,
    AcademicSchedulePublishedAssignment,
    AcademicSchedulePublishedBlock,
    Classroom,
    HistoricalScheduleImport,
    TeacherAvailability,
)
from app.models.activity_log import ActivityLog
from app.models.practice_planilla import PracticePlanillaOutput
from app.models.user import User
from app.services import academic_management_service
from app.services.effective_schedule_service import effective_schedule_slots
from app.services.historical_assistential_import import (
    HistoricalImportError,
    ImportInputs,
    apply_import,
    build_preview,
    safe_preview,
)
from app.services.payroll_schedule_source_service import payroll_schedule_sources


def _write_workbook(path: Path, sheet_name: str, rows: range, red_rows: set[int]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    for row in rows:
        sheet.cell(row, 1, f"synthetic-row-{row}")
        sheet.cell(row, 2, 20)
        if row in red_rows:
            sheet.cell(row, 1).font = Font(color="FFFF0000")
    workbook.save(path)


def _fixture(tmp_path: Path) -> tuple[ImportInputs, str]:
    designation = tmp_path / "designation.xlsx"
    enrollment = tmp_path / "enrollment.xlsx"
    curriculum = tmp_path / "curriculum.txt"
    profiles_path = tmp_path / "profiles.json"
    decisions_path = tmp_path / "decisions.json"
    _write_workbook(designation, "DOCENTES ASISTENCIALES", range(2, 39), {6, 7, 8, 13})
    _write_workbook(enrollment, "GRUPOS", range(2, 40), set())
    curriculum.write_text(
        "MRF 0100 Anatomy I\nMRF 0200 Anatomy II\nCIR 0300 Surgery I\n"
        "CIR 0400 Surgery II\n{ malformed CR: 2 item",
        encoding="utf-8",
    )

    profiles: dict[str, dict[str, str]] = {}
    blocks = []
    replacement_rows = iter((35, 36, 37))
    for index in range(33):
        base_ci = f"SYN-{1000 + index}"
        profiles[base_ci] = {
            "full_name": f"Synthetic Teacher {index}",
            "email": f"teacher{index}@example.invalid",
        }
        code = tuple(("MRF 0100", "MRF 0200", "CIR 0300", "CIR 0400"))[index % 4]
        semester = {"MRF 0100": 1, "MRF 0200": 2, "CIR 0300": 3, "CIR 0400": 4}[code]
        assignments = [{
            "teacher_ci": base_ci,
            "effective_from": "2026-08-20",
            "effective_to": "2026-09-01" if index < 4 else None,
            "source_rows": [index + 2],
        }]
        if index < 3:
            replacement_ci = f"SYN-{2000 + index}"
            profiles[replacement_ci] = {
                "full_name": f"Replacement Teacher {index}",
                "phone": f"555000{index}",
            }
            replacement_row = next(replacement_rows)
            assignments.append({
                "teacher_ci": replacement_ci,
                "effective_from": "2026-09-02",
                "effective_to": None,
                "source_rows": [replacement_row, 38] if index == 2 else [replacement_row],
            })
        minute = (index * 5) % 60
        hour = 7 + (index * 5) // 60
        start = f"{hour:02d}:{minute:02d}"
        end_minutes = hour * 60 + minute + 5
        end = f"{end_minutes // 60:02d}:{end_minutes % 60:02d}"
        extra = index == 32
        blocks.append({
            "key": f"synthetic-{index}",
            "subject_code": code,
            "source_subject_label": f"Synthetic source subject {index}",
            "semester": semester,
            "group_code": "EXTRA" if extra else f"G-{index:02d}",
            "weekday": "monday",
            "start_time": start,
            "end_time": end,
            "source_rows": [index + 2],
            "source_real": 57,
            "classroom": {
                "code": "FIELD" if extra else f"COB-NEW-{index:02d}",
                "name": "Sin aula / campo asistencial" if extra else f"NUEVO / Aula {index}",
                "building": "" if extra else "NUEVO",
                "type": "other" if extra else "classroom",
                "capacity": None if extra else 20,
                "enrollment_row": None if extra else index + 2,
            },
            "assignments": assignments,
        })
    mariana_ci = "SYN-2000"
    marien_ci = "SYN-1001"
    nataly_ci = "SYN-1002"
    kenia_ci = "SYN-1003"
    pedro_ci = "SYN-1004"
    corrected_scopes = (
        (0, "MRF 0200", 2, "M-02"),
        (1, "MRF 0100", 1, "T-02"),
        (2, "MRF 0100", 1, "N-02"),
        (3, "MRF 0200", 2, "M-03"),
        (4, "MRF 0200", 2, "T-02"),
        (5, "CIR 0400", 4, "M-04"),
        (6, "CIR 0400", 4, "M-01"),
    )
    for index, code, semester, group in corrected_scopes:
        blocks[index].update(subject_code=code, semester=semester, group_code=group)
    blocks[0]["assignments"][1].update(teacher_ci=mariana_ci, source_rows=[35, 38])
    blocks[1]["assignments"][1]["teacher_ci"] = mariana_ci
    blocks[2]["assignments"][1]["teacher_ci"] = mariana_ci
    blocks[5]["assignments"][0]["teacher_ci"] = pedro_ci
    blocks[6]["assignments"][0]["teacher_ci"] = kenia_ci
    blocks[6].update(start_time="08:00 AM", end_time="11:20 AM")
    profiles_path.write_text(json.dumps(profiles), encoding="utf-8")
    decisions_path.write_text(json.dumps({
        "academic_period": "II/2026",
        "default_effective_date": "2026-08-20",
        "inclusive_until": "2026-09-01",
        "replacement_from": "2026-09-02",
        "physical_source_rows": list(range(2, 39)),
        "red_source_rows": [6, 7, 8, 13],
        "confirmed_corrections": {
            "mariana_ci": mariana_ci,
            "marien_ci": marien_ci,
            "nataly_ci": nataly_ci,
            "kenia_ci": kenia_ci,
            "pedro_ci": pedro_ci,
        },
        "blocks": blocks,
    }), encoding="utf-8")
    paths = {
        "designation": designation,
        "curriculum": curriculum,
        "enrollment": enrollment,
        "profiles": profiles_path,
        "decisions": decisions_path,
    }
    hashes = {name: hashlib.sha256(path.read_bytes()).hexdigest() for name, path in paths.items()}
    return ImportInputs(**paths, expected_hashes=hashes), next(iter(profiles))


def _admin(db_session) -> User:
    actor = User(
        ci="ADMIN-HIST-1", full_name="Historical Import Admin", password_hash="not-used",
        role="admin", is_active=True,
    )
    db_session.add(actor)
    db_session.flush()
    return actor


def _preview(db_session, inputs: ImportInputs) -> dict:
    return build_preview(
        db_session, inputs, actor_ci="ADMIN-HIST-1", policy="historical_availability_not_recorded",
        effective_date=date(2026, 8, 20), historical_availability_unrecorded=True,
    )


def test_preview_apply_two_revisions_without_availability_or_designations(db_session, tmp_path):
    _admin(db_session)
    inputs, _ci = _fixture(tmp_path)
    preview = _preview(db_session, inputs)
    public_preview = safe_preview(preview)
    serialized = json.dumps(public_preview)
    assert preview["can_apply"] is True
    assert preview["counts"] == {
        "physical_rows": 37, "logical_intervals": 36, "initial_blocks": 33,
        "current_blocks": 32, "replacements": 3,
    }
    assert "@example.invalid" not in serialized
    assert "555000" not in serialized
    assert "SYN-" not in serialized

    result = apply_import(
        db_session, inputs, actor_ci="ADMIN-HIST-1", policy="historical_availability_not_recorded",
        effective_date=date(2026, 8, 20), historical_availability_unrecorded=True,
        expected_digest=preview["digest"],
    )
    assert result["replayed"] is False
    assert db_session.query(TeacherAvailability).count() == 0
    assert db_session.query(AcademicSchedulePublication).count() == 2
    publications = db_session.query(AcademicSchedulePublication).order_by(
        AcademicSchedulePublication.effective_from
    ).all()
    assert [len(item.blocks) for item in publications] == [33, 32]
    assert db_session.query(AcademicScheduleAssignment).count() == 65
    assert db_session.query(AcademicSchedulePublishedAssignment).count() == 65
    field_location = db_session.query(Classroom).filter_by(code="FIELD").one()
    assert field_location.classroom_type == "other"
    assert field_location.capacity is None
    assert [
        sum(len(block.assignments) for block in publication.blocks)
        for publication in publications
    ] == [33, 32]
    assert db_session.query(HistoricalScheduleImport).count() == 1

    before = effective_schedule_slots(db_session, "II/2026", date(2026, 9, 1), include_practice=True)
    after = effective_schedule_slots(db_session, "II/2026", date(2026, 9, 2), include_practice=True)
    assert len(before) == 33
    assert len(after) == 32
    assert all(item.source_type == "published" and item.activity_type == "practice" for item in after)
    payroll = payroll_schedule_sources(
        db_session, academic_period="II/2026", period_start=date(2026, 9, 1),
        period_end=date(2026, 9, 7), activity_kind="practice",
    )
    assert payroll
    assert all(item.source_kind == "published" for item in payroll)

    event = db_session.query(ActivityLog).filter_by(action="historical_assistential_import").one()
    assert event.user_ci is None and event.user_name is None
    assert not any(field in json.dumps(event.details).lower() for field in ("email", "phone", "account_number"))


def test_normal_planner_still_requires_availability(db_session, tmp_path):
    _admin(db_session)
    inputs, ci = _fixture(tmp_path)
    preview = _preview(db_session, inputs)
    apply_import(
        db_session, inputs, actor_ci="ADMIN-HIST-1", policy="historical_availability_not_recorded",
        effective_date=date(2026, 8, 20), historical_availability_unrecorded=True,
        expected_digest=preview["digest"],
    )
    block = db_session.query(AcademicSchedulePublishedBlock).first()
    draft_block = db_session.get(AcademicScheduleBlock, block.source_block_id)
    draft = draft_block.draft
    with pytest.raises(HTTPException, match="disponibilidad"):
        academic_management_service._ensure_teacher_availability(db_session, draft, draft_block, ci)


def test_digest_binding_idempotent_replay_and_authority_flags(db_session, tmp_path):
    _admin(db_session)
    inputs, _ci = _fixture(tmp_path)
    with pytest.raises(HistoricalImportError, match="historical-availability"):
        build_preview(
            db_session, inputs, actor_ci="ADMIN-HIST-1", policy="historical_availability_not_recorded",
            effective_date=date(2026, 8, 20), historical_availability_unrecorded=False,
        )
    with pytest.raises(HistoricalImportError, match="--policy"):
        build_preview(
            db_session, inputs, actor_ci="ADMIN-HIST-1",
            policy="contact person@example.test at 70000000",
            effective_date=date(2026, 8, 20), historical_availability_unrecorded=True,
        )
    preview = _preview(db_session, inputs)
    with pytest.raises(HistoricalImportError, match="changed after preview"):
        apply_import(
            db_session, inputs, actor_ci="ADMIN-HIST-1", policy="historical_availability_not_recorded",
            effective_date=date(2026, 8, 20), historical_availability_unrecorded=True,
            expected_digest="0" * 64,
        )
    first = apply_import(
        db_session, inputs, actor_ci="ADMIN-HIST-1", policy="historical_availability_not_recorded",
        effective_date=date(2026, 8, 20), historical_availability_unrecorded=True,
        expected_digest=preview["digest"],
    )
    replay = apply_import(
        db_session, inputs, actor_ci="ADMIN-HIST-1", policy="historical_availability_not_recorded",
        effective_date=date(2026, 8, 20), historical_availability_unrecorded=True,
        expected_digest=preview["digest"],
    )
    assert replay["replayed"] is True
    assert replay["initial_publication_id"] == first["initial_publication_id"]
    assert db_session.query(AcademicSchedulePublication).count() == 2
    with pytest.raises(HistoricalImportError, match="Replay binding"):
        apply_import(
            db_session, inputs, actor_ci="ADMIN-HIST-1",
            policy="historical_availability_not_recorded",
            effective_date=date(2026, 8, 20), historical_availability_unrecorded=True,
            expected_digest="0" * 64,
        )


def test_apply_rejects_same_count_pre_state_mutation(db_session, tmp_path):
    actor = _admin(db_session)
    inputs, _ci = _fixture(tmp_path)
    preview = _preview(db_session, inputs)
    actor.full_name = "Changed Administrator Name"
    db_session.flush()
    with pytest.raises(HistoricalImportError, match="changed after preview"):
        apply_import(
            db_session, inputs, actor_ci="ADMIN-HIST-1",
            policy="historical_availability_not_recorded",
            effective_date=date(2026, 8, 20), historical_availability_unrecorded=True,
            expected_digest=preview["digest"],
        )


def test_replay_rejects_same_count_live_state_mutation(db_session, tmp_path):
    _admin(db_session)
    inputs, _ci = _fixture(tmp_path)
    preview = _preview(db_session, inputs)
    apply_import(
        db_session, inputs, actor_ci="ADMIN-HIST-1",
        policy="historical_availability_not_recorded",
        effective_date=date(2026, 8, 20), historical_availability_unrecorded=True,
        expected_digest=preview["digest"],
    )
    publication = db_session.query(AcademicSchedulePublication).order_by(
        AcademicSchedulePublication.effective_from
    ).first()
    publication.content_digest = "f" * 64
    db_session.flush()
    with pytest.raises(HistoricalImportError, match="Replay binding"):
        apply_import(
            db_session, inputs, actor_ci="ADMIN-HIST-1",
            policy="historical_availability_not_recorded",
            effective_date=date(2026, 8, 20), historical_availability_unrecorded=True,
            expected_digest=preview["digest"],
        )


def test_retroactive_evidence_blocks_without_partial_writes(db_session, tmp_path):
    _admin(db_session)
    inputs, _ci = _fixture(tmp_path)
    db_session.add(PracticePlanillaOutput(
        month=8, year=2026, generated_at=datetime(2026, 8, 31), file_path=None,
        total_teachers=1, total_hours=1, total_payment=1, status="generated",
    ))
    db_session.flush()
    preview = _preview(db_session, inputs)
    assert preview["can_apply"] is False
    assert preview["operational_blockers"]["practice_payroll"] == 1
    with pytest.raises(HistoricalImportError, match="blocked"):
        apply_import(
            db_session, inputs, actor_ci="ADMIN-HIST-1", policy="historical_availability_not_recorded",
            effective_date=date(2026, 8, 20), historical_availability_unrecorded=True,
            expected_digest=preview["digest"],
        )
    assert db_session.query(AcademicSchedulePublication).count() == 0
