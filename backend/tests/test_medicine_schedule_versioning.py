from datetime import date, time
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import Workbook
from sqlalchemy.exc import IntegrityError

from app.database import Base
from app.main import app
from app.models.activity_log import ActivityLog
from app.models.attendance import AttendanceRecord
from app.models.designation import Designation
from app.models.medicine_schedule import (MedicineCorrection, MedicineImportIssue, MedicineMeeting, MedicineOffering,
                                          MedicineScheduleVersion, MedicineVersionEvent)
from app.models.planilla import PlanillaOutput
from app.models.teacher import Teacher
from app.models.user import User
from app.routers.medicine_schedules import router
from app.schemas.medicine_schedule import MedicineIssuePreview
from app.services import app_settings_service
from app.services.medicine_schedule_parser import parse_medicine_workbook
from app.services.medicine_schedule_version_service import (MedicineVersionConflict, MedicineVersionError,
                                                            medicine_schedule_version_service as service)


MEDICINE_TABLES = {
    "medicine_schedule_versions", "medicine_offerings", "medicine_meetings",
    "medicine_import_issues", "medicine_corrections", "medicine_version_events",
    "medicine_simulations",
}

def _actor(db_session, suffix="ADMIN", role="admin", active=True) -> User:
    actor = User(ci=f"MED-VERSION-{suffix}", full_name="Medicine Actor",
                 password_hash="unused", role=role, is_active=active)
    db_session.add(actor)
    db_session.flush()
    return actor

def _preview(errors=0, warnings=0):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "1RO"
    sheet.append(["PRIMER SEMESTRE | GRUPO M1 | MAÑANA"])
    sheet.append([])
    sheet.append(["HORARIO", "LUNES", "MARTES", "MIERCOLES"])
    sheet.append(["08:00 - 09:30", "Raw Subject"])
    sheet.append([None, "Raw Teacher"])
    sheet.append([None, "TEORÍA"])
    stream = BytesIO()
    workbook.save(stream)
    preview = parse_medicine_workbook(stream.getvalue())
    preview.issues.extend(
        MedicineIssuePreview(severity="error", code="canonical_error",
                             message="Explicit correction required", location={"sheet": "1RO", "cell": "B4"})
        for index in range(errors)
    )
    preview.issues.extend(
        MedicineIssuePreview(severity="warning", code="review_warning",
                             message="Explicit acceptance required", location={"cell": f"C{index}"})
        for index in range(warnings)
    )
    return preview

def _persist(db_session, actor, *, errors=0, warnings=0, period="I/2099"):
    return service.persist_preview(
        db_session, _preview(errors, warnings), period, actor.id,
        source_file_path="/isolated/medicine.xlsx", description="Version fixture",
    )

def test_medicine_models_are_registered_and_isolated():
    assert MEDICINE_TABLES <= set(Base.metadata.tables)
    foreign_tables = {
        foreign_key.column.table.name
        for name in MEDICINE_TABLES
        for foreign_key in Base.metadata.tables[name].foreign_keys
    }
    assert foreign_tables <= MEDICINE_TABLES | {"users"}

def test_feature_is_default_disabled_and_registered(client, db_session):
    app_settings_service.invalidate_cache()
    assert app_settings_service.get_medicine_schedule_assistant_enabled(db_session) is False
    assert router.prefix == "/api/medicine-schedules"
    assert any(route.path == "/api/medicine-schedules/status" for route in app.routes)
    assert client.get("/api/medicine-schedules/status").status_code == 404
    settings = client.get("/api/admin/settings")
    assert settings.status_code == 200
    assert settings.json()["medicine_schedule_assistant_enabled"] is False
    assert client.put("/api/admin/settings", json={"medicine_schedule_assistant_enabled": True}).json()["medicine_schedule_assistant_enabled"] is True
    assert client.get("/api/medicine-schedules/status").json() == {"enabled": True}

def test_persist_preview_preserves_lineage_and_legacy_rows(db_session):
    actor = _actor(db_session)
    teacher = Teacher(ci="LEGACY-MED", full_name="Legacy Teacher")
    db_session.add(teacher)
    db_session.flush()
    designation = Designation(teacher_ci=teacher.ci, subject="Legacy Subject", semester="1",
                              group_code="M1", academic_period="I/2099", schedule_json=[])
    payroll = PlanillaOutput(month=12, year=2099, total_teachers=1, total_hours=1,
                             total_payment=Decimal("70"), status="generated")
    db_session.add_all([designation, payroll])
    db_session.flush()
    attendance = AttendanceRecord(
        teacher_ci=teacher.ci, designation_id=designation.id, date=date(2099, 12, 1),
        scheduled_start=time(8), scheduled_end=time(9), status="ATTENDED",
        academic_hours=1, late_minutes=0, month=12, year=2099,
    )
    db_session.add(attendance)
    db_session.flush()
    legacy = (designation.subject, attendance.status, payroll.total_payment, actor.full_name)
    version = _persist(db_session, actor, warnings=1)
    offering = db_session.query(MedicineOffering).filter_by(version_id=version.id).one()
    meeting = db_session.query(MedicineMeeting).filter_by(offering_id=offering.id).one()
    issue = db_session.query(MedicineImportIssue).filter_by(version_id=version.id).one()
    assert (version.academic_period, version.description, version.status, version.is_active) == (
        "I/2099", "Version fixture", "preview", False)
    assert version.workbook_sha256 and version.parser_schema_version == "medicine-v1"
    assert version.uploaded_by == actor.id and version.created_at is not None
    assert (offering.subject_raw, offering.raw_payload["subject"], offering.raw_payload["subject_cell"]) == ("Raw Subject", "Raw Subject", "B4")
    assert meeting.teacher_raw == "Raw Teacher" and meeting.source_cell == "B4"
    assert issue.severity == "warning" and issue.state == "open"
    assert db_session.query(MedicineVersionEvent).filter_by(version_id=version.id, event_type="upload").count() == 1
    assert db_session.query(ActivityLog).filter_by(category="medicine_schedule", action="upload").count() == 1
    assert legacy == (designation.subject, attendance.status, payroll.total_payment, actor.full_name)

def test_parser_to_persist_correct_accept_activate_and_restore(db_session, monkeypatch):
    actor = _actor(db_session)
    blocked = _persist(db_session, actor, errors=10, period="BLOCKED/2099")
    assert db_session.query(MedicineImportIssue).filter_by(version_id=blocked.id, severity="error").count() == 10
    with pytest.raises(MedicineVersionError, match="10 unresolved"):
        service.activate(db_session, blocked.id, actor.id)
    for suffix, role, active in (("DOCENTE", "docente", True), ("INACTIVE", "admin", False)):
        invalid_actor = _actor(db_session, suffix, role, active)
        with pytest.raises(MedicineVersionError, match="Active administrator"):
            service.activate(db_session, blocked.id, invalid_actor.id)
    first = _persist(db_session, actor, errors=2, warnings=1)
    offering = db_session.query(MedicineOffering).filter_by(version_id=first.id).one()
    meeting = db_session.query(MedicineMeeting).filter_by(offering_id=offering.id).one()
    errors = db_session.query(MedicineImportIssue).filter_by(version_id=first.id, severity="error").all()
    warning = db_session.query(MedicineImportIssue).filter_by(version_id=first.id, severity="warning").one()
    with pytest.raises(MedicineVersionError, match="3 unresolved"):
        service.activate(db_session, first.id, actor.id)
    with pytest.raises(MedicineVersionError, match="Warning cannot"):
        service.accept_warning(db_session, first.id, errors[0].id, actor.id)
    with pytest.raises(MedicineVersionError, match="Unsupported"):
        service.correct_field(db_session, first.id, "offering", offering.id,
                              "subject_raw", "Forbidden", actor.id)
    subject_payload = offering.raw_payload
    before_rejection = (offering.subject_key, db_session.query(MedicineCorrection).count(), db_session.query(MedicineVersionEvent).filter_by(event_type="correction").count(), db_session.query(ActivityLog).filter_by(action="correction").count())
    for location, payload in (({"sheet": "1RO", "cell": "A4"}, subject_payload), ([], subject_payload), ({"sheet": "1RO", "cell": "B4"}, [])):
        errors[0].code, errors[0].location, offering.raw_payload = "invalid_time", location, payload
        db_session.flush()
        with pytest.raises(MedicineVersionError, match="causally bound"):
            service.correct_field(db_session, first.id, "offering", offering.id, "subject_key", "BYPASS", actor.id, [errors[0].id])
        assert before_rejection == (offering.subject_key, db_session.query(MedicineCorrection).count(), db_session.query(MedicineVersionEvent).filter_by(event_type="correction").count(), db_session.query(ActivityLog).filter_by(action="correction").count())
        assert errors[0].state == "open" and not db_session.new and not db_session.dirty
    offering.raw_payload, errors[0].location = subject_payload, {"sheet": "1RO", "cell": "B4"}
    db_session.flush()
    subject_correction = service.correct_field(db_session, first.id, "offering", offering.id, "subject_key",
                                               "CANONICAL SUBJECT", actor.id, [errors[0].id])
    teacher_before = (meeting.teacher_key, db_session.query(MedicineCorrection).count(), db_session.query(MedicineVersionEvent).filter_by(event_type="correction").count(), db_session.query(ActivityLog).filter_by(action="correction").count())
    with pytest.raises(MedicineVersionError, match="causally bound"):
        service.correct_field(db_session, first.id, "meeting", meeting.id, "teacher_key", "BYPASS", actor.id, [errors[1].id])
    assert teacher_before == (meeting.teacher_key, db_session.query(MedicineCorrection).count(), db_session.query(MedicineVersionEvent).filter_by(event_type="correction").count(), db_session.query(ActivityLog).filter_by(action="correction").count())
    assert errors[1].state == "open" and not db_session.new and not db_session.dirty
    errors[1].location = {"sheet": "1RO", "cell": "B5"}
    db_session.flush()
    service.correct_field(db_session, first.id, "meeting", meeting.id, "teacher_key", "CANONICAL TEACHER", actor.id, [errors[1].id])
    accepted = service.accept_warning(db_session, first.id, warning.id, actor.id)
    service.activate(db_session, first.id, actor.id)
    assert (offering.subject_raw, offering.subject_key) == ("Raw Subject", "CANONICAL SUBJECT")
    assert (meeting.teacher_raw, meeting.teacher_key) == ("Raw Teacher", "CANONICAL TEACHER")
    assert subject_correction.before_value == {"value": "RAW SUBJECT"}
    assert subject_correction.after_value == {"value": "CANONICAL SUBJECT"} and subject_correction.actor_id == actor.id
    assert accepted.accepted_by == actor.id and accepted.accepted_at is not None
    assert first.is_active and first.locked_at is not None and first.status == "active"
    assert db_session.query(MedicineImportIssue).filter_by(version_id=first.id, state="open").count() == 0
    with pytest.raises(MedicineVersionError, match="Only validated"):
        service.restore(db_session, first.id, actor.id)
    with pytest.raises(MedicineVersionError, match="immutable"):
        service.correct_field(db_session, first.id, "offering", offering.id,
                              "subject_key", "LATE CHANGE", actor.id)
    second = _persist(db_session, actor, period="II/2099")
    service.activate(db_session, second.id, actor.id)
    assert second.is_active and not first.is_active and first.status == "inactive"
    service.restore(db_session, first.id, actor.id)
    assert first.is_active and not second.is_active
    assert db_session.query(MedicineScheduleVersion).filter_by(is_active=True).count() == 1
    assert {event.event_type for event in db_session.query(MedicineVersionEvent).all()} >= {
        "upload", "correction", "warning_acceptance", "activation", "restore"}
    assert {entry.action for entry in db_session.query(ActivityLog).all()} >= {
        "upload", "correction", "warning_acceptance", "activation", "restore"}
    assert db_session.query(MedicineCorrection).filter_by(version_id=first.id).count() == 2
    third = _persist(db_session, actor, period="III/2099")
    def conflict():
        raise IntegrityError("active index", {}, Exception("duplicate"))
    monkeypatch.setattr(db_session, "begin_nested", conflict)
    with pytest.raises(MedicineVersionConflict, match="became active"):
        service.activate(db_session, third.id, actor.id)
    assert db_session.get(User, actor.id) is actor and third.status == "preview"
