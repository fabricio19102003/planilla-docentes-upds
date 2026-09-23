"""Digest-bound importer for the one II/2026 assistential schedule recovery.

The public planner remains strict. This module writes the same immutable catalog,
draft, and publication rows through a deliberately separate, CLI-only boundary.
"""
from __future__ import annotations

import hashlib
import hmac
import json
import re
import unicodedata
from dataclasses import dataclass
from datetime import date, time
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.models.academic_management import (
    AcademicGroup,
    AcademicProgram,
    AcademicScheduleAssignment,
    AcademicScheduleBlock,
    AcademicScheduleDraft,
    AcademicSchedulePublication,
    AcademicSchedulePublishedAssignment,
    AcademicSchedulePublishedBlock,
    AcademicSubject,
    Classroom,
    HistoricalScheduleImport,
    SubjectOffering,
    TeacherAvailability,
)
from app.models.activity_log import ActivityLog
from app.models.attendance import AttendanceRecord
from app.models.billing_publication import BillingPublication
from app.models.contract import ContractDocument, ContractLine
from app.models.designation import Designation
from app.models.practice_attendance import PracticeAttendanceLog
from app.models.practice_planilla import PracticePlanillaOutput
from app.models.teacher import Teacher
from app.models.user import User
from app.services.teacher_profile_import_service import PROFILE_FIELDS, _canonical_profile, _same


PERIOD = "II/2026"
PROGRAM_CODE = "MED 510-01"
PROGRAM_NAME = "Medicina"
INITIAL_DATE = date(2026, 8, 20)
CUTOVER_DATE = date(2026, 9, 2)
SUBJECTS = {
    "MRF 0100": ("Anatomy I", 1),
    "MRF 0200": ("Anatomy II", 2),
    "CIR 0300": ("Surgery I", 3),
    "CIR 0400": ("Surgery II", 4),
}
PROFILE_ALLOWED = {"full_name", *PROFILE_FIELDS}
SHA256_RE = re.compile(r"^[0-9a-f]{64}$")
WEEKDAYS = {"monday", "tuesday", "wednesday", "thursday", "friday", "saturday"}
POLICY = "historical_availability_not_recorded"
POLICY_DESCRIPTION = (
    "Historical II/2026 availability was not recorded; source-backed schedule recovery authorized."
)
IMPORT_LOCK_KEY = hashlib.sha256(f"historical-assistential-import:{PERIOD}".encode()).hexdigest()


class HistoricalImportError(ValueError):
    pass


@dataclass(frozen=True)
class ImportInputs:
    designation: Path
    curriculum: Path
    enrollment: Path
    profiles: Path
    decisions: Path
    expected_hashes: dict[str, str]


def _json_bytes(value: Any) -> bytes:
    return json.dumps(
        value, ensure_ascii=False, sort_keys=True, separators=(",", ":")
    ).encode("utf-8")


def _sha(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def _normalized(value: object) -> str:
    raw = unicodedata.normalize("NFKD", str(value or "").strip().casefold())
    return " ".join("".join(c for c in raw if not unicodedata.combining(c)).split())


def mask_ci(ci: str) -> str:
    compact = "".join(character for character in ci if character.isalnum())
    return f"***{compact[-3:]}" if compact else "***"


def _parse_date(value: object, label: str) -> date:
    try:
        return date.fromisoformat(str(value))
    except ValueError as exc:
        raise HistoricalImportError(f"{label} must be an ISO date") from exc


def _parse_time(value: object, label: str) -> time:
    raw = " ".join(str(value).strip().upper().split())
    formats = (
        (r"^(\d{1,2}):(\d{2})\s*(AM|PM)$", True),
        (r"^(\d{1,2}):(\d{2})$", False),
    )
    for pattern, twelve_hour in formats:
        match = re.fullmatch(pattern, raw)
        if not match:
            continue
        hour, minute = int(match.group(1)), int(match.group(2))
        if minute > 59 or (twelve_hour and not 1 <= hour <= 12) or (not twelve_hour and hour > 23):
            break
        if twelve_hour:
            marker = match.group(3)
            hour = hour % 12 + (12 if marker == "PM" else 0)
        result = time(hour, minute)
        if result > time(23, 5):
            raise HistoricalImportError(f"{label} exceeds the historical 23:05 limit")
        return result
    raise HistoricalImportError(f"{label} is ambiguous or impossible; add an explicit correction")


def _is_red(cell: Any) -> bool:
    colors = (cell.font.color, cell.fill.fgColor, cell.fill.start_color)
    for color in colors:
        if color is None:
            continue
        rgb = str(getattr(color, "rgb", "") or "").upper()
        indexed = getattr(color, "indexed", None)
        if rgb.endswith("FF0000") or indexed == 10:
            return True
    return False


def _sheet_evidence(raw: bytes, sheet_name: str) -> dict[int, dict[str, Any]]:
    workbook = load_workbook(BytesIO(raw), read_only=False, data_only=False)
    if sheet_name not in workbook.sheetnames:
        raise HistoricalImportError(f"Workbook is missing sheet {sheet_name!r}")
    sheet = workbook[sheet_name]
    evidence: dict[int, dict[str, Any]] = {}
    for row in sheet.iter_rows():
        populated = [cell for cell in row if cell.value not in (None, "")]
        if populated:
            evidence[row[0].row] = {
                "cells": [cell.coordinate for cell in populated],
                "red": any(_is_red(cell) for cell in populated),
                "normalized_values": [_normalized(cell.value) for cell in populated],
            }
    return evidence


def _read_inputs(inputs: ImportInputs) -> tuple[dict[str, bytes], dict[str, str]]:
    paths = {
        "designation": inputs.designation,
        "curriculum": inputs.curriculum,
        "enrollment": inputs.enrollment,
        "profiles": inputs.profiles,
        "decisions": inputs.decisions,
    }
    raw = {name: path.read_bytes() for name, path in paths.items()}
    hashes = {name: _sha(payload) for name, payload in raw.items()}
    expected = {name: value.lower() for name, value in inputs.expected_hashes.items()}
    if set(expected) != set(paths) or any(not SHA256_RE.fullmatch(value) for value in expected.values()):
        raise HistoricalImportError("All five expected SHA-256 values are required")
    changed = sorted(name for name in paths if not hmac.compare_digest(hashes[name], expected[name]))
    if changed:
        raise HistoricalImportError("Source SHA-256 mismatch: " + ", ".join(changed))
    return raw, hashes


def _load_json(raw: bytes, label: str) -> dict[str, Any]:
    try:
        value = json.loads(raw.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise HistoricalImportError(f"{label} must be valid UTF-8 JSON") from exc
    if not isinstance(value, dict):
        raise HistoricalImportError(f"{label} must be a JSON object")
    return value


def _validate_curriculum(raw: bytes) -> None:
    text_value = raw.decode("utf-8", errors="replace")
    missing = [code for code in SUBJECTS if code not in text_value]
    if missing:
        raise HistoricalImportError("Curriculum is missing exact subject codes: " + ", ".join(missing))


def _profile_plan(db: Session, payload: dict[str, Any]) -> tuple[list[dict[str, Any]], list[str]]:
    changes: list[dict[str, Any]] = []
    conflicts: list[str] = []
    for row_number, (ci, raw_profile) in enumerate(sorted(payload.items()), start=1):
        if not isinstance(ci, str) or not ci.strip() or not isinstance(raw_profile, dict):
            raise HistoricalImportError("Profiles must be keyed by nonempty authoritative CI")
        unsupported = sorted(set(raw_profile) - PROFILE_ALLOWED)
        if unsupported:
            raise HistoricalImportError(
                f"Profile {mask_ci(ci)} contains unsupported fields: {', '.join(unsupported)}"
            )
        full_name = " ".join(str(raw_profile.get("full_name") or "").split())
        if len(full_name.split()) < 2:
            raise HistoricalImportError(
                f"Profile {mask_ci(ci)} requires an authoritative multi-token full_name"
            )
        validation_errors: list[str] = []
        canonical_fields = _canonical_profile(
            {name: raw_profile.get(name) for name in PROFILE_FIELDS},
            row_number,
            validation_errors,
        )
        if validation_errors:
            raise HistoricalImportError(
                f"Profile {mask_ci(ci)} contains invalid supported field values"
            )
        canonical_profile = {"full_name": full_name, **canonical_fields}
        teacher = db.get(Teacher, ci)
        field_actions: dict[str, str] = {}
        if teacher is None:
            field_actions["full_name"] = "create"
            for field_name in PROFILE_FIELDS:
                field_actions[field_name] = "create" if canonical_profile.get(field_name) not in (None, "") else "noop"
            action = "create"
        else:
            action = "update"
            field_actions["full_name"] = "noop" if _normalized(teacher.full_name) == _normalized(full_name) else "conflict"
            for field_name in PROFILE_FIELDS:
                incoming = canonical_profile.get(field_name)
                current = getattr(teacher, field_name)
                if incoming in (None, ""):
                    field_actions[field_name] = "noop"
                elif current in (None, ""):
                    field_actions[field_name] = "fill"
                elif _same(field_name, str(current), str(incoming)):
                    field_actions[field_name] = "noop"
                else:
                    field_actions[field_name] = "conflict"
        if "conflict" in field_actions.values():
            conflicts.append(mask_ci(ci))
        changes.append({
            "ci_masked": mask_ci(ci), "ci": ci, "action": action,
            "fields": field_actions, "profile": canonical_profile,
        })
    return changes, conflicts


def _intervals_overlap(left: dict[str, Any], right: dict[str, Any]) -> bool:
    left_end = left["effective_to"] or date.max
    right_end = right["effective_to"] or date.max
    return left["effective_from"] <= right_end and right["effective_from"] <= left_end


def _parse_plan(
    decisions: dict[str, Any], designation_evidence: dict[int, dict[str, Any]],
    enrollment_evidence: dict[int, dict[str, Any]], profiles: dict[str, Any],
) -> tuple[list[dict[str, Any]], dict[str, int]]:
    if decisions.get("academic_period") != PERIOD:
        raise HistoricalImportError(f"decisions.academic_period must be exactly {PERIOD}")
    if _parse_date(decisions.get("default_effective_date"), "default_effective_date") != INITIAL_DATE:
        raise HistoricalImportError("default_effective_date must be 2026-08-20")
    physical_rows = decisions.get("physical_source_rows")
    if not isinstance(physical_rows, list) or len(set(physical_rows)) != 37:
        raise HistoricalImportError("physical_source_rows must identify exactly 37 distinct workbook rows")
    missing_rows = sorted(set(physical_rows) - set(designation_evidence))
    if missing_rows:
        raise HistoricalImportError(f"Designation source rows are missing: {missing_rows}")
    red_rows = decisions.get("red_source_rows")
    if not isinstance(red_rows, list) or any(not designation_evidence.get(row, {}).get("red") for row in red_rows):
        raise HistoricalImportError("red_source_rows must be backed by workbook red style evidence")
    if decisions.get("inclusive_until") != "2026-09-01" or decisions.get("replacement_from") != "2026-09-02":
        raise HistoricalImportError("The inclusive cutoff and replacement date decisions are required")

    raw_blocks = decisions.get("blocks")
    if not isinstance(raw_blocks, list):
        raise HistoricalImportError("decisions.blocks must be a list")
    blocks: list[dict[str, Any]] = []
    assignment_count = replacement_count = 0
    source_rows_seen: set[int] = set()
    for index, raw_block in enumerate(raw_blocks, start=1):
        if not isinstance(raw_block, dict):
            raise HistoricalImportError(f"Block {index} must be an object")
        code = str(raw_block.get("subject_code") or "")
        if code not in SUBJECTS:
            raise HistoricalImportError(f"Block {index} has an unsupported subject code")
        canonical_name, canonical_semester = SUBJECTS[code]
        if int(raw_block.get("semester", 0)) != canonical_semester:
            raise HistoricalImportError(f"Block {index} has the wrong canonical semester")
        weekday = str(raw_block.get("weekday") or "").lower()
        if weekday not in WEEKDAYS:
            raise HistoricalImportError(f"Block {index} has an invalid weekday")
        start = _parse_time(raw_block.get("start_time"), f"block {index} start_time")
        end = _parse_time(raw_block.get("end_time"), f"block {index} end_time")
        if start >= end:
            raise HistoricalImportError(f"Block {index} must have start_time < end_time")
        source_rows = raw_block.get("source_rows")
        if not isinstance(source_rows, list) or not source_rows or any(row not in designation_evidence for row in source_rows):
            raise HistoricalImportError(f"Block {index} requires valid designation source_rows")
        source_rows_seen.update(source_rows)
        classroom = raw_block.get("classroom")
        if not isinstance(classroom, dict):
            raise HistoricalImportError(f"Block {index} requires classroom evidence")
        group_code = str(raw_block.get("group_code") or "").strip().upper()
        if group_code == "EXTRA":
            if classroom.get("type") != "other" or classroom.get("name") != "Sin aula / campo asistencial":
                raise HistoricalImportError("EXTRA must use the special field-practice location")
            capacity = None
        else:
            enrollment_row = classroom.get("enrollment_row")
            if enrollment_row not in enrollment_evidence:
                raise HistoricalImportError(f"Block {index} lacks enrollment classroom lineage")
            capacity = classroom.get("capacity")
            if isinstance(capacity, bool) or not isinstance(capacity, int) or capacity <= 0:
                raise HistoricalImportError(f"Block {index} requires positive physical classroom capacity")
            if str(capacity) not in enrollment_evidence[enrollment_row]["normalized_values"]:
                raise HistoricalImportError(f"Block {index} classroom capacity is not present in its source row")
        assignments: list[dict[str, Any]] = []
        for raw_assignment in raw_block.get("assignments") or []:
            ci = str(raw_assignment.get("teacher_ci") or "")
            if ci not in profiles:
                raise HistoricalImportError(f"Block {index} references a CI absent from private profiles")
            effective_from = _parse_date(raw_assignment.get("effective_from"), "assignment effective_from")
            effective_to = (
                _parse_date(raw_assignment["effective_to"], "assignment effective_to")
                if raw_assignment.get("effective_to") else None
            )
            if effective_to and effective_to < effective_from:
                raise HistoricalImportError("Assignment effective_to precedes effective_from")
            assignment_rows = raw_assignment.get("source_rows") or source_rows
            if any(row not in designation_evidence for row in assignment_rows):
                raise HistoricalImportError("Assignment has invalid source lineage")
            source_rows_seen.update(assignment_rows)
            assignments.append({
                "teacher_ci": ci, "effective_from": effective_from,
                "effective_to": effective_to, "source_rows": sorted(set(assignment_rows)),
            })
            assignment_count += 1
            replacement_count += effective_from == CUTOVER_DATE
        assignments.sort(key=lambda item: (item["effective_from"], item["teacher_ci"]))
        for left, right in zip(assignments, assignments[1:]):
            if _intervals_overlap(left, right):
                raise HistoricalImportError(f"Block {index} has overlapping assignment intervals")
        blocks.append({
            "key": str(raw_block.get("key") or f"block-{index}"),
            "subject_code": code, "subject_name": canonical_name,
            "source_subject_label": str(raw_block.get("source_subject_label") or ""),
            "semester": canonical_semester, "group_code": group_code,
            "weekday": weekday, "start_time": start, "end_time": end,
            "classroom": {
                "code": str(classroom.get("code") or "").strip().upper(),
                "name": str(classroom.get("name") or "").strip(),
                "campus": "Cobija", "capacity": capacity,
                "type": str(classroom.get("type") or "classroom"),
                "building": str(classroom.get("building") or "").strip().upper(),
                "enrollment_row": classroom.get("enrollment_row"),
            },
            "source_rows": sorted(set(source_rows)), "assignments": assignments,
            "official_workload": 60, "source_real": raw_block.get("source_real"),
        })
    if source_rows_seen != set(physical_rows):
        raise HistoricalImportError("Every physical designation row must have exact lineage")
    if len(blocks) != 33 or assignment_count != 36 or replacement_count != 3:
        raise HistoricalImportError("Expected exactly 33 blocks, 36 logical intervals, and 3 replacements")

    # Cross-block group/classroom and teacher conflicts retain normal integrity rules.
    for position, left in enumerate(blocks):
        for right in blocks[position + 1:]:
            if left["weekday"] != right["weekday"] or not (
                left["start_time"] < right["end_time"] and left["end_time"] > right["start_time"]
            ):
                continue
            if left["group_code"] == right["group_code"] and left["semester"] == right["semester"]:
                raise HistoricalImportError("A group has overlapping historical blocks")
            if left["classroom"]["code"] == right["classroom"]["code"] and left["classroom"]["type"] != "other":
                raise HistoricalImportError("A classroom has overlapping historical blocks")
            for left_assignment in left["assignments"]:
                for right_assignment in right["assignments"]:
                    if left_assignment["teacher_ci"] == right_assignment["teacher_ci"] and _intervals_overlap(left_assignment, right_assignment):
                        raise HistoricalImportError("A teacher has overlapping historical blocks")

    current = [
        block for block in blocks
        if any(
            item["effective_from"] <= CUTOVER_DATE
            and (item["effective_to"] is None or item["effective_to"] >= CUTOVER_DATE)
            for item in block["assignments"]
        )
    ]
    if len(current) != 32:
        raise HistoricalImportError("The 2026-09-02 revision must contain exactly 32 assigned blocks")
    _validate_confirmed_corrections(blocks, decisions.get("confirmed_corrections"), profiles)
    return blocks, {
        "physical_rows": 37, "logical_intervals": 36, "initial_blocks": 33,
        "current_blocks": 32, "replacements": 3,
    }


def _validate_confirmed_corrections(
    blocks: list[dict[str, Any]], corrections: object, profiles: dict[str, Any]
) -> None:
    if not isinstance(corrections, dict):
        raise HistoricalImportError("confirmed_corrections is required")
    required = {"mariana_ci", "marien_ci", "nataly_ci", "kenia_ci", "pedro_ci"}
    if set(corrections) != required:
        raise HistoricalImportError("confirmed_corrections must provide the five authoritative CI references")
    if any(not isinstance(corrections[key], str) or corrections[key] not in profiles for key in required):
        raise HistoricalImportError("Every confirmed correction CI must be an exact private-profile key")

    by_scope = {
        (block["subject_code"], block["group_code"]): block for block in blocks
    }

    def assignment(block: dict[str, Any], ci_key: str, start: date) -> dict[str, Any] | None:
        return next((
            item for item in block.get("assignments", [])
            if item["teacher_ci"] == corrections[ci_key] and item["effective_from"] == start
        ), None)

    m02 = by_scope.get(("MRF 0200", "M-02"))
    mariana_m02 = assignment(m02 or {}, "mariana_ci", CUTOVER_DATE) if m02 else None
    if mariana_m02 is None or len(mariana_m02["source_rows"]) != 2:
        raise HistoricalImportError("Anatomy II M-02 must collapse the duplicate into one Mariana interval")
    for group, prior_key in (("T-02", "marien_ci"), ("N-02", "nataly_ci")):
        block = by_scope.get(("MRF 0100", group))
        prior = assignment(block or {}, prior_key, INITIAL_DATE) if block else None
        replacement = assignment(block or {}, "mariana_ci", CUTOVER_DATE) if block else None
        if prior is None or prior["effective_to"] != date(2026, 9, 1) or replacement is None:
            raise HistoricalImportError(f"Anatomy I {group} must retain the confirmed 2026-09-02 replacement")
    m03 = by_scope.get(("MRF 0200", "M-03"))
    kenia_m03 = assignment(m03 or {}, "kenia_ci", INITIAL_DATE) if m03 else None
    if kenia_m03 is None or kenia_m03["effective_to"] != date(2026, 9, 1) or len(m03["assignments"]) != 1:
        raise HistoricalImportError("Anatomy II M-03 must end with no successor")
    if assignment(by_scope.get(("MRF 0200", "T-02")) or {}, "pedro_ci", INITIAL_DATE) is None:
        raise HistoricalImportError("Pedro Anatomy II T-02 must remain in semester 2")
    if assignment(by_scope.get(("CIR 0400", "M-04")) or {}, "pedro_ci", INITIAL_DATE) is None:
        raise HistoricalImportError("Pedro Surgery II M-04 must remain in semester 4")
    m01 = by_scope.get(("CIR 0400", "M-01"))
    if (
        assignment(m01 or {}, "kenia_ci", INITIAL_DATE) is None
        or m01["end_time"] != time(11, 20)
    ):
        raise HistoricalImportError("Kenia Surgery II M-01 must end at 11:20 AM")


def _state_value(value: Any) -> Any:
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, bytes):
        return {"sha256": _sha(value), "size": len(value)}
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    if isinstance(value, (dict, list)):
        return value
    return str(value)


def _state_rows(query: Any, model: type[Any], *, exclude: set[str] | None = None) -> list[dict[str, Any]]:
    excluded = exclude or set()
    primary_keys = list(model.__table__.primary_key.columns)
    rows = query.order_by(*primary_keys).all()
    return [
        {
            column.name: _state_value(getattr(row, column.name))
            for column in model.__table__.columns
            if column.name not in excluded
        }
        for row in rows
    ]


def _state_fingerprint(
    db: Session, profile_cis: list[str], blocks: list[dict[str, Any]], actor_id: int
) -> str:
    subject_names = [name for name, _semester in SUBJECTS.values()]
    classroom_codes = sorted({block["classroom"]["code"] for block in blocks})
    contract_lines = db.query(ContractLine).join(ContractDocument).filter(
        ContractDocument.academic_period == PERIOD,
        ContractLine.activity_kind == "practice",
        ContractLine.subject_label.in_(subject_names),
    )
    contract_ids = sorted({row.contract_id for row in contract_lines.all()})
    state = {
        "actor": _state_rows(db.query(User).filter(User.id == actor_id), User),
        "teachers": _state_rows(db.query(Teacher).filter(Teacher.ci.in_(profile_cis)), Teacher),
        "availability": _state_rows(db.query(TeacherAvailability).filter(
            TeacherAvailability.academic_period == PERIOD,
            TeacherAvailability.teacher_ci.in_(profile_cis),
        ), TeacherAvailability),
        "programs": _state_rows(db.query(AcademicProgram).filter(
            AcademicProgram.code == PROGRAM_CODE
        ), AcademicProgram),
        "subjects": _state_rows(db.query(AcademicSubject).filter(
            AcademicSubject.code.in_(SUBJECTS)
        ), AcademicSubject),
        "offerings": _state_rows(db.query(SubjectOffering).filter(
            SubjectOffering.academic_period == PERIOD
        ), SubjectOffering),
        "groups": _state_rows(db.query(AcademicGroup).filter(
            AcademicGroup.academic_period == PERIOD
        ), AcademicGroup),
        "classrooms": _state_rows(db.query(Classroom).filter(
            Classroom.code.in_(classroom_codes)
        ), Classroom),
        "drafts": _state_rows(db.query(AcademicScheduleDraft).filter(
            AcademicScheduleDraft.academic_period == PERIOD
        ), AcademicScheduleDraft),
        "blocks": _state_rows(db.query(AcademicScheduleBlock).join(AcademicScheduleDraft).filter(
            AcademicScheduleDraft.academic_period == PERIOD
        ), AcademicScheduleBlock),
        "assignments": _state_rows(db.query(AcademicScheduleAssignment).join(
            AcademicScheduleBlock
        ).join(AcademicScheduleDraft).filter(
            AcademicScheduleDraft.academic_period == PERIOD
        ), AcademicScheduleAssignment),
        "publications": _state_rows(db.query(AcademicSchedulePublication).filter(
            AcademicSchedulePublication.academic_period == PERIOD
        ), AcademicSchedulePublication),
        "published_blocks": _state_rows(db.query(AcademicSchedulePublishedBlock).join(
            AcademicSchedulePublication
        ).filter(AcademicSchedulePublication.academic_period == PERIOD), AcademicSchedulePublishedBlock),
        "published_assignments": _state_rows(db.query(AcademicSchedulePublishedAssignment).join(
            AcademicSchedulePublishedBlock
        ).join(AcademicSchedulePublication).filter(
            AcademicSchedulePublication.academic_period == PERIOD
        ), AcademicSchedulePublishedAssignment),
        "receipts": _state_rows(db.query(HistoricalScheduleImport).filter(
            HistoricalScheduleImport.academic_period == PERIOD
        ), HistoricalScheduleImport, exclude={"applied_state_digest"}),
        "legacy_practice": _state_rows(db.query(Designation).filter(
            Designation.academic_period == PERIOD,
            Designation.designation_type == "practice",
        ), Designation),
        "regular_attendance": _state_rows(db.query(AttendanceRecord).join(Designation).filter(
            AttendanceRecord.date >= INITIAL_DATE,
            Designation.academic_period == PERIOD,
            Designation.designation_type == "practice",
            Designation.subject.in_(subject_names),
        ), AttendanceRecord),
        "practice_attendance": _state_rows(db.query(PracticeAttendanceLog).join(Designation).filter(
            PracticeAttendanceLog.date >= INITIAL_DATE,
            Designation.academic_period == PERIOD,
            Designation.designation_type == "practice",
            Designation.subject.in_(subject_names),
        ), PracticeAttendanceLog),
        "practice_payroll": _state_rows(db.query(PracticePlanillaOutput).filter(
            (PracticePlanillaOutput.year > 2026)
            | ((PracticePlanillaOutput.year == 2026) & (PracticePlanillaOutput.month >= 8))
        ), PracticePlanillaOutput),
        "billing": _state_rows(db.query(BillingPublication).filter(
            (BillingPublication.year > 2026)
            | ((BillingPublication.year == 2026) & (BillingPublication.month >= 8))
        ), BillingPublication),
        "contract_documents": _state_rows(db.query(ContractDocument).filter(
            ContractDocument.id.in_(contract_ids)
        ), ContractDocument),
        "contract_lines": _state_rows(contract_lines, ContractLine),
        "audit": _state_rows(db.query(ActivityLog).filter(
            ActivityLog.action == "historical_assistential_import"
        ), ActivityLog),
    }
    return _sha(_json_bytes(state))


def _operational_blockers(db: Session) -> dict[str, int]:
    subject_names = [name for name, _semester in SUBJECTS.values()]
    counts = {
        "regular_attendance": db.query(AttendanceRecord.id).join(Designation).filter(
            AttendanceRecord.date >= INITIAL_DATE,
            Designation.academic_period == PERIOD,
            Designation.designation_type == "practice",
            Designation.subject.in_(subject_names),
        ).count(),
        "practice_attendance": db.query(PracticeAttendanceLog.id).join(Designation).filter(
            PracticeAttendanceLog.date >= INITIAL_DATE,
            Designation.academic_period == PERIOD,
            Designation.designation_type == "practice",
            Designation.subject.in_(subject_names),
        ).count(),
        "practice_payroll": db.query(PracticePlanillaOutput.id).filter(
            (PracticePlanillaOutput.year > 2026) | ((PracticePlanillaOutput.year == 2026) & (PracticePlanillaOutput.month >= 8))
        ).count(),
        "billing": db.query(BillingPublication.id).filter(
            BillingPublication.planilla_type == "practice",
            (BillingPublication.year > 2026) | ((BillingPublication.year == 2026) & (BillingPublication.month >= 8))
        ).count(),
        "contracts": db.query(ContractLine.id).join(ContractDocument).filter(
            ContractDocument.academic_period == PERIOD,
            ContractLine.activity_kind == "practice",
            ContractLine.subject_label.in_(subject_names),
        ).count(),
    }
    return {name: count for name, count in counts.items() if count}


def _target_state_errors(db: Session) -> list[str]:
    errors: list[str] = []
    target_counts = {
        "program": db.query(AcademicProgram).filter(AcademicProgram.code == PROGRAM_CODE).count(),
        "subjects": db.query(AcademicSubject).filter(AcademicSubject.code.in_(SUBJECTS)).count(),
        "period_offerings": db.query(SubjectOffering).filter(SubjectOffering.academic_period == PERIOD).count(),
        "period_groups": db.query(AcademicGroup).filter(AcademicGroup.academic_period == PERIOD).count(),
        "period_drafts": db.query(AcademicScheduleDraft).filter(AcademicScheduleDraft.academic_period == PERIOD).count(),
        "period_publications": db.query(AcademicSchedulePublication).filter(
            AcademicSchedulePublication.academic_period == PERIOD
        ).count(),
    }
    if any(target_counts.values()):
        errors.append("Target catalogs/drafts/publications already exist without this import receipt; partial merge is refused")
    if db.query(Designation.id).filter(
        Designation.academic_period == PERIOD,
        Designation.designation_type == "practice",
    ).first():
        errors.append("Legacy practice designations already exist; synthetic Designation merge is refused")
    return errors


def build_preview(
    db: Session, inputs: ImportInputs, *, actor_ci: str, policy: str,
    effective_date: date, historical_availability_unrecorded: bool,
) -> dict[str, Any]:
    if not historical_availability_unrecorded:
        raise HistoricalImportError("--historical-availability-unrecorded is required")
    if effective_date != INITIAL_DATE or effective_date >= date.today():
        raise HistoricalImportError("The explicit effective date must be the past date 2026-08-20")
    if policy != POLICY:
        raise HistoricalImportError(f"--policy must be exactly {POLICY}")
    actor = db.query(User).filter(User.ci == actor_ci).one_or_none()
    if actor is None or actor.role != "admin" or not actor.is_active:
        raise HistoricalImportError("An active administrator actor is required")
    raw, source_hashes = _read_inputs(inputs)
    profiles = _load_json(raw["profiles"], "profiles")
    decisions = _load_json(raw["decisions"], "decisions")
    _validate_curriculum(raw["curriculum"])
    designation_evidence = _sheet_evidence(raw["designation"], "DOCENTES ASISTENCIALES")
    enrollment_evidence = _sheet_evidence(raw["enrollment"], "GRUPOS")
    blocks, counts = _parse_plan(decisions, designation_evidence, enrollment_evidence, profiles)
    profile_changes, profile_conflicts = _profile_plan(db, profiles)
    blockers = _operational_blockers(db)
    errors = _target_state_errors(db)
    if profile_conflicts:
        errors.append("Teacher profile conflicts exist for: " + ", ".join(profile_conflicts))
    if blockers:
        errors.append("Retroactive operational evidence exists; import is fail-closed")
    state_digest = _state_fingerprint(db, sorted(profiles), blocks, actor.id)
    input_digest = _sha(_json_bytes({
        "schema_version": 2, "academic_period": PERIOD,
        "effective_date": effective_date.isoformat(), "policy": policy,
        "historical_availability_unrecorded": True, "actor_id": actor.id,
        "source_hashes": source_hashes,
    }))
    idempotency_key = _sha(_json_bytes({
        "period": PERIOD, "effective_date": effective_date.isoformat(),
        "policy": policy, "sources": source_hashes,
    }))
    existing = db.query(HistoricalScheduleImport).filter_by(idempotency_key=idempotency_key).one_or_none()
    safe_profiles = [
        {"ci_masked": item["ci_masked"], "action": item["action"], "fields": item["fields"]}
        for item in profile_changes
    ]
    canonical = {
        "schema_version": 2, "academic_period": PERIOD,
        "effective_date": effective_date.isoformat(), "cutover_date": CUTOVER_DATE.isoformat(),
        "historical_availability_unrecorded": True, "actor_id": actor.id,
        "policy": policy, "policy_description": POLICY_DESCRIPTION,
        "source_hashes": source_hashes, "idempotency_key": idempotency_key,
        "input_digest": input_digest, "state_digest": state_digest,
        "counts": counts, "profiles": safe_profiles,
        "operational_blockers": blockers, "errors": errors,
        "lineage": {
            "designation_rows": decisions["physical_source_rows"],
            "red_rows": decisions["red_source_rows"],
            "designation_cells": {
                str(row): designation_evidence[row]["cells"] for row in decisions["physical_source_rows"]
            },
        },
        "replay": existing.result if existing else None,
    }
    digest = _sha(_json_bytes(canonical))
    canonical["digest"] = digest
    canonical["can_apply"] = not errors
    canonical["_private"] = {
        "actor": actor, "blocks": blocks, "profile_changes": profile_changes,
        "profiles": profiles, "decisions": decisions,
    }
    return canonical


def safe_preview(preview: dict[str, Any]) -> dict[str, Any]:
    return {key: value for key, value in preview.items() if key != "_private"}


def _acquire_lock(db: Session, key: str) -> None:
    if db.bind is not None and db.bind.dialect.name == "postgresql":
        signed = int(key[:16], 16)
        if signed >= 2**63:
            signed -= 2**64
        db.execute(text("SELECT pg_advisory_xact_lock(:key)"), {"key": signed})


def _apply_profiles(db: Session, changes: list[dict[str, Any]]) -> None:
    for change in changes:
        raw_profile = change["profile"]
        teacher = db.get(Teacher, change["ci"])
        if teacher is None:
            teacher = Teacher(ci=change["ci"], full_name=raw_profile["full_name"])
            db.add(teacher)
        for field_name in PROFILE_FIELDS:
            if change["fields"].get(field_name) in {"create", "fill"} and raw_profile.get(field_name) not in (None, ""):
                setattr(teacher, field_name, raw_profile[field_name])
    db.flush()


def _create_publication(
    db: Session, *, draft: AcademicScheduleDraft, blocks: list[dict[str, Any]],
    catalog: dict[str, Any], actor_id: int, effective_from: date, snapshot_date: date,
) -> AcademicSchedulePublication:
    draft_blocks: list[tuple[AcademicScheduleBlock, dict[str, Any]]] = []
    for item in blocks:
        block = AcademicScheduleBlock(
            draft_id=draft.id, offering_id=catalog["offerings"][item["subject_code"]].id,
            group_id=catalog["groups"][(item["semester"], item["group_code"])].id,
            classroom_id=catalog["classrooms"][item["classroom"]["code"]].id,
            activity_type="practice", weekday=item["weekday"],
            start_time=item["start_time"], end_time=item["end_time"],
        )
        db.add(block)
        db.flush()
        selected = [assignment for assignment in item["assignments"] if (
            assignment["effective_from"] <= snapshot_date
            and (assignment["effective_to"] is None or assignment["effective_to"] >= snapshot_date)
        )]
        if len(selected) != 1:
            raise HistoricalImportError("Each publication block must have exactly one snapshot owner")
        for assignment in selected:
            db.add(AcademicScheduleAssignment(
                block_id=block.id, teacher_ci=assignment["teacher_ci"],
                effective_from=assignment["effective_from"], effective_to=assignment["effective_to"],
            ))
        draft_blocks.append((block, item))
    db.flush()
    content_digest = _sha(_json_bytes({
        "effective_from": effective_from.isoformat(),
        "blocks": [
            {
                "key": item["key"], "subject": item["subject_code"], "group": item["group_code"],
                "weekday": item["weekday"], "start": item["start_time"].isoformat(),
                "end": item["end_time"].isoformat(),
            }
            for _block, item in draft_blocks
        ],
    }))
    publication = AcademicSchedulePublication(
        program_id=draft.program_id, academic_period=PERIOD, effective_from=effective_from,
        sequence=1, content_digest=content_digest, source_draft_id=draft.id, created_by=actor_id,
    )
    db.add(publication)
    db.flush()
    for block, item in draft_blocks:
        published = AcademicSchedulePublishedBlock(
            publication_id=publication.id, source_block_id=block.id,
            source_offering_id=block.offering_id,
            source_subject_id=catalog["subjects"][item["subject_code"]].id,
            source_group_id=block.group_id, source_classroom_id=block.classroom_id,
            subject_code=item["subject_code"], subject_name=item["subject_name"],
            group_code=item["group_code"], semester=item["semester"],
            classroom_code=item["classroom"]["code"], classroom_name=item["classroom"]["name"],
            activity_type="practice", weekday=item["weekday"], start_time=item["start_time"],
            end_time=item["end_time"], notes=json.dumps({
                "source_sheet": "DOCENTES ASISTENCIALES", "source_rows": item["source_rows"],
                "source_subject_label": item["source_subject_label"],
                "official_workload": 60, "source_real": item["source_real"],
            }, ensure_ascii=False, sort_keys=True),
        )
        db.add(published)
        db.flush()
        assignments = db.query(AcademicScheduleAssignment).filter_by(block_id=block.id).order_by(
            AcademicScheduleAssignment.effective_from
        ).all()
        for assignment in assignments:
            teacher = db.get(Teacher, assignment.teacher_ci)
            db.add(AcademicSchedulePublishedAssignment(
                publication_block_id=published.id, source_assignment_id=assignment.id,
                teacher_ci=assignment.teacher_ci, teacher_name=teacher.full_name,
                effective_from=assignment.effective_from, effective_to=assignment.effective_to,
            ))
    draft.status = "published"
    db.flush()
    return publication


def apply_import(
    db: Session, inputs: ImportInputs, *, actor_ci: str, policy: str,
    effective_date: date, historical_availability_unrecorded: bool,
    expected_digest: str,
) -> dict[str, Any]:
    # The caller owns the single outer transaction.
    initial = build_preview(
        db, inputs, actor_ci=actor_ci, policy=policy, effective_date=effective_date,
        historical_availability_unrecorded=historical_availability_unrecorded,
    )
    # Serialize every candidate for this target, even when different source files
    # produce different receipt keys.
    _acquire_lock(db, IMPORT_LOCK_KEY)
    current = build_preview(
        db, inputs, actor_ci=actor_ci, policy=policy, effective_date=effective_date,
        historical_availability_unrecorded=historical_availability_unrecorded,
    )
    replay = db.query(HistoricalScheduleImport).filter_by(
        idempotency_key=initial["idempotency_key"]
    ).one_or_none()
    if replay is not None:
        bindings_match = (
            hmac.compare_digest(expected_digest, replay.preview_digest)
            and hmac.compare_digest(current["input_digest"], replay.input_digest)
            and hmac.compare_digest(current["state_digest"], replay.applied_state_digest)
            and current["policy"] == replay.policy
            and current["source_hashes"] == replay.source_hashes
            and current["actor_id"] == replay.actor_id
            and current["academic_period"] == replay.academic_period
            and current["effective_date"] == replay.effective_from.isoformat()
        )
        if not bindings_match:
            raise HistoricalImportError("Replay binding does not match the original import receipt")
        return {**replay.result, "replayed": True}
    preview = current
    if not hmac.compare_digest(preview["digest"], expected_digest):
        raise HistoricalImportError("Inputs or database state changed after preview")
    if not preview["can_apply"]:
        raise HistoricalImportError("Import is blocked: " + "; ".join(preview["errors"]))
    private = preview["_private"]
    actor: User = private["actor"]
    _apply_profiles(db, private["profile_changes"])

    program = AcademicProgram(code=PROGRAM_CODE, name=PROGRAM_NAME, active=True)
    db.add(program)
    db.flush()
    subjects: dict[str, AcademicSubject] = {}
    offerings: dict[str, SubjectOffering] = {}
    for code, (name, semester) in SUBJECTS.items():
        subject = AcademicSubject(code=code, name=name, active=True)
        db.add(subject)
        db.flush()
        subjects[code] = subject
        offering = SubjectOffering(
            subject_id=subject.id, program_id=program.id, academic_period=PERIOD,
            semester=semester, theory_hours=0, practice_hours=60, active=True,
        )
        db.add(offering)
        db.flush()
        offerings[code] = offering
    groups: dict[tuple[int, str], AcademicGroup] = {}
    classrooms: dict[str, Classroom] = {}
    for block in private["blocks"]:
        group_key = (block["semester"], block["group_code"])
        if group_key not in groups:
            group = AcademicGroup(
                program_id=program.id, academic_period=PERIOD, semester=block["semester"],
                shift=block["group_code"].split("-")[0], code=block["group_code"],
                expected_size=None, active=True,
            )
            db.add(group)
            db.flush()
            groups[group_key] = group
        classroom_data = block["classroom"]
        if classroom_data["code"] not in classrooms:
            classroom = Classroom(
                code=classroom_data["code"], name=classroom_data["name"], campus="Cobija",
                capacity=classroom_data["capacity"], classroom_type=classroom_data["type"],
                resources=[], active=True,
            )
            db.add(classroom)
            db.flush()
            classrooms[classroom_data["code"]] = classroom
    catalog = {"subjects": subjects, "offerings": offerings, "groups": groups, "classrooms": classrooms}
    initial_draft = AcademicScheduleDraft(
        program_id=program.id, academic_period=PERIOD,
        name="Historical assistential import — 2026-08-20", normalized_name="historical assistential import — 2026-08-20",
        status="draft",
    )
    current_draft = AcademicScheduleDraft(
        program_id=program.id, academic_period=PERIOD,
        name="Historical assistential import — 2026-09-02", normalized_name="historical assistential import — 2026-09-02",
        status="draft",
    )
    db.add_all([initial_draft, current_draft])
    db.flush()
    initial_publication = _create_publication(
        db, draft=initial_draft, blocks=private["blocks"], catalog=catalog,
        actor_id=actor.id, effective_from=INITIAL_DATE, snapshot_date=INITIAL_DATE,
    )
    current_blocks = [
        block for block in private["blocks"]
        if any(item["effective_from"] <= CUTOVER_DATE and (
            item["effective_to"] is None or item["effective_to"] >= CUTOVER_DATE
        ) for item in block["assignments"])
    ]
    current_publication = _create_publication(
        db, draft=current_draft, blocks=current_blocks, catalog=catalog,
        actor_id=actor.id, effective_from=CUTOVER_DATE, snapshot_date=CUTOVER_DATE,
    )
    result = {
        "academic_period": PERIOD, "preview_digest": preview["digest"],
        "idempotency_key": preview["idempotency_key"],
        "initial_publication_id": initial_publication.id,
        "current_publication_id": current_publication.id,
        "counts": preview["counts"], "replayed": False,
    }
    receipt = HistoricalScheduleImport(
        idempotency_key=preview["idempotency_key"], preview_digest=preview["digest"],
        input_digest=preview["input_digest"], pre_state_digest=preview["state_digest"],
        applied_state_digest="",
        academic_period=PERIOD, effective_from=INITIAL_DATE, actor_id=actor.id,
        policy=preview["policy"], source_hashes=preview["source_hashes"],
        counts=preview["counts"], result=result,
    )
    db.add(receipt)
    db.add(ActivityLog(
        user_id=actor.id, user_ci=None, user_name=None, user_role="admin",
        action="historical_assistential_import", category="academic_schedule",
        description="Applied digest-bound II/2026 historical assistential schedule import",
        details={
            "actor_id": actor.id, "source_sha256": preview["source_hashes"],
            "preview_digest": preview["digest"], "policy": preview["policy"],
            "policy_description": POLICY_DESCRIPTION,
            "counts": preview["counts"], "historical_availability_unrecorded": True,
        }, status="success",
    ))
    db.flush()
    receipt.applied_state_digest = _state_fingerprint(
        db, sorted(private["profiles"]), private["blocks"], actor.id
    )
    db.flush()
    return result
